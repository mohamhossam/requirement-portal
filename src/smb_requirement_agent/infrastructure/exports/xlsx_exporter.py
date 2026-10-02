"""Readable, formula-safe XLSX adapter for the neutral export contract."""

from __future__ import annotations

import re
from collections.abc import Iterable
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from smb_requirement_agent.application.errors import BacklogExportFormatError
from smb_requirement_agent.application.exports import (
    ExportArchitecture,
    ExportFormat,
    ExportJourneyNeighbour,
    NeutralBacklogExport,
)

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_MAX_CELL_CHARACTERS = 32_767
type CellValue = str | int | bool | None


class XlsxBacklogExporter:
    @property
    def format(self) -> ExportFormat:
        return ExportFormat.XLSX

    @property
    def extension(self) -> str:
        return "xlsx"

    @property
    def media_type(self) -> str:
        return XLSX_MEDIA_TYPE

    def render(self, document: NeutralBacklogExport) -> bytes:
        workbook = Workbook()
        workbook.remove(workbook.worksheets[0])
        workbook.properties.creator = "SMB AI Requirement Breakdown Agent"
        workbook.properties.created = document.manifest.revision_created_at
        workbook.properties.modified = document.manifest.revision_created_at
        self._manifest(workbook, document)
        self._epic(workbook, document)
        self._features(workbook, document)
        self._stories(workbook, document)
        self._criteria(workbook, document)
        self._systems(workbook, document)
        self._dependencies(workbook, document)
        self._connected(workbook, document)
        self._product_context(workbook, document)
        self._journey_steps(workbook, document)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _manifest(self, workbook: Workbook, document: NeutralBacklogExport) -> None:
        value = document.manifest
        approval = value.final_approval
        sheet = workbook.create_sheet("Manifest")
        _append(sheet, ("Field", "Value"), header=True)
        rows: tuple[tuple[CellValue, CellValue], ...] = (
            ("schema_version", document.schema_version),
            ("requirement_id", value.requirement_id),
            ("breakdown_revision", value.breakdown_revision),
            ("revision_created_at", _iso(value.revision_created_at)),
            ("final_approval_id", approval.id),
            ("subject_fingerprint", approval.subject_fingerprint),
            ("approved_by_id", approval.recorded_by.id),
            ("approved_by_name", approval.recorded_by.display_name),
            ("approved_by_email", approval.recorded_by.email),
            ("approved_at", _iso(approval.recorded_at)),
            ("approval_rationale", approval.rationale),
            ("epic_count", value.counts.epics),
            ("feature_count", value.counts.features),
            ("story_count", value.counts.stories),
            ("acceptance_criterion_count", value.counts.acceptance_criteria),
        )
        for row in rows:
            _append(sheet, row)
        _finish(sheet)

    def _epic(self, workbook: Workbook, document: NeutralBacklogExport) -> None:
        epic = document.epic
        sheet = workbook.create_sheet("Epic")
        _append(
            sheet,
            (
                "id",
                "requirement_id",
                "name",
                "outcome",
                "business_case",
                "status",
                "generated_at",
                "model",
                "prompt_version",
            ),
            header=True,
        )
        _append(
            sheet,
            (
                epic.id,
                epic.requirement_id,
                epic.name,
                epic.outcome,
                epic.business_case,
                epic.status,
                _iso(epic.provenance.generated_at),
                epic.provenance.model,
                epic.provenance.prompt_version,
            ),
        )
        _finish(sheet)

    def _features(self, workbook: Workbook, document: NeutralBacklogExport) -> None:
        sheet = workbook.create_sheet("Features")
        _append(
            sheet,
            (
                "sequence",
                "id",
                "epic_id",
                "name",
                "outcome",
                "delivery_drop",
                "splitting_pattern",
                "splitting_rationale",
                "status",
                "generated_at",
                "model",
                "prompt_version",
            ),
            header=True,
        )
        for feature in document.epic.features:
            _append(
                sheet,
                (
                    feature.sequence,
                    feature.id,
                    feature.epic_id,
                    feature.name,
                    feature.outcome,
                    feature.delivery_drop,
                    feature.splitting_pattern,
                    feature.splitting_rationale,
                    feature.status,
                    _iso(feature.provenance.generated_at),
                    feature.provenance.model,
                    feature.provenance.prompt_version,
                ),
            )
        _finish(sheet)

    def _stories(self, workbook: Workbook, document: NeutralBacklogExport) -> None:
        sheet = workbook.create_sheet("Stories")
        _append(
            sheet,
            (
                "feature_sequence",
                "sequence",
                "id",
                "feature_id",
                "role",
                "action",
                "value",
                "voice",
                "status",
                "generated_at",
                "model",
                "prompt_version",
            ),
            header=True,
        )
        for feature in document.epic.features:
            for story in feature.stories:
                _append(
                    sheet,
                    (
                        feature.sequence,
                        story.sequence,
                        story.id,
                        story.feature_id,
                        story.role,
                        story.action,
                        story.value,
                        story.voice,
                        story.status,
                        _iso(story.provenance.generated_at),
                        story.provenance.model,
                        story.provenance.prompt_version,
                    ),
                )
        _finish(sheet)

    def _criteria(self, workbook: Workbook, document: NeutralBacklogExport) -> None:
        sheet = workbook.create_sheet("Acceptance Criteria")
        _append(
            sheet,
            ("feature_id", "story_id", "sequence", "given", "when", "then"),
            header=True,
        )
        for feature in document.epic.features:
            for story in feature.stories:
                for criterion in story.acceptance_criteria:
                    _append(
                        sheet,
                        (
                            feature.id,
                            story.id,
                            criterion.sequence,
                            criterion.given,
                            criterion.when,
                            criterion.then,
                        ),
                    )
        _finish(sheet)

    def _systems(self, workbook: Workbook, document: NeutralBacklogExport) -> None:
        sheet = workbook.create_sheet("Systems")
        _append(
            sheet,
            (
                "artifact_type",
                "artifact_id",
                "knowledge_version",
                "mapped_at",
                "system_id",
                "system_name",
                "catalogued",
                "squad_ids",
                "squad_names",
                "value_stream_names",
                "product_names",
                "capability_ids",
                "capability_names",
                "capability_domains",
                "capability_components",
            ),
            header=True,
        )
        for kind, item_id, architecture in _architectures(document):
            for system in architecture.systems:
                _append(
                    sheet,
                    (
                        kind,
                        item_id,
                        architecture.knowledge_version,
                        _iso(architecture.mapped_at),
                        system.id,
                        system.name,
                        system.catalogued,
                        "; ".join(item.id for item in system.squads),
                        "; ".join(item.name for item in system.squads),
                        "; ".join(item.name for item in system.value_streams),
                        "; ".join(item.name for item in system.products),
                        "; ".join(item.id for item in system.capabilities),
                        "; ".join(item.name for item in system.capabilities),
                        "; ".join(
                            " > ".join(item.domain_path)
                            for item in system.capabilities
                            if item.domain_path
                        ),
                        "; ".join(
                            f"{item.name}: {item.component}"
                            for item in system.capabilities
                            if item.component
                        ),
                    ),
                )
        _finish(sheet)

    def _dependencies(self, workbook: Workbook, document: NeutralBacklogExport) -> None:
        sheet = workbook.create_sheet("Dependencies")
        _append(
            sheet,
            (
                "artifact_type",
                "artifact_id",
                "source_system_id",
                "target_system_id",
                "description",
                "kind",
            ),
            header=True,
        )
        for kind, item_id, architecture in _architectures(document):
            for dependency in architecture.dependencies:
                _append(
                    sheet,
                    (
                        kind,
                        item_id,
                        dependency.source_system_id,
                        dependency.target_system_id,
                        dependency.description,
                        dependency.kind,
                    ),
                )
        _finish(sheet)

    def _connected(self, workbook: Workbook, document: NeutralBacklogExport) -> None:
        """One row per relationship that connects a mapped system to one outside the mapping."""
        sheet = workbook.create_sheet("Connected Systems")
        _append(
            sheet,
            (
                "artifact_type",
                "artifact_id",
                "system_id",
                "system_name",
                "squad_names",
                "mapped_system_id",
                "direction",
                "description",
                "kind",
            ),
            header=True,
        )
        for kind, item_id, architecture in _architectures(document):
            systems = {item.id: item for item in architecture.adjacent_systems}
            for dependency in architecture.adjacent_dependencies:
                upstream = dependency.source_system_id in systems
                system = systems[
                    dependency.source_system_id if upstream else dependency.target_system_id
                ]
                _append(
                    sheet,
                    (
                        kind,
                        item_id,
                        system.id,
                        system.name,
                        "; ".join(item.name for item in system.squads),
                        dependency.target_system_id if upstream else dependency.source_system_id,
                        "depends_on_mapped" if upstream else "mapped_depends_on",
                        dependency.description,
                        dependency.kind,
                    ),
                )
        _finish(sheet)

    def _product_context(self, workbook: Workbook, document: NeutralBacklogExport) -> None:
        """One row per system responsibility of an offering the item names (ADR-0097)."""
        sheet = workbook.create_sheet("Product Context")
        _append(
            sheet,
            (
                "artifact_type",
                "artifact_id",
                "product_id",
                "product_name",
                "order_type",
                "matched_terms",
                "component_name",
                "system_id",
                "system_name",
                "role",
                "responsibility",
            ),
            header=True,
        )
        for kind, item_id, architecture in _architectures(document):
            for context in architecture.product_contexts:
                facts = (
                    kind,
                    item_id,
                    context.product_id,
                    context.product_name,
                    context.order_type,
                    "; ".join(context.matched_terms),
                )
                if not context.responsibilities:
                    _append(sheet, (*facts, None, None, None, None, None))
                for duty in context.responsibilities:
                    _append(
                        sheet,
                        (
                            *facts,
                            duty.component_name,
                            duty.system_id,
                            duty.system_name,
                            duty.role,
                            duty.description,
                        ),
                    )
        _finish(sheet)

    def _journey_steps(self, workbook: Workbook, document: NeutralBacklogExport) -> None:
        """One row per journey activity a mapped system performs or supports (ADR-0097)."""
        sheet = workbook.create_sheet("Journey Steps")
        _append(
            sheet,
            (
                "artifact_type",
                "artifact_id",
                "system_id",
                "journey_id",
                "journey_name",
                "fulfils",
                "activity_number",
                "activity_name",
                "part",
                "before",
                "after",
            ),
            header=True,
        )
        for kind, item_id, architecture in _architectures(document):
            for step in architecture.journey_steps:
                _append(
                    sheet,
                    (
                        kind,
                        item_id,
                        step.system_id,
                        step.journey_id,
                        step.journey_name,
                        " > ".join(step.fulfils),
                        step.number,
                        step.name,
                        "performs" if step.performs else "supports",
                        "; ".join(_activity(other) for other in step.before),
                        "; ".join(_activity(other) for other in step.after),
                    ),
                )
        _finish(sheet)


def _activity(item: ExportJourneyNeighbour) -> str:
    named = f"{item.number}. {item.name}"
    return f"{named} ({item.system_name})" if item.system_name else named


def _architectures(
    document: NeutralBacklogExport,
) -> Iterable[tuple[str, str, ExportArchitecture]]:
    for feature in document.epic.features:
        if feature.architecture is not None:
            yield "feature", feature.id, feature.architecture
        for story in feature.stories:
            if story.architecture is not None:
                yield "story", story.id, story.architecture


def _iso(value: datetime) -> str:
    return value.astimezone(UTC).isoformat().replace("+00:00", "Z")


def _validated(value: CellValue) -> CellValue:
    if not isinstance(value, str):
        return value
    if len(value) > _MAX_CELL_CHARACTERS or _ILLEGAL_XML.search(value):
        raise BacklogExportFormatError(
            "This content cannot be represented losslessly in XLSX; export the revision as JSON."
        )
    return value


def _append(sheet: Worksheet, values: tuple[CellValue, ...], *, header: bool = False) -> None:
    row_number = 1 if sheet.max_row == 1 and sheet.cell(1, 1).value is None else sheet.max_row + 1
    for column, raw_value in enumerate(values, start=1):
        value = _validated(raw_value)
        cell: Cell = sheet.cell(row=row_number, column=column, value=value)
        if isinstance(value, str):
            cell.data_type = "s"
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if header:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="274C5B")


def _finish(sheet: Worksheet, *, filter_rows: bool = True) -> None:
    sheet.freeze_panes = "A2"
    if filter_rows and sheet.max_row >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    for column_index, column_cells in enumerate(sheet.columns, start=1):
        width = min(
            60,
            max(12, *(len(str(cell.value)) + 2 for cell in column_cells if cell.value is not None)),
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = width
