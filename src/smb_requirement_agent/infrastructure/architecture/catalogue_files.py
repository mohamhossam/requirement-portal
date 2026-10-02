"""Excel, YAML and JSON catalogue files, all mapping to one catalogue structure.

YAML and JSON share the packaged catalogue's shape (``systems`` and
``dependencies``). Excel spreads the same content over one sheet per kind so
people can fill it in without learning a data format. Every parse error names
where it happened (sheet and row, or entry), because a maintainer has to find
and fix it in the file.
"""

from __future__ import annotations

import io
import json
import zipfile
from collections.abc import Callable, Iterator, Sequence
from functools import partial
from typing import Any

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException
from smb_kernel.documents.extraction_base import (
    MAX_OFFICE_PARTS,
    MAX_OFFICE_UNCOMPRESSED_BYTES,
    MAX_WORKBOOK_ROWS,
)

from smb_requirement_agent.application.ports.catalogue_file import (
    CatalogueContent,
    CatalogueFileFormat,
)
from smb_requirement_agent.domain.architecture.journeys import (
    Activity,
    ActivityIntegration,
    FlowRule,
    Journey,
    flow_rule_kind,
)
from smb_requirement_agent.domain.architecture.knowledge import (
    ArchitectureKnowledge,
    CapabilityDomain,
    InvalidKnowledgeError,
    KnowledgeCapability,
    LandscapeDomain,
    RelationshipKind,
    SystemComponent,
    SystemDefinition,
    SystemRelationship,
)
from smb_requirement_agent.domain.architecture.products import (
    ComponentResponsibility,
    OfferingComponent,
    OfferingPoint,
    OrderType,
    ProductOffering,
    SourceConfidence,
)

SYSTEMS = "Systems"
COMPONENTS = "Components"
CAPABILITIES = "Capabilities"
CONSTRAINTS = "Constraints"
RELATIONSHIPS = "Relationships"
DOMAINS = "Domains"
LANDSCAPE = "LandscapeDomains"
PRODUCTS = "Products"
ORDER_TYPES = "OrderTypes"
OFFERING_COMPONENTS = "OfferingComponents"
RESPONSIBILITIES = "Responsibilities"
PRODUCT_POINTS = "ProductPoints"
JOURNEYS = "Journeys"
ACTIVITIES = "Activities"
FLOW_RULES = "FlowRules"
ACTIVITY_INTEGRATIONS = "ActivityIntegrations"
INSTRUCTIONS = "Instructions"
_DOMAIN_HEADERS = ("domain_id", "name", "name_ar", "parent_id", "description")
_HEADERS: dict[str, tuple[str, ...]] = {
    SYSTEMS: ("system_id", "name", "name_ar", "aliases", "description", "landscape_domain_id"),
    COMPONENTS: (
        "system_id",
        "component_id",
        "name",
        "name_ar",
        "aliases",
        "technology",
        "description",
    ),
    CAPABILITIES: (
        "system_id",
        "capability_id",
        "name",
        "triggers",
        "domain_id",
        "component_id",
    ),
    CONSTRAINTS: ("system_id", "constraint"),
    RELATIONSHIPS: ("source_system_id", "target_system_id", "description", "kind"),
    DOMAINS: _DOMAIN_HEADERS,
    LANDSCAPE: _DOMAIN_HEADERS,
    PRODUCTS: (
        "product_id",
        "name",
        "code",
        "family",
        "version",
        "lifecycle",
        "proposition",
        "rules",
        "confidence",
        "source",
    ),
    ORDER_TYPES: ("product_id", "code", "name", "enabled", "description", "confidence", "source"),
    OFFERING_COMPONENTS: (
        "product_id",
        "component_id",
        "name",
        "code",
        "type",
        "mandatory",
        "customer_visible",
        "description",
        "commercial_spec",
        "technical_spec",
        "technical_details",
        "confidence",
        "source",
    ),
    RESPONSIBILITIES: (
        "product_id",
        "component_id",
        "system_id",
        "role",
        "description",
        "order_types",
        "confidence",
        "source",
    ),
    PRODUCT_POINTS: ("product_id", "kind", "name", "description", "confidence", "source"),
    JOURNEYS: (
        "journey_id",
        "name",
        "product_id",
        "order_type_code",
        "description",
        "confidence",
        "source",
    ),
    ACTIVITIES: (
        "journey_id",
        "number",
        "name",
        "phase",
        "track",
        "performing_system_id",
        "supporting_system_ids",
        "system_function",
        "mode",
        "customer_visible",
        "description",
        "component_ids",
        "input",
        "output",
        "etom",
        "confidence",
        "source",
    ),
    FLOW_RULES: (
        "journey_id",
        "kind",
        "from_activity",
        "to_activity",
        "condition",
        "branch",
        "parallel_group",
        "rejoin_at",
        "confidence",
        "source",
    ),
    ACTIVITY_INTEGRATIONS: (
        "journey_id",
        "from_activity",
        "to_activity",
        "interaction",
        "interface",
        "payload",
        "timing",
        "correlation_key",
        "confidence",
        "source",
    ),
}
# Headers a sheet cannot do without. Columns added later stay optional, so
# workbooks filled from an older template still import.
_REQUIRED_HEADERS: dict[str, tuple[str, ...]] = {
    SYSTEMS: ("system_id", "name"),
    COMPONENTS: ("system_id", "component_id", "name"),
    CAPABILITIES: ("system_id", "capability_id", "name", "triggers"),
    CONSTRAINTS: _HEADERS[CONSTRAINTS],
    RELATIONSHIPS: ("source_system_id", "target_system_id", "description"),
    DOMAINS: ("domain_id", "name"),
    LANDSCAPE: ("domain_id", "name"),
    PRODUCTS: ("product_id", "name"),
    ORDER_TYPES: ("product_id", "code", "name"),
    OFFERING_COMPONENTS: ("product_id", "component_id", "name"),
    RESPONSIBILITIES: ("product_id", "component_id", "system_id", "role", "description"),
    PRODUCT_POINTS: ("product_id", "kind", "name"),
    JOURNEYS: ("journey_id", "name"),
    ACTIVITIES: ("journey_id", "number", "name"),
    FLOW_RULES: ("journey_id", "kind", "from_activity", "to_activity"),
    ACTIVITY_INTEGRATIONS: ("journey_id", "from_activity", "to_activity"),
}
_KINDS = ", ".join(kind.value for kind in RelationshipKind)
_REQUIRED_SHEETS = (SYSTEMS,)
_LIST_SEPARATOR = ";"
_INSTRUCTIONS = (
    ("How to fill this catalogue",),
    ("Systems: one row per system. system_id is a short stable key such as bcrm.",),
    ("aliases and triggers: several values separated by semicolons (;).",),
    ("Capabilities: one row per capability; triggers are the phrases that point to it.",),
    ("Constraints: one row per constraint of a system.",),
    ("Relationships: one row per dependency between two systems listed on Systems.",),
    (f"kind (optional): how the source depends on the target, one of {_KINDS}.",),
    ("Domains (optional): business areas such as Order capture; parent_id nests one in another.",),
    ("Capabilities domain_id (optional): the domain a capability belongs to.",),
    ("Components (optional): the parts of a system, such as a module or service.",),
    ("Capabilities component_id (optional): the component of its system that delivers it.",),
    (
        "LandscapeDomains (optional): where systems sit, such as Customer; parent_id makes a "
        "sub-domain such as Customer > Assisted.",
    ),
    ("Systems description and landscape_domain_id (optional): what it is for, and where.",),
    ("Products (optional): product offerings such as Business Pro Plus; rules separated by ;.",),
    ("OrderTypes: how each offering is ordered; enabled is yes or no.",),
    ("OfferingComponents: the parts of each offering; mandatory and customer_visible yes or no.",),
    (
        "Responsibilities: which system does what for a component, in which role; order_types "
        "names order type codes separated by ;.",
    ),
    ("ProductPoints: kind value (what it gives customers) or audience (who it is for).",),
    ("confidence (optional): confirmed, inferred or gap, as the source says.",),
    ("Journeys (optional): the activities that fulfil an offering's order type.",),
    (
        "Activities: one row per numbered activity; track is MAIN or a side track; systems and "
        "component_ids separated by ;.",
    ),
    ("FlowRules: kind decision, loop or parallel, from one activity number to another.",),
    ("ActivityIntegrations: how one activity hands over to another.",),
    ("Row 1 of each sheet holds the headers; keep them as they are.",),
)


class CatalogueFileAdapter:
    def read(self, file_format: CatalogueFileFormat, content: bytes) -> CatalogueContent:
        if file_format is CatalogueFileFormat.XLSX:
            return _read_workbook(content)
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise InvalidKnowledgeError("The catalogue file must be UTF-8 text.") from exc
        try:
            raw = (
                json.loads(text)
                if file_format is CatalogueFileFormat.JSON
                else yaml.safe_load(text)
            )
        except (json.JSONDecodeError, yaml.YAMLError) as exc:
            raise InvalidKnowledgeError(
                f"The {file_format.value.upper()} file could not be parsed: {_position(exc)}."
            ) from exc
        return content_from_mapping(raw)

    def write(self, file_format: CatalogueFileFormat, release: ArchitectureKnowledge) -> bytes:
        if file_format is CatalogueFileFormat.XLSX:
            return _workbook(release)
        mapping = release_to_mapping(release)
        if file_format is CatalogueFileFormat.JSON:
            return json.dumps(mapping, ensure_ascii=False, indent=2).encode()
        return yaml.safe_dump(mapping, allow_unicode=True, sort_keys=False).encode()

    def template(self) -> bytes:
        return _workbook(None)


def _position(exc: json.JSONDecodeError | yaml.YAMLError) -> str:
    if isinstance(exc, json.JSONDecodeError):
        return f"line {exc.lineno}, column {exc.colno}"
    mark = getattr(exc, "problem_mark", None)
    return f"line {mark.line + 1}" if mark is not None else "invalid syntax"


def _text(value: object, where: str, field: str) -> str:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise InvalidKnowledgeError(f"{where}: {field} is required.")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if not isinstance(value, str | int | float):
        raise InvalidKnowledgeError(f"{where}: {field} must be text.")
    return str(value).strip()


def _optional_text(value: object, where: str, field: str) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    return _text(value, where, field)


def _kind(value: object, where: str) -> RelationshipKind:
    """An optional relationship kind; "Calls API" and "calls_api" both read as calls_api."""
    text = _optional_text(value, where, "kind")
    if text is None:
        return RelationshipKind.UNSPECIFIED
    try:
        return RelationshipKind(text.strip().casefold().replace(" ", "_"))
    except ValueError as exc:
        raise InvalidKnowledgeError(f"{where}: kind must be one of {_KINDS}.") from exc


def _text_list(value: object, where: str, field: str) -> tuple[str, ...]:
    if value is None:
        return ()
    if isinstance(value, list):
        return tuple(_text(item, where, field) for item in value)
    raise InvalidKnowledgeError(f"{where}: {field} must be a list.")


def _entries(raw: dict[str, Any], key: str) -> list[dict[str, Any]]:
    value = raw.get(key, [])
    if not isinstance(value, list) or not all(isinstance(item, dict) for item in value):
        raise InvalidKnowledgeError(f"'{key}' must be a list of entries.")
    return value


def content_from_mapping(raw: object) -> CatalogueContent:
    """The YAML/JSON shape: ``systems`` with nested capabilities and components, plus
    ``dependencies``."""
    if not isinstance(raw, dict):
        raise InvalidKnowledgeError("The catalogue file must contain an object at the top level.")
    systems = []
    for number, item in enumerate(_entries(raw, "systems"), start=1):
        where = f"systems entry {number}"
        capabilities = []
        raw_capabilities = item.get("capabilities") or []
        if not isinstance(raw_capabilities, list):
            raise InvalidKnowledgeError(f"{where}: capabilities must be a list.")
        for position, capability in enumerate(raw_capabilities, start=1):
            place = f"{where}, capability {position}"
            if not isinstance(capability, dict):
                raise InvalidKnowledgeError(f"{place} must be an object.")
            try:
                capabilities.append(
                    KnowledgeCapability(
                        _text(capability.get("id"), place, "id"),
                        _text(capability.get("name"), place, "name"),
                        _text_list(capability.get("triggers"), place, "triggers"),
                        _optional_text(capability.get("domain"), place, "domain"),
                        _optional_text(capability.get("component"), place, "component"),
                    )
                )
            except InvalidKnowledgeError as exc:
                raise _located(place, exc) from exc
        components = []
        try:
            raw_components = _entries(item, "components")
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
        for position, component in enumerate(raw_components, start=1):
            place = f"{where}, component {position}"
            try:
                components.append(
                    SystemComponent(
                        id=_text(component.get("id"), place, "id"),
                        name=_text(component.get("name"), place, "name"),
                        name_ar=_optional_text(component.get("name_ar"), place, "name_ar"),
                        description=_optional_text(
                            component.get("description"), place, "description"
                        ),
                        aliases=_text_list(component.get("aliases"), place, "aliases"),
                        technology=_optional_text(component.get("technology"), place, "technology"),
                    )
                )
            except InvalidKnowledgeError as exc:
                raise _located(place, exc) from exc
        try:
            systems.append(
                SystemDefinition(
                    id=_text(item.get("id"), where, "id"),
                    name=_text(item.get("name"), where, "name"),
                    aliases=_text_list(item.get("aliases"), where, "aliases"),
                    capabilities=tuple(capabilities),
                    constraints=_text_list(item.get("constraints"), where, "constraints"),
                    name_ar=_optional_text(item.get("name_ar"), where, "name_ar"),
                    components=tuple(components),
                    description=_optional_text(item.get("description"), where, "description"),
                    landscape_domain_id=_optional_text(
                        item.get("landscape_domain"), where, "landscape_domain"
                    ),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    relationships = []
    for number, item in enumerate(_entries(raw, "dependencies"), start=1):
        where = f"dependencies entry {number}"
        try:
            relationships.append(
                SystemRelationship(
                    _text(item.get("source_system_id"), where, "source_system_id"),
                    _text(item.get("target_system_id"), where, "target_system_id"),
                    _text(item.get("description"), where, "description"),
                    _kind(item.get("kind"), where),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    return _content(
        systems,
        relationships,
        _mapped_domains(raw, "capability_domains", CapabilityDomain),
        _mapped_domains(raw, "landscape_domains", LandscapeDomain),
        _offerings(_entries(raw, "products")),
        _journeys(_entries(raw, "journeys")),
    )


class _LocatedError(InvalidKnowledgeError):
    """An error that already says which row or entry it is about."""


def _part[T](place: str, build: Callable[[], T]) -> T:
    try:
        return build()
    except _LocatedError:
        raise
    except InvalidKnowledgeError as exc:
        raise _LocatedError(str(_located(place, exc))) from exc


def _activity(step: dict[str, Any], place: str) -> Activity:
    return Activity(
        number=_text(step.get("number"), place, "number"),
        name=_text(step.get("name"), place, "name"),
        phase=_optional_text(step.get("phase"), place, "phase"),
        track=_optional_text(step.get("track"), place, "track"),
        performing_system_id=_optional_text(step.get("system"), place, "system"),
        supporting_system_ids=_text_list(step.get("supporting"), place, "supporting"),
        system_function=_optional_text(step.get("system_function"), place, "system_function"),
        mode=_optional_text(step.get("mode"), place, "mode"),
        customer_visible=_flag(step.get("customer_visible"), place, "customer_visible", None),
        description=_optional_text(step.get("description"), place, "description"),
        component_ids=_text_list(step.get("components"), place, "components"),
        input=_optional_text(step.get("input"), place, "input"),
        output=_optional_text(step.get("output"), place, "output"),
        etom=_optional_text(step.get("etom"), place, "etom"),
        confidence=_trust(step.get("confidence"), place),
        source=_optional_text(step.get("source"), place, "source"),
    )


def _rule(rule: dict[str, Any], place: str) -> FlowRule:
    return FlowRule(
        kind=flow_rule_kind(_text(rule.get("kind"), place, "kind")),
        from_activity=_text(rule.get("from"), place, "from"),
        to_activity=_text(rule.get("to"), place, "to"),
        condition=_optional_text(rule.get("condition"), place, "condition"),
        branch=_optional_text(rule.get("branch"), place, "branch"),
        parallel_group=_optional_text(rule.get("parallel_group"), place, "parallel_group"),
        rejoin_at=_optional_text(rule.get("rejoin_at"), place, "rejoin_at"),
        confidence=_trust(rule.get("confidence"), place),
        source=_optional_text(rule.get("source"), place, "source"),
    )


def _link(link: dict[str, Any], place: str) -> ActivityIntegration:
    return ActivityIntegration(
        from_activity=_text(link.get("from"), place, "from"),
        to_activity=_text(link.get("to"), place, "to"),
        interaction=_optional_text(link.get("interaction"), place, "interaction"),
        interface=_optional_text(link.get("interface"), place, "interface"),
        payload=_optional_text(link.get("payload"), place, "payload"),
        timing=_optional_text(link.get("timing"), place, "timing"),
        correlation_key=_optional_text(link.get("correlation_key"), place, "correlation_key"),
        confidence=_trust(link.get("confidence"), place),
        source=_optional_text(link.get("source"), place, "source"),
    )


# Journeys (ADR-0096), read the same way as offerings; each activity, rule and
# integration says which row it is when it is refused.
def _journeys(entries: list[dict[str, Any]]) -> list[Journey]:
    journeys = []
    for number, item in enumerate(entries, start=1):
        where = _where(item, f"journeys entry {number}")
        try:
            activities = [
                _part(place, partial(_activity, step, place))
                for position, step in enumerate(_sub_entries(item, "activities", where), 1)
                for place in [_where(step, f"{where}, activity {position}")]
            ]
            rules = [
                _part(place, partial(_rule, rule, place))
                for position, rule in enumerate(_sub_entries(item, "flow_rules", where), 1)
                for place in [_where(rule, f"{where}, flow rule {position}")]
            ]
            links = [
                _part(place, partial(_link, link, place))
                for position, link in enumerate(_sub_entries(item, "integrations", where), 1)
                for place in [_where(link, f"{where}, integration {position}")]
            ]
            journeys.append(
                Journey(
                    id=_text(item.get("id"), where, "id"),
                    name=_text(item.get("name"), where, "name"),
                    product_id=_optional_text(item.get("product"), where, "product"),
                    order_type_code=_optional_text(item.get("order_type"), where, "order_type"),
                    description=_optional_text(item.get("description"), where, "description"),
                    activities=tuple(activities),
                    flow_rules=tuple(rules),
                    integrations=tuple(links),
                    confidence=_trust(item.get("confidence"), where),
                    source=_optional_text(item.get("source"), where, "source"),
                )
            )
        except _LocatedError:
            raise
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    return journeys


def _journey_mapping(journey: Journey) -> dict[str, Any]:
    return {
        "id": journey.id,
        "name": journey.name,
        **_present(
            product=journey.product_id,
            order_type=journey.order_type_code,
            description=journey.description,
        ),
        **_sourced(journey),
        **_present(
            activities=[
                {
                    "number": step.number,
                    "name": step.name,
                    **_present(
                        phase=step.phase,
                        track=step.track,
                        system=step.performing_system_id,
                        supporting=list(step.supporting_system_ids),
                        system_function=step.system_function,
                        mode=step.mode,
                        customer_visible=step.customer_visible,
                        description=step.description,
                        components=list(step.component_ids),
                        input=step.input,
                        output=step.output,
                        etom=step.etom,
                    ),
                    **_sourced(step),
                }
                for step in journey.activities
            ],
            flow_rules=[
                {
                    "kind": rule.kind.value,
                    "from": rule.from_activity,
                    "to": rule.to_activity,
                    **_present(
                        condition=rule.condition,
                        branch=rule.branch,
                        parallel_group=rule.parallel_group,
                        rejoin_at=rule.rejoin_at,
                    ),
                    **_sourced(rule),
                }
                for rule in journey.flow_rules
            ],
            integrations=[
                {
                    "from": link.from_activity,
                    "to": link.to_activity,
                    **_present(
                        interaction=link.interaction,
                        interface=link.interface,
                        payload=link.payload,
                        timing=link.timing,
                        correlation_key=link.correlation_key,
                    ),
                    **_sourced(link),
                }
                for link in journey.integrations
            ],
        ),
    }


# Product offerings (ADR-0095). Workbook rows are gathered into the YAML/JSON shape and
# read by the same code, each entry keeping where it came from for its errors.
_YES = frozenset({"yes", "y", "true", "1"})
_NO = frozenset({"no", "n", "false", "0"})
_VALUE_KINDS = frozenset({"value", "values", "customer value"})
_AUDIENCE_KINDS = frozenset({"audience", "audiences", "who it is for"})


def _flag(value: object, where: str, field: str, default: bool | None) -> bool | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().casefold()
    if text in _YES:
        return True
    if text in _NO:
        return False
    raise InvalidKnowledgeError(f"{where}: {field} must be yes or no.")


def _trust(value: object, where: str) -> SourceConfidence | None:
    """How sure the source is: confirmed, inferred or gap, in any case."""
    text = _optional_text(value, where, "confidence")
    if text is None:
        return None
    try:
        return SourceConfidence(text.casefold())
    except ValueError as exc:
        raise InvalidKnowledgeError(
            f"{where}: confidence must be confirmed, inferred or gap."
        ) from exc


def _where(item: dict[str, Any], fallback: str) -> str:
    return str(item.get("_where") or fallback)


def _sub_entries(item: dict[str, Any], key: str, where: str) -> list[dict[str, Any]]:
    try:
        return _entries(item, key)
    except InvalidKnowledgeError as exc:
        raise _located(where, exc) from exc


def _point(item: dict[str, Any], where: str) -> OfferingPoint:
    return OfferingPoint(
        _text(item.get("name"), where, "name"),
        _optional_text(item.get("description"), where, "description"),
        _trust(item.get("confidence"), where),
        _optional_text(item.get("source"), where, "source"),
    )


def _offerings(entries: list[dict[str, Any]]) -> list[ProductOffering]:
    offerings = []
    for number, item in enumerate(entries, start=1):
        where = _where(item, f"products entry {number}")
        try:
            order_types = [
                OrderType(
                    _text(order.get("code"), place, "code"),
                    _text(order.get("name"), place, "name"),
                    bool(_flag(order.get("enabled"), place, "enabled", True)),
                    _optional_text(order.get("description"), place, "description"),
                    _trust(order.get("confidence"), place),
                    _optional_text(order.get("source"), place, "source"),
                )
                for position, order in enumerate(_sub_entries(item, "order_types", where), 1)
                for place in [_where(order, f"{where}, order type {position}")]
            ]
            components = []
            for position, part in enumerate(_sub_entries(item, "components", where), 1):
                place = _where(part, f"{where}, component {position}")
                responsibilities = [
                    ComponentResponsibility(
                        _text(duty.get("system"), spot, "system"),
                        _text(duty.get("role"), spot, "role"),
                        _text(duty.get("description"), spot, "description"),
                        _text_list(duty.get("order_types"), spot, "order_types"),
                        _trust(duty.get("confidence"), spot),
                        _optional_text(duty.get("source"), spot, "source"),
                    )
                    for index, duty in enumerate(_sub_entries(part, "responsibilities", place), 1)
                    for spot in [_where(duty, f"{place}, responsibility {index}")]
                ]
                components.append(
                    OfferingComponent(
                        id=_text(part.get("id"), place, "id"),
                        name=_text(part.get("name"), place, "name"),
                        code=_optional_text(part.get("code"), place, "code"),
                        kind=_optional_text(part.get("type"), place, "type"),
                        mandatory=_flag(part.get("mandatory"), place, "mandatory", None),
                        customer_visible=_flag(
                            part.get("customer_visible"), place, "customer_visible", None
                        ),
                        description=_optional_text(part.get("description"), place, "description"),
                        commercial_spec=_optional_text(
                            part.get("commercial_spec"), place, "commercial_spec"
                        ),
                        technical_spec=_optional_text(
                            part.get("technical_spec"), place, "technical_spec"
                        ),
                        technical_details=_optional_text(
                            part.get("technical_details"), place, "technical_details"
                        ),
                        responsibilities=tuple(responsibilities),
                        confidence=_trust(part.get("confidence"), place),
                        source=_optional_text(part.get("source"), place, "source"),
                    )
                )
            offerings.append(
                ProductOffering(
                    id=_text(item.get("id"), where, "id"),
                    name=_text(item.get("name"), where, "name"),
                    code=_optional_text(item.get("code"), where, "code"),
                    family=_optional_text(item.get("family"), where, "family"),
                    version=_optional_text(item.get("version"), where, "version"),
                    lifecycle=_optional_text(item.get("lifecycle"), where, "lifecycle"),
                    proposition=_optional_text(item.get("proposition"), where, "proposition"),
                    rules=_text_list(item.get("rules"), where, "rules"),
                    order_types=tuple(order_types),
                    components=tuple(components),
                    values=tuple(
                        _point(point, _where(point, f"{where}, value {index}"))
                        for index, point in enumerate(_sub_entries(item, "values", where), 1)
                    ),
                    audiences=tuple(
                        _point(point, _where(point, f"{where}, audience {index}"))
                        for index, point in enumerate(_sub_entries(item, "audiences", where), 1)
                    ),
                    confidence=_trust(item.get("confidence"), where),
                    source=_optional_text(item.get("source"), where, "source"),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    return offerings


def _present(**values: object) -> dict[str, Any]:
    """The keys that say something; files leave out what is unknown."""
    return {key: value for key, value in values.items() if value not in (None, "", [], ())}


def _sourced(item: Any) -> dict[str, Any]:
    return _present(
        confidence=item.confidence.value if item.confidence else None, source=item.source
    )


def _offering_mapping(offering: ProductOffering) -> dict[str, Any]:
    return {
        "id": offering.id,
        "name": offering.name,
        **_present(
            code=offering.code,
            family=offering.family,
            version=offering.version,
            lifecycle=offering.lifecycle,
            proposition=offering.proposition,
            rules=list(offering.rules),
        ),
        **_sourced(offering),
        **_present(
            order_types=[
                {
                    "code": order.code,
                    "name": order.name,
                    "enabled": order.enabled,
                    **_present(description=order.description),
                    **_sourced(order),
                }
                for order in offering.order_types
            ],
            components=[
                {
                    "id": part.id,
                    "name": part.name,
                    **_present(
                        code=part.code,
                        type=part.kind,
                        mandatory=part.mandatory,
                        customer_visible=part.customer_visible,
                        description=part.description,
                        commercial_spec=part.commercial_spec,
                        technical_spec=part.technical_spec,
                        technical_details=part.technical_details,
                    ),
                    **_sourced(part),
                    **_present(
                        responsibilities=[
                            {
                                "system": duty.system_id,
                                "role": duty.role,
                                "description": duty.description,
                                **_present(order_types=list(duty.order_types)),
                                **_sourced(duty),
                            }
                            for duty in part.responsibilities
                        ]
                    ),
                }
                for part in offering.components
            ],
            values=[
                {"name": point.name, **_present(description=point.description), **_sourced(point)}
                for point in offering.values
            ],
            audiences=[
                {"name": point.name, **_present(description=point.description), **_sourced(point)}
                for point in offering.audiences
            ],
        ),
    }


def _mapped_domains[DomainT: (CapabilityDomain, LandscapeDomain)](
    raw: dict[str, Any], key: str, kind: type[DomainT]
) -> list[DomainT]:
    domains = []
    for number, item in enumerate(_entries(raw, key), start=1):
        where = f"{key} entry {number}"
        try:
            domains.append(
                kind(
                    _text(item.get("id"), where, "id"),
                    _text(item.get("name"), where, "name"),
                    _optional_text(item.get("name_ar"), where, "name_ar"),
                    _optional_text(item.get("parent_id"), where, "parent_id"),
                    _optional_text(item.get("description"), where, "description"),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    return domains


def _content(
    systems: list[SystemDefinition],
    relationships: list[SystemRelationship],
    domains: list[CapabilityDomain],
    landscape: list[LandscapeDomain],
    offerings: list[ProductOffering],
    journeys: list[Journey],
) -> CatalogueContent:
    """The file's content, refusing a placement in a domain the file does not list."""
    known = {item.id for item in domains}
    places = {item.id for item in landscape}
    for system in systems:
        for capability in system.capabilities:
            if capability.domain_id is not None and capability.domain_id not in known:
                raise InvalidKnowledgeError(
                    f"{system.id}, capability {capability.id}: domain {capability.domain_id!r} "
                    "is not listed among the capability domains."
                )
        if system.landscape_domain_id is not None and system.landscape_domain_id not in places:
            raise InvalidKnowledgeError(
                f"{system.id}: landscape domain {system.landscape_domain_id!r} is not listed "
                "among the landscape domains."
            )
    return CatalogueContent(
        tuple(systems),
        tuple(relationships),
        tuple(domains),
        tuple(landscape),
        tuple(offerings),
        tuple(journeys),
    )


def _domain_mapping(item: CapabilityDomain | LandscapeDomain) -> dict[str, Any]:
    return {
        "id": item.id,
        "name": item.name,
        **({"name_ar": item.name_ar} if item.name_ar else {}),
        **({"parent_id": item.parent_id} if item.parent_id else {}),
        **({"description": item.description} if item.description else {}),
    }


def _component_mapping(component: SystemComponent) -> dict[str, Any]:
    return {
        "id": component.id,
        "name": component.name,
        **({"name_ar": component.name_ar} if component.name_ar else {}),
        **({"aliases": list(component.aliases)} if component.aliases else {}),
        **({"technology": component.technology} if component.technology else {}),
        **({"description": component.description} if component.description else {}),
    }


def release_to_mapping(release: ArchitectureKnowledge) -> dict[str, Any]:
    return {
        "version": release.id,
        "systems": [
            {
                "id": item.id,
                "name": item.name,
                **({"name_ar": item.name_ar} if item.name_ar else {}),
                "aliases": list(item.aliases),
                **({"description": item.description} if item.description else {}),
                **(
                    {"landscape_domain": item.landscape_domain_id}
                    if item.landscape_domain_id
                    else {}
                ),
                "capabilities": [
                    {
                        "id": cap.id,
                        "name": cap.name,
                        "triggers": list(cap.triggers),
                        **({"domain": cap.domain_id} if cap.domain_id else {}),
                        **({"component": cap.component_id} if cap.component_id else {}),
                    }
                    for cap in item.capabilities
                ],
                **(
                    {"components": [_component_mapping(part) for part in item.components]}
                    if item.components
                    else {}
                ),
                **({"constraints": list(item.constraints)} if item.constraints else {}),
            }
            for item in release.systems
        ],
        "dependencies": [
            {
                "source_system_id": item.source_system_id,
                "target_system_id": item.target_system_id,
                "description": item.description,
                **(
                    {"kind": item.kind.value}
                    if item.kind is not RelationshipKind.UNSPECIFIED
                    else {}
                ),
            }
            for item in release.relationships
        ],
        **(
            {"capability_domains": [_domain_mapping(item) for item in release.capability_domains]}
            if release.capability_domains
            else {}
        ),
        **(
            {"landscape_domains": [_domain_mapping(item) for item in release.landscape_domains]}
            if release.landscape_domains
            else {}
        ),
        **(
            {"products": [_offering_mapping(item) for item in release.products]}
            if release.products
            else {}
        ),
        **(
            {"journeys": [_journey_mapping(item) for item in release.journeys]}
            if release.journeys
            else {}
        ),
    }


def _check_archive(content: bytes) -> None:
    """Refuse archive bombs before openpyxl inflates anything."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            parts = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise InvalidKnowledgeError("The Excel file is damaged or not an .xlsx workbook.") from exc
    if len(parts) > MAX_OFFICE_PARTS or (
        sum(item.file_size for item in parts) > MAX_OFFICE_UNCOMPRESSED_BYTES
    ):
        raise InvalidKnowledgeError("The Excel file is too large to import.")


def _rows(workbook: Any, sheet: str) -> Iterator[tuple[int, dict[str, object]]]:
    if sheet not in workbook.sheetnames:
        if sheet in _REQUIRED_SHEETS:
            raise InvalidKnowledgeError(f"The workbook needs a '{sheet}' sheet.")
        return
    rows = workbook[sheet].iter_rows(values_only=True)
    header_row = next(rows, None) or ()
    headers = [str(value).strip().casefold() if value is not None else "" for value in header_row]
    if any(name not in headers for name in _REQUIRED_HEADERS[sheet]):
        raise InvalidKnowledgeError(
            f"{sheet} row 1 must have the headers {', '.join(_HEADERS[sheet])}."
        )
    for number, values in enumerate(rows, start=2):
        if number > MAX_WORKBOOK_ROWS:
            raise InvalidKnowledgeError(f"{sheet} has more rows than can be imported.")
        cells: dict[str, object] = {
            header: value for header, value in zip(headers, values, strict=False) if header
        }
        if all(value is None or str(value).strip() == "" for value in cells.values()):
            continue
        yield number, cells


def _split(value: object, where: str, field: str) -> tuple[str, ...]:
    text = _optional_text(value, where, field)
    if text is None:
        return ()
    return tuple(part.strip() for part in text.split(_LIST_SEPARATOR) if part.strip())


def _read_workbook(content: bytes) -> CatalogueContent:
    _check_archive(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise InvalidKnowledgeError("The Excel file could not be opened.") from exc
    try:
        systems: dict[str, dict[str, Any]] = {}
        for number, cells in _rows(workbook, SYSTEMS):
            where = f"{SYSTEMS} row {number}"
            system_id = _text(cells.get("system_id"), where, "system_id")
            if system_id in systems:
                raise InvalidKnowledgeError(f"{where}: system {system_id!r} is listed twice.")
            systems[system_id] = {
                "where": where,
                "name": _text(cells.get("name"), where, "name"),
                "name_ar": _optional_text(cells.get("name_ar"), where, "name_ar"),
                "aliases": _split(cells.get("aliases"), where, "aliases"),
                "description": _optional_text(cells.get("description"), where, "description"),
                "landscape_domain_id": _optional_text(
                    cells.get("landscape_domain_id"), where, "landscape_domain_id"
                ),
                "capabilities": [],
                "constraints": [],
                "components": [],
            }

        def owner(cells: dict[str, object], where: str) -> dict[str, Any]:
            system_id = _text(cells.get("system_id"), where, "system_id")
            if system_id not in systems:
                raise InvalidKnowledgeError(
                    f"{where}: system {system_id!r} is not listed on the {SYSTEMS} sheet."
                )
            return systems[system_id]

        for number, cells in _rows(workbook, CAPABILITIES):
            where = f"{CAPABILITIES} row {number}"
            try:
                owner(cells, where)["capabilities"].append(
                    KnowledgeCapability(
                        _text(cells.get("capability_id"), where, "capability_id"),
                        _text(cells.get("name"), where, "name"),
                        _split(cells.get("triggers"), where, "triggers"),
                        _optional_text(cells.get("domain_id"), where, "domain_id"),
                        _optional_text(cells.get("component_id"), where, "component_id"),
                    )
                )
            except InvalidKnowledgeError as exc:
                raise _located(where, exc) from exc
        for number, cells in _rows(workbook, COMPONENTS):
            where = f"{COMPONENTS} row {number}"
            try:
                owner(cells, where)["components"].append(
                    SystemComponent(
                        id=_text(cells.get("component_id"), where, "component_id"),
                        name=_text(cells.get("name"), where, "name"),
                        name_ar=_optional_text(cells.get("name_ar"), where, "name_ar"),
                        description=_optional_text(cells.get("description"), where, "description"),
                        aliases=_split(cells.get("aliases"), where, "aliases"),
                        technology=_optional_text(cells.get("technology"), where, "technology"),
                    )
                )
            except InvalidKnowledgeError as exc:
                raise _located(where, exc) from exc
        for number, cells in _rows(workbook, CONSTRAINTS):
            where = f"{CONSTRAINTS} row {number}"
            owner(cells, where)["constraints"].append(
                _text(cells.get("constraint"), where, "constraint")
            )
        relationships = []
        for number, cells in _rows(workbook, RELATIONSHIPS):
            where = f"{RELATIONSHIPS} row {number}"
            try:
                relationships.append(
                    SystemRelationship(
                        _text(cells.get("source_system_id"), where, "source_system_id"),
                        _text(cells.get("target_system_id"), where, "target_system_id"),
                        _text(cells.get("description"), where, "description"),
                        _kind(cells.get("kind"), where),
                    )
                )
            except InvalidKnowledgeError as exc:
                raise _located(where, exc) from exc
        domains = _sheet_domains(workbook, DOMAINS, CapabilityDomain)
        landscape = _sheet_domains(workbook, LANDSCAPE, LandscapeDomain)
        offerings = _offerings(_sheet_offerings(workbook))
        journeys = _journeys(_sheet_journeys(workbook))
    finally:
        workbook.close()
    definitions = []
    for system_id, data in systems.items():
        try:
            definitions.append(
                SystemDefinition(
                    id=system_id,
                    name=data["name"],
                    aliases=data["aliases"],
                    capabilities=tuple(data["capabilities"]),
                    constraints=tuple(data["constraints"]),
                    name_ar=data["name_ar"],
                    components=tuple(data["components"]),
                    description=data["description"],
                    landscape_domain_id=data["landscape_domain_id"],
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(data["where"], exc) from exc
    return _content(definitions, relationships, domains, landscape, offerings, journeys)


def _sheet_journeys(workbook: Any) -> list[dict[str, Any]]:
    """The journey sheets gathered into the YAML/JSON shape, each entry knowing its row."""
    journeys: dict[str, dict[str, Any]] = {}
    for number, cells in _rows(workbook, JOURNEYS):
        where = f"{JOURNEYS} row {number}"
        journey_id = _text(cells.get("journey_id"), where, "journey_id")
        if journey_id in journeys:
            raise InvalidKnowledgeError(f"{where}: journey {journey_id!r} is listed twice.")
        journeys[journey_id] = {
            "_where": where,
            "id": journey_id,
            "name": cells.get("name"),
            "product": cells.get("product_id"),
            "order_type": cells.get("order_type_code"),
            "description": cells.get("description"),
            "confidence": cells.get("confidence"),
            "source": cells.get("source"),
            "activities": [],
            "flow_rules": [],
            "integrations": [],
        }

    def journey(cells: dict[str, object], where: str) -> dict[str, Any]:
        journey_id = _text(cells.get("journey_id"), where, "journey_id")
        if journey_id not in journeys:
            raise InvalidKnowledgeError(
                f"{where}: journey {journey_id!r} is not listed on the {JOURNEYS} sheet."
            )
        return journeys[journey_id]

    for number, cells in _rows(workbook, ACTIVITIES):
        where = f"{ACTIVITIES} row {number}"
        journey(cells, where)["activities"].append(
            {
                **cells,
                "_where": where,
                "system": cells.get("performing_system_id"),
                "supporting": list(
                    _split(cells.get("supporting_system_ids"), where, "supporting_system_ids")
                ),
                "components": list(_split(cells.get("component_ids"), where, "component_ids")),
            }
        )
    for number, cells in _rows(workbook, FLOW_RULES):
        where = f"{FLOW_RULES} row {number}"
        journey(cells, where)["flow_rules"].append(
            {
                **cells,
                "_where": where,
                "from": cells.get("from_activity"),
                "to": cells.get("to_activity"),
            }
        )
    for number, cells in _rows(workbook, ACTIVITY_INTEGRATIONS):
        where = f"{ACTIVITY_INTEGRATIONS} row {number}"
        journey(cells, where)["integrations"].append(
            {
                **cells,
                "_where": where,
                "from": cells.get("from_activity"),
                "to": cells.get("to_activity"),
            }
        )
    return list(journeys.values())


def _sheet_offerings(workbook: Any) -> list[dict[str, Any]]:
    """The offering sheets gathered into the YAML/JSON shape, each entry knowing its row."""
    offerings: dict[str, dict[str, Any]] = {}
    for number, cells in _rows(workbook, PRODUCTS):
        where = f"{PRODUCTS} row {number}"
        product_id = _text(cells.get("product_id"), where, "product_id")
        if product_id in offerings:
            raise InvalidKnowledgeError(f"{where}: product {product_id!r} is listed twice.")
        offerings[product_id] = {
            **{key: cells.get(key) for key in _HEADERS[PRODUCTS][1:] if key != "rules"},
            "_where": where,
            "id": product_id,
            "rules": list(_split(cells.get("rules"), where, "rules")),
            "order_types": [],
            "components": [],
            "values": [],
            "audiences": [],
        }

    def offering(cells: dict[str, object], where: str) -> dict[str, Any]:
        product_id = _text(cells.get("product_id"), where, "product_id")
        if product_id not in offerings:
            raise InvalidKnowledgeError(
                f"{where}: product {product_id!r} is not listed on the {PRODUCTS} sheet."
            )
        return offerings[product_id]

    for number, cells in _rows(workbook, ORDER_TYPES):
        where = f"{ORDER_TYPES} row {number}"
        offering(cells, where)["order_types"].append({**cells, "_where": where})
    parts: dict[tuple[str, str], dict[str, Any]] = {}
    for number, cells in _rows(workbook, OFFERING_COMPONENTS):
        where = f"{OFFERING_COMPONENTS} row {number}"
        owner = offering(cells, where)
        part = {
            **cells,
            "_where": where,
            "id": cells.get("component_id"),
            "responsibilities": [],
        }
        parts[(owner["id"], _text(cells.get("component_id"), where, "component_id"))] = part
        owner["components"].append(part)
    for number, cells in _rows(workbook, RESPONSIBILITIES):
        where = f"{RESPONSIBILITIES} row {number}"
        key = (
            offering(cells, where)["id"],
            _text(cells.get("component_id"), where, "component_id"),
        )
        if key not in parts:
            raise InvalidKnowledgeError(
                f"{where}: component {key[1]!r} is not listed on the {OFFERING_COMPONENTS} sheet."
            )
        parts[key]["responsibilities"].append(
            {
                **cells,
                "_where": where,
                "system": cells.get("system_id"),
                "order_types": list(_split(cells.get("order_types"), where, "order_types")),
            }
        )
    for number, cells in _rows(workbook, PRODUCT_POINTS):
        where = f"{PRODUCT_POINTS} row {number}"
        kind = (_text(cells.get("kind"), where, "kind")).casefold()
        if kind not in _VALUE_KINDS | _AUDIENCE_KINDS:
            raise InvalidKnowledgeError(f"{where}: kind must be value or audience.")
        bucket = "values" if kind in _VALUE_KINDS else "audiences"
        offering(cells, where)[bucket].append({**cells, "_where": where})
    return list(offerings.values())


def _sheet_domains[DomainT: (CapabilityDomain, LandscapeDomain)](
    workbook: Any, sheet: str, kind: type[DomainT]
) -> list[DomainT]:
    domains = []
    for number, cells in _rows(workbook, sheet):
        where = f"{sheet} row {number}"
        try:
            domains.append(
                kind(
                    _text(cells.get("domain_id"), where, "domain_id"),
                    _text(cells.get("name"), where, "name"),
                    _optional_text(cells.get("name_ar"), where, "name_ar"),
                    _optional_text(cells.get("parent_id"), where, "parent_id"),
                    _optional_text(cells.get("description"), where, "description"),
                )
            )
        except InvalidKnowledgeError as exc:
            raise _located(where, exc) from exc
    return domains


def _located(where: str, exc: InvalidKnowledgeError) -> InvalidKnowledgeError:
    message = str(exc)
    return exc if message.startswith(where) else InvalidKnowledgeError(f"{where}: {message}")


def _append(sheet: Any, values: Sequence[object], *, header: bool = False) -> None:
    sheet.append(list(values))
    if header:
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)


def _yes(value: bool | None) -> str | None:
    return None if value is None else "yes" if value else "no"


def _confidence(item: Any) -> str | None:
    return item.confidence.value if item.confidence else None


def _write_offering(sheets: dict[str, Any], product: ProductOffering) -> None:
    _append(
        sheets[PRODUCTS],
        (
            product.id,
            product.name,
            product.code,
            product.family,
            product.version,
            product.lifecycle,
            product.proposition,
            f"{_LIST_SEPARATOR} ".join(product.rules),
            _confidence(product),
            product.source,
        ),
    )
    for order in product.order_types:
        _append(
            sheets[ORDER_TYPES],
            (
                product.id,
                order.code,
                order.name,
                _yes(order.enabled),
                order.description,
                _confidence(order),
                order.source,
            ),
        )
    for part in product.components:
        _append(
            sheets[OFFERING_COMPONENTS],
            (
                product.id,
                part.id,
                part.name,
                part.code,
                part.kind,
                _yes(part.mandatory),
                _yes(part.customer_visible),
                part.description,
                part.commercial_spec,
                part.technical_spec,
                part.technical_details,
                _confidence(part),
                part.source,
            ),
        )
        for duty in part.responsibilities:
            _append(
                sheets[RESPONSIBILITIES],
                (
                    product.id,
                    part.id,
                    duty.system_id,
                    duty.role,
                    duty.description,
                    f"{_LIST_SEPARATOR} ".join(duty.order_types),
                    _confidence(duty),
                    duty.source,
                ),
            )
    for kind, points in (("value", product.values), ("audience", product.audiences)):
        for point in points:
            _append(
                sheets[PRODUCT_POINTS],
                (product.id, kind, point.name, point.description, _confidence(point), point.source),
            )


def _write_journey(sheets: dict[str, Any], journey: Journey) -> None:
    joined = f"{_LIST_SEPARATOR} ".join
    _append(
        sheets[JOURNEYS],
        (
            journey.id,
            journey.name,
            journey.product_id,
            journey.order_type_code,
            journey.description,
            _confidence(journey),
            journey.source,
        ),
    )
    for step in journey.activities:
        _append(
            sheets[ACTIVITIES],
            (
                journey.id,
                step.number,
                step.name,
                step.phase,
                step.track,
                step.performing_system_id,
                joined(step.supporting_system_ids),
                step.system_function,
                step.mode,
                _yes(step.customer_visible),
                step.description,
                joined(step.component_ids),
                step.input,
                step.output,
                step.etom,
                _confidence(step),
                step.source,
            ),
        )
    for rule in journey.flow_rules:
        _append(
            sheets[FLOW_RULES],
            (
                journey.id,
                rule.kind.value,
                rule.from_activity,
                rule.to_activity,
                rule.condition,
                rule.branch,
                rule.parallel_group,
                rule.rejoin_at,
                _confidence(rule),
                rule.source,
            ),
        )
    for link in journey.integrations:
        _append(
            sheets[ACTIVITY_INTEGRATIONS],
            (
                journey.id,
                link.from_activity,
                link.to_activity,
                link.interaction,
                link.interface,
                link.payload,
                link.timing,
                link.correlation_key,
                _confidence(link),
                link.source,
            ),
        )


def _workbook(release: ArchitectureKnowledge | None) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    instructions = workbook.create_sheet(INSTRUCTIONS)
    for line in _INSTRUCTIONS:
        instructions.append(list(line))
    instructions["A1"].font = Font(bold=True)
    instructions.column_dimensions["A"].width = 90
    sheets = {name: workbook.create_sheet(name) for name in _HEADERS}
    for name, sheet in sheets.items():
        _append(sheet, _HEADERS[name], header=True)
        sheet.freeze_panes = "A2"
        for column in "ABCDEFGHIJKLMNOPQ":
            sheet.column_dimensions[column].width = 32
    systems = release.systems if release is not None else ()
    for system in systems:
        _append(
            sheets[SYSTEMS],
            (
                system.id,
                system.name,
                system.name_ar,
                f"{_LIST_SEPARATOR} ".join(system.aliases),
                system.description,
                system.landscape_domain_id,
            ),
        )
        for capability in system.capabilities:
            _append(
                sheets[CAPABILITIES],
                (
                    system.id,
                    capability.id,
                    capability.name,
                    f"{_LIST_SEPARATOR} ".join(capability.triggers),
                    capability.domain_id,
                    capability.component_id,
                ),
            )
        for component in system.components:
            _append(
                sheets[COMPONENTS],
                (
                    system.id,
                    component.id,
                    component.name,
                    component.name_ar,
                    f"{_LIST_SEPARATOR} ".join(component.aliases),
                    component.technology,
                    component.description,
                ),
            )
        for constraint in system.constraints:
            _append(sheets[CONSTRAINTS], (system.id, constraint))
    for relationship in release.relationships if release is not None else ():
        _append(
            sheets[RELATIONSHIPS],
            (
                relationship.source_system_id,
                relationship.target_system_id,
                relationship.description,
                relationship.kind.value,
            ),
        )
    for sheet, tree in (
        (DOMAINS, release.capability_domains if release is not None else ()),
        (LANDSCAPE, release.landscape_domains if release is not None else ()),
    ):
        for domain in tree:
            _append(
                sheets[sheet],
                (domain.id, domain.name, domain.name_ar, domain.parent_id, domain.description),
            )
    for product in release.products if release is not None else ():
        _write_offering(sheets, product)
    for journey in release.journeys if release is not None else ():
        _write_journey(sheets, journey)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
