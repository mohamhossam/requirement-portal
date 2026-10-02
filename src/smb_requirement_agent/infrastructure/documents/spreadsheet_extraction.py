"""Excel (XLSX) worksheets, merged cells and embedded images."""

from __future__ import annotations

import io
import posixpath
import zipfile
from dataclasses import dataclass
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from smb_requirement_agent.application.errors import (
    DocumentExtractionError,
    UnsupportedDocumentError,
)
from smb_requirement_agent.application.ports.document_extractor import (
    ExtractedDocument,
)
from smb_requirement_agent.domain.document.entities import (
    DocumentExtractionWarning,
)
from smb_requirement_agent.domain.document.value_objects import (
    EvidenceBlockKind,
    ExtractionWarningSeverity,
)
from smb_requirement_agent.infrastructure.documents.extraction_base import (
    MAX_WORKBOOK_COLUMNS,
    MAX_WORKBOOK_NONEMPTY_CELLS,
    MAX_WORKBOOK_ROWS,
    MAX_WORKBOOK_SHEETS,
    XLSX_EXTRACTION_VERSION,
    EvidenceBuilder,
    ExtractionBase,
    R,
    S,
)


@dataclass(frozen=True)
class _SpreadsheetMerge:
    first_column: int
    first_row: int
    last_column: int
    last_row: int

    @property
    def anchor(self) -> str:
        return f"{get_column_letter(self.first_column)}{self.first_row}"

    @property
    def reference(self) -> str:
        return f"{self.anchor}:{get_column_letter(self.last_column)}{self.last_row}"


@dataclass(frozen=True)
class _WorksheetMerges:
    rows: dict[int, tuple[_SpreadsheetMerge, ...]]
    max_row: int
    max_column: int
    cell_count: int

    def cells_for_row(self, row_number: int) -> dict[int, _SpreadsheetMerge]:
        return {
            column: merged
            for merged in self.rows.get(row_number, ())
            for column in range(merged.first_column, merged.last_column + 1)
        }


class SpreadsheetExtraction(ExtractionBase):
    def _xlsx(self, content: bytes) -> ExtractedDocument:
        builder = EvidenceBuilder([], [], [])
        try:
            with self._office_archive(content, "XLSX") as archive:
                names = set(archive.namelist())
                if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                    raise DocumentExtractionError("XLSX package is missing required content.")
                if any(name.startswith("xl/externalLinks/") for name in names):
                    builder.warnings.append(
                        DocumentExtractionWarning(
                            "external_workbook_links",
                            ExtractionWarningSeverity.WARNING,
                            "External workbook links were not followed.",
                        )
                    )
                media = tuple(
                    name
                    for name in sorted(names)
                    if name.startswith("xl/media/") and not name.endswith("/")
                )
                charts = tuple(
                    name
                    for name in sorted(names)
                    if name.startswith("xl/charts/chart") and name.endswith(".xml")
                )
                media_contexts, chart_contexts, sheet_merges, sheet_paths = self._xlsx_structure(
                    archive
                )
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            cached_workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True, keep_links=False
            )
            try:
                if len(workbook.worksheets) > MAX_WORKBOOK_SHEETS:
                    raise UnsupportedDocumentError(
                        f"XLSX contains more than {MAX_WORKBOOK_SHEETS} worksheets."
                    )
                nonempty_cells = 0
                visited_cells = 0
                unavailable_formula_cache = False
                for sheet in workbook.worksheets:
                    cached_sheet = cached_workbook[sheet.title]
                    merges = sheet_merges[sheet.title]
                    if sheet.max_row is None or sheet.max_column is None:
                        sheet.calculate_dimension(force=True)
                    max_row = max(sheet.max_row or 1, merges.max_row)
                    max_column = max(sheet.max_column or 1, merges.max_column)
                    if max_row > MAX_WORKBOOK_ROWS or max_column > MAX_WORKBOOK_COLUMNS:
                        raise UnsupportedDocumentError(
                            "XLSX worksheet dimensions exceed the safe extraction bounds."
                        )
                    hidden = sheet.sheet_state != "visible"
                    path = sheet_paths[sheet.title]
                    location = path[0].removeprefix("Hidden worksheet: ")
                    heading = builder.add(
                        EvidenceBlockKind.HEADING, path[0], path, text=sheet.title
                    )
                    if hidden:
                        builder.warnings.append(
                            DocumentExtractionWarning(
                                "hidden_worksheet",
                                ExtractionWarningSeverity.WARNING,
                                f"{location} is {sheet.sheet_state} "
                                "and is excluded from AI analysis unless explicitly selected.",
                                heading.id,
                            )
                        )
                    if visited_cells + max_row * max_column > self._max_visited_spreadsheet_cells:
                        raise UnsupportedDocumentError(
                            "XLSX contains too many visited cells for safe extraction."
                        )
                    if merges.rows:
                        builder.warnings.append(
                            DocumentExtractionWarning(
                                "worksheet_merged_cells",
                                ExtractionWarningSeverity.WARNING,
                                f"{location}: compare merged ranges "
                                "with the original. "
                                "Bracketed merge markers are structural annotations; "
                                "anchor wording is not copied into continuation cells.",
                                heading.id,
                            )
                        )
                    rows = iter(sheet.iter_rows(max_row=max_row, max_col=max_column))
                    cached_rows = iter(cached_sheet.iter_rows(max_row=max_row, max_col=max_column))
                    # Read-only parsers can omit empty trailing merge rows.
                    for row_number in range(1, max_row + 1):
                        row, cached_row = next(rows, ()), next(cached_rows, ())
                        values: list[str] = []
                        merged_cells = merges.cells_for_row(row_number)
                        for column_number in range(1, max_column + 1):
                            cell = row[column_number - 1] if column_number <= len(row) else None
                            cached_cell = (
                                cached_row[column_number - 1]
                                if column_number <= len(cached_row)
                                else None
                            )
                            visited_cells += 1
                            if visited_cells > self._max_visited_spreadsheet_cells:
                                raise UnsupportedDocumentError(
                                    "XLSX contains too many visited cells for safe extraction."
                                )
                            coordinate = f"{get_column_letter(column_number)}{row_number}"
                            merged = merged_cells.get(column_number)
                            value = cell.value if cell is not None else None
                            if merged is not None and coordinate != merged.anchor:
                                if value is not None:
                                    raise DocumentExtractionError(
                                        "XLSX merged continuation cell contains a value. "
                                        "Unmerge or correct the original workbook before uploading."
                                    )
                                values.append(
                                    f"{coordinate}=[merged {merged.reference} continuation; "
                                    f"anchor {merged.anchor}]"
                                )
                                continue
                            if value is None:
                                if merged is not None:
                                    values.append(
                                        f"{coordinate}=[empty cell] "
                                        f"[merged {merged.reference} anchor]"
                                    )
                                continue
                            nonempty_cells += 1
                            if nonempty_cells > MAX_WORKBOOK_NONEMPTY_CELLS:
                                raise UnsupportedDocumentError(
                                    "XLSX contains too many non-empty cells for safe extraction."
                                )
                            display = str(value)
                            if cell is not None and cell.data_type == "f":
                                cached_value = (
                                    cached_cell.value if cached_cell is not None else None
                                )
                                if cached_value is None:
                                    unavailable_formula_cache = True
                                    display = f"{display} (cached value unavailable)"
                                else:
                                    display = f"{display} (cached value: {cached_value})"
                            if merged is not None:
                                display += f" [merged {merged.reference} anchor]"
                            values.append(f"{coordinate}={display}")
                        if values:
                            builder.add(
                                EvidenceBlockKind.WORKSHEET_RANGE,
                                f"{location}!{row_number}:{row_number}",
                                path,
                                text=" | ".join(values),
                            )
            finally:
                workbook.close()
                cached_workbook.close()
            if unavailable_formula_cache:
                builder.warnings.append(
                    DocumentExtractionWarning(
                        "formula_cache_unavailable",
                        ExtractionWarningSeverity.WARNING,
                        "One or more formulas had no cached value; formulas were not executed.",
                    )
                )
            if charts:
                with self._office_archive(content, "XLSX") as archive:
                    chart_number = 0
                    for chart_path in charts:
                        chart = self._parse_xml(
                            archive.read(chart_path), f"Workbook chart {chart_path}"
                        )
                        title_parts = [
                            (item.text or "").strip()
                            for item in chart.iter()
                            if item.tag.endswith("}t") and (item.text or "").strip()
                        ]
                        ranges = [
                            (item.text or "").strip()
                            for item in chart.iter()
                            if item.tag.endswith("}f") and (item.text or "").strip()
                        ]
                        contexts = chart_contexts.get(chart_path) or {("Workbook charts",)}
                        for context in sorted(contexts):
                            chart_number += 1
                            builder.add(
                                EvidenceBlockKind.WORKSHEET_RANGE,
                                f"Chart {chart_number}",
                                context,
                                text=(
                                    f"Title: {' '.join(title_parts) or '(untitled)'}; "
                                    f"Referenced ranges: {', '.join(ranges) or '(none)'}"
                                ),
                            )
            if media:
                with self._office_archive(content, "XLSX") as archive:
                    image_number = 0
                    for package_path in media:
                        mime_type = self._image_mime(package_path)
                        image_bytes = archive.read(package_path)
                        if mime_type is None:
                            missing = builder.add(
                                EvidenceBlockKind.EXTERNAL_REFERENCE,
                                f"Workbook media {len(builder.blocks) + 1}",
                                ("Workbook media",),
                                text="Image content unavailable: unsupported workbook image.",
                            )
                            builder.warnings.append(
                                DocumentExtractionWarning(
                                    "unsupported_workbook_image",
                                    ExtractionWarningSeverity.BLOCKING,
                                    f"Workbook image {package_path} is not PNG or JPEG.",
                                    missing.id,
                                )
                            )
                            continue
                        try:
                            mime_type, image_bytes, decorative = self._sanitize_image(
                                image_bytes, mime_type
                            )
                        except DocumentExtractionError:
                            missing = builder.add(
                                EvidenceBlockKind.EXTERNAL_REFERENCE,
                                f"Workbook media {len(builder.blocks) + 1}",
                                ("Workbook media",),
                                text="Image content unavailable: unsafe or invalid workbook image.",
                            )
                            builder.warnings.append(
                                DocumentExtractionWarning(
                                    "invalid_workbook_image",
                                    ExtractionWarningSeverity.BLOCKING,
                                    f"Workbook image {package_path} could not be decoded safely.",
                                    missing.id,
                                )
                            )
                            continue
                        if decorative:
                            continue
                        contexts = media_contexts.get(package_path) or {("Workbook media",)}
                        for context in sorted(contexts):
                            image_number += 1
                            asset_id = self._asset_id(
                                f"{package_path}|{'/'.join(context)}", image_bytes
                            )
                            builder.add(
                                EvidenceBlockKind.IMAGE,
                                f"Workbook image {image_number}",
                                context,
                                asset=(asset_id, mime_type, package_path, image_bytes),
                            )
        except zipfile.BadZipFile as exc:
            raise DocumentExtractionError("XLSX package is corrupt.") from exc
        except (ValueError, KeyError, OSError) as exc:
            raise DocumentExtractionError("XLSX extraction failed.") from exc
        builder.warnings.append(
            DocumentExtractionWarning(
                "worksheet_name_structure",
                ExtractionWarningSeverity.WARNING,
                "Worksheet locations use workbook tab positions, including empty and hidden tabs. "
                "Names are independently reviewable heading text; excluding or correcting a name "
                "does not change row locations or hidden-sheet selection. Formula and chart text "
                "may explicitly reference names; review those passages separately.",
            )
        )
        text = self._non_blank(self._blocks_text(builder.blocks), "XLSX")
        return self._result(text, builder, extraction_version=XLSX_EXTRACTION_VERSION)

    def _xlsx_structure(
        self,
        archive: zipfile.ZipFile,
    ) -> tuple[
        dict[str, set[tuple[str, ...]]],
        dict[str, set[tuple[str, ...]]],
        dict[str, _WorksheetMerges],
        dict[str, tuple[str, ...]],
    ]:
        media: dict[str, set[tuple[str, ...]]] = {}
        charts: dict[str, set[tuple[str, ...]]] = {}
        merges: dict[str, _WorksheetMerges] = {}
        paths: dict[str, tuple[str, ...]] = {}
        merged_cells = 0
        workbook = self._parse_xml(archive.read("xl/workbook.xml"), "XLSX workbook")
        workbook_relationships = self._relationships(archive, "xl/_rels/workbook.xml.rels")
        names = set(archive.namelist())
        for number, sheet in enumerate(workbook.findall(f".//{{{S}}}sheet"), 1):
            relationship_id = sheet.get(f"{{{R}}}id", "")
            relationship = workbook_relationships.get(relationship_id)
            if relationship is None or relationship[1]:
                raise DocumentExtractionError("XLSX worksheet relationship is invalid.")
            sheet_path = posixpath.normpath(posixpath.join("xl", relationship[0])).lstrip("/")
            if sheet_path not in names:
                raise DocumentExtractionError("XLSX worksheet part is missing.")
            sheet_name = sheet.get("name", "Worksheet")
            state = sheet.get("state", "visible")
            context = (
                f"Hidden worksheet: Worksheet {number}"
                if state != "visible"
                else f"Worksheet {number}",
            )
            paths[sheet_name] = context
            sheet_root = self._parse_xml(archive.read(sheet_path), f"XLSX worksheet {sheet_name}")
            merged = self._xlsx_merges(
                sheet_root, self._max_visited_spreadsheet_cells - merged_cells
            )
            merges[sheet_name] = merged
            merged_cells += merged.cell_count
            sheet_relationships = self._relationships(
                archive, ExtractionBase._relationship_part(sheet_path)
            )
            for drawing in sheet_root.findall(f".//{{{S}}}drawing"):
                drawing_relationship = sheet_relationships.get(drawing.get(f"{{{R}}}id", ""))
                if drawing_relationship is None or drawing_relationship[1]:
                    continue
                drawing_path = posixpath.normpath(
                    posixpath.join(posixpath.dirname(sheet_path), drawing_relationship[0])
                ).lstrip("/")
                drawing_relationships = self._relationships(
                    archive, ExtractionBase._relationship_part(drawing_path)
                )
                for target, external in drawing_relationships.values():
                    if external:
                        continue
                    package_path = posixpath.normpath(
                        posixpath.join(posixpath.dirname(drawing_path), target)
                    ).lstrip("/")
                    if package_path.startswith("xl/media/"):
                        media.setdefault(package_path, set()).add(context)
                    elif package_path.startswith("xl/charts/"):
                        charts.setdefault(package_path, set()).add(context)
        return media, charts, merges, paths

    @staticmethod
    def _xlsx_merges(root: ElementTree.Element, remaining_cells: int) -> _WorksheetMerges:
        rows: dict[int, list[_SpreadsheetMerge]] = {}
        max_row = max_column = cells = 0
        for node in root.findall(f"{{{S}}}mergeCells/{{{S}}}mergeCell"):
            try:
                first_column, first_row, last_column, last_row = range_boundaries(
                    node.get("ref", "")
                )
            except ValueError as exc:
                raise DocumentExtractionError("XLSX merged range is invalid.") from exc
            if (
                first_column is None
                or first_row is None
                or last_column is None
                or last_row is None
                or first_column < 1
                or first_row < 1
                or last_column < first_column
                or last_row < first_row
                or (first_column == last_column and first_row == last_row)
            ):
                raise DocumentExtractionError("XLSX merged range is invalid.")
            if last_column > MAX_WORKBOOK_COLUMNS or last_row > MAX_WORKBOOK_ROWS:
                raise UnsupportedDocumentError(
                    "XLSX merged range exceeds the safe extraction bounds."
                )
            cells += (last_column - first_column + 1) * (last_row - first_row + 1)
            if cells > remaining_cells:
                raise UnsupportedDocumentError("XLSX merged ranges exceed the visited cells limit.")
            merged = _SpreadsheetMerge(first_column, first_row, last_column, last_row)
            max_row, max_column = max(max_row, last_row), max(max_column, last_column)
            for row in range(first_row, last_row + 1):
                rows.setdefault(row, []).append(merged)
        ordered: dict[int, tuple[_SpreadsheetMerge, ...]] = {}
        for row, ranges in rows.items():
            ranges.sort(key=lambda merged: merged.first_column)
            if any(
                left.last_column >= right.first_column
                for left, right in zip(ranges, ranges[1:], strict=False)
            ):
                raise DocumentExtractionError("XLSX merged ranges overlap.")
            ordered[row] = tuple(ranges)
        return _WorksheetMerges(ordered, max_row, max_column, cells)
