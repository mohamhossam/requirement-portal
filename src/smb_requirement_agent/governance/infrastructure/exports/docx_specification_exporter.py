"""Word adapter for the Business Requirements Document.

python-docx writes zip entries stamped with the wall clock, so the archive is rewritten with
fixed timestamps: the same revision always renders to the same bytes.
"""

from __future__ import annotations

from datetime import UTC
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.document import Document as WordDocument

from smb_requirement_agent.governance.application.errors import SpecificationExportFormatError
from smb_requirement_agent.governance.application.specification import (
    SPECIFICATION_GENERATOR,
    SPECIFICATION_TITLE,
    RequirementSpecification,
    SpecificationFormat,
)
from smb_requirement_agent.governance.infrastructure.exports.specification_outline import (
    Block,
    Bullets,
    Heading,
    Paragraph,
    Table,
    outline,
)
from smb_requirement_agent.governance.infrastructure.exports.xml_text import require_xml_safe

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_EPOCH = (1980, 1, 1, 0, 0, 0)


class DocxSpecificationExporter:
    @property
    def format(self) -> SpecificationFormat:
        return SpecificationFormat.DOCX

    @property
    def extension(self) -> str:
        return "docx"

    @property
    def media_type(self) -> str:
        return DOCX_MEDIA_TYPE

    def render(self, document: RequirementSpecification) -> bytes:
        word = Document()
        _properties(word, document)
        for block in outline(document):
            _render(word, block)
        buffer = BytesIO()
        word.save(buffer)
        return _stable_zip(buffer.getvalue())


def _properties(word: WordDocument, document: RequirementSpecification) -> None:
    manifest = document.manifest
    stamp = manifest.revision_created_at.astimezone(UTC).replace(tzinfo=None)
    properties = word.core_properties
    properties.title = _safe(f"{SPECIFICATION_TITLE} — {manifest.title}")
    properties.subject = _safe(manifest.requirement_id)
    properties.category = SPECIFICATION_TITLE
    properties.author = SPECIFICATION_GENERATOR
    properties.last_modified_by = SPECIFICATION_GENERATOR
    properties.comments = f"Document schema {document.schema_version}"
    properties.created = stamp
    properties.modified = stamp
    properties.revision = 1


def _render(word: WordDocument, block: Block) -> None:
    match block:
        case Heading(level, text):
            word.add_heading(_safe(text), min(level, 9))
        case Paragraph(text, note):
            paragraph = word.add_paragraph()
            run = paragraph.add_run(_safe(text))
            run.italic = note
        case Bullets(items):
            for item in items:
                word.add_paragraph(_safe(item), style="List Bullet")
        case Table(headers, rows):
            table = word.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            for cell, header in zip(table.rows[0].cells, headers, strict=True):
                cell.text = ""
                cell.paragraphs[0].add_run(_safe(header)).bold = True
            for row in rows:
                for cell, value in zip(table.add_row().cells, row, strict=True):
                    cell.text = _safe(value)


def _safe(value: str) -> str:
    return require_xml_safe(
        value,
        lambda: SpecificationExportFormatError(
            "This content cannot be represented in a Word document; download it as Markdown."
        ),
    )


def _stable_zip(payload: bytes) -> bytes:
    """Rewrite the archive with sorted entries and fixed timestamps."""
    with ZipFile(BytesIO(payload)) as source:
        names = sorted(source.namelist(), key=lambda name: (name != "[Content_Types].xml", name))
        output = BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as target:
            for name in names:
                entry = ZipInfo(name, date_time=_EPOCH)
                entry.compress_type = ZIP_DEFLATED
                entry.external_attr = 0o600 << 16
                target.writestr(entry, source.read(name))
    return output.getvalue()
