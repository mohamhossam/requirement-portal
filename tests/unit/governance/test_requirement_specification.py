"""The Business Requirements Document: contract, renderers and transport (ADR-0119)."""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from smb_requirement_agent.governance.application.errors import (
    RequirementSpecificationNotAvailableError,
    SpecificationExportFormatError,
)
from smb_requirement_agent.governance.application.specification import (
    GLOSSARY,
    RequirementSpecification,
    SpecificationActor,
    SpecificationBacklog,
    SpecificationCounts,
    SpecificationCriterion,
    SpecificationEpic,
    SpecificationEvidence,
    SpecificationFeature,
    SpecificationFormat,
    SpecificationManifest,
    SpecificationNonFunctionalRequirement,
    SpecificationObjectives,
    SpecificationResolution,
    SpecificationResolvable,
    SpecificationRisk,
    SpecificationScope,
    SpecificationSignature,
    SpecificationSourceDocument,
    SpecificationStakeholder,
    SpecificationStatement,
    SpecificationStory,
    SpecificationSummary,
    SpecificationUseCase,
    SpecificationVersion,
    SpecificationVersionKind,
    StatementOrigin,
)
from smb_requirement_agent.governance.application.use_cases.export_requirement_specification import (  # noqa: E501
    ExportRequirementSpecification,
    build_requirement_specification,
    is_specification_available,
    select_requirement_revision,
)
from smb_requirement_agent.governance.domain.revision.entities import RevisionNumber
from smb_requirement_agent.governance.infrastructure.exports.docx_specification_exporter import (
    DocxSpecificationExporter,
)
from smb_requirement_agent.governance.infrastructure.exports.markdown_specification_exporter import (  # noqa: E501
    MarkdownSpecificationExporter,
)
from smb_requirement_agent.interfaces.api.container import Container
from smb_requirement_agent.shared_kernel.identifiers import RequirementId
from tests.unit.workflow_helpers import approve_fake_breakdown, confirm_fake_analysis, post_analysis

NOW = datetime(2026, 9, 4, 8, 30, tzinfo=UTC)
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _confirm_only(client: TestClient) -> tuple[str, int]:
    """A Requirement whose analysis is confirmed and whose backlog does not exist yet."""
    requirement_id = client.post(
        "/requirements",
        json={
            "title": "Confirmed need",
            "description": "Describe the need in business language.",
            "desired_outcome": "Customers order in one visit.",
            "channels": ["Web"],
            "systems": ["CRM"],
            "business_rules": ["Orders need a signed contract."],
            "constraints": ["Go live before the fiscal year ends."],
        },
    ).json()["id"]
    assert post_analysis(client, requirement_id).succeeded
    confirm_fake_analysis(client, requirement_id)
    history = client.get(f"/requirements/{requirement_id}/revisions").json()
    revision = next(
        item["number"]
        for item in reversed(history["breakdown_revisions"])
        if item["specification_available"]
    )
    return str(requirement_id), int(revision)


def _specification(*, with_backlog: bool, resolved: bool = True) -> RequirementSpecification:
    """A document built by hand, so a renderer is tested on its own."""
    owner = SpecificationActor("Amina Owner")
    resolution = SpecificationResolution("Yes, in every channel.", owner, NOW) if resolved else None
    backlog = (
        SpecificationBacklog(
            epic=SpecificationEpic("Order in one visit", "Fewer visits", "Retention"),
            features=(
                SpecificationFeature(
                    id="FR-01",
                    name="Ordering",
                    outcome="Customers can order online.",
                    delivery_drop="mvp",
                    kind="business",
                    acceptance_criteria=("An order completes online.",),
                    stories=(
                        SpecificationStory(
                            id="FR-01.1",
                            role="SMB customer",
                            action="place an order",
                            value="I do not visit a shop",
                            kind="business",
                            criteria=(
                                SpecificationCriterion("a signed contract", "I order", "it ships"),
                            ),
                        ),
                    ),
                ),
            ),
            approval=SpecificationSignature(owner, NOW, "Ready for delivery."),
            dependencies=("CRM exposes the contract status.",),
            risks=(SpecificationRisk("warning", "CRM release slips."),),
        )
        if with_backlog
        else None
    )
    return RequirementSpecification(
        schema_version="1.0",
        manifest=SpecificationManifest(
            requirement_id="req-1",
            title="Order in one visit",
            requirement_version=2,
            breakdown_revision=7,
            revision_created_at=NOW,
            analysis_round=1,
            analysis_generated_at=NOW,
            analysis_model="fake",
            analysis_prompt_version="v1",
            confirmation=SpecificationSignature(owner, NOW),
            backlog_approval=backlog.approval if backlog else None,
        ),
        summary=SpecificationSummary(
            "Customers must | be able to order\nwithout a visit.",
            SpecificationCounts(1, 1, 1, 1, 0, 1, 0, 1 if resolved else 0),
        ),
        objectives=SpecificationObjectives(
            "Customers order in one visit.", StatementOrigin.CAPTURED, ("Visits halve",)
        ),
        scope=SpecificationScope("Small businesses", ("Web",), ("CRM",)),
        stakeholders=(SpecificationStakeholder("Amina Owner", "Requirement Owner", NOW),),
        business_requirements=(
            SpecificationStatement(
                "BR-01",
                "Orders are placed online.",
                StatementOrigin.ANALYSIS,
                (SpecificationEvidence("brief.docx", "Section 2"),),
            ),
        ),
        business_rules=(
            SpecificationStatement(
                "RL-01", "Orders need a signed contract.", StatementOrigin.CAPTURED
            ),
        ),
        constraints=(
            SpecificationStatement(
                "CN-01", "Go live before the fiscal year ends.", StatementOrigin.ACCEPTED_PROPOSAL
            ),
        ),
        assumptions=(
            SpecificationResolvable("AS-01", "Every channel sells the plan.", None, (), resolution),
        ),
        dependencies=(),
        open_questions=(
            SpecificationResolvable(
                "OQ-01", "Which channels?", "Scope of the first drop.", (), resolution
            ),
        ),
        ambiguities=(),
        product_impact=None,
        functional_requirements=backlog,
        source_documents=(SpecificationSourceDocument("brief.docx", "v1", "ab" * 32),),
        versions=(
            SpecificationVersion(SpecificationVersionKind.REQUIREMENT, 1, NOW, "Need captured"),
            SpecificationVersion(SpecificationVersionKind.BREAKDOWN, 7, NOW, "Analysis confirmed"),
        ),
        non_functional_requirements=(
            SpecificationNonFunctionalRequirement(
                "NF-01", "Order confirmation within 5 seconds.", "performance"
            ),
        ),
        business_use_cases=(
            SpecificationUseCase(
                id="BUC-01",
                name="Customer orders a plan",
                actor="SMB customer",
                trigger=None,
                goal="The plan is active without a shop visit.",
                main_steps=("Choose a plan", "Place the order"),
                alternate_paths=("Eligibility fails: the customer is told why.",),
                business_rules=("Orders need a signed contract.",),
            ),
        ),
    )


def _headings(content: bytes) -> list[str]:
    word = Document(BytesIO(content))
    return [
        paragraph.text
        for paragraph in word.paragraphs
        if paragraph.style is not None
        and paragraph.style.name is not None
        and paragraph.style.name.startswith(("Heading", "Title"))
    ]


# --- builder --------------------------------------------------------------------------------


def test_confirmed_analysis_without_backlog_builds_a_document_without_functional_requirements(
    client: TestClient, container: Container
) -> None:
    requirement_id, revision_number = _confirm_only(client)
    revisions = container.breakdown_repository
    breakdown = revisions.get_breakdown_revision(
        RequirementId(requirement_id), RevisionNumber(revision_number)
    )
    assert breakdown is not None and is_specification_available(breakdown)
    requirement_revision = select_requirement_revision(
        revisions.list_requirement_revisions(RequirementId(requirement_id)), breakdown
    )
    assert requirement_revision is not None

    document = build_requirement_specification(
        breakdown,
        requirement_revision,
        revisions.list_requirement_revisions(RequirementId(requirement_id)),
        revisions.list_breakdown_revisions(RequirementId(requirement_id)),
    )

    assert document.schema_version == "1.4"
    assert document.manifest.title == "Confirmed need"
    assert document.manifest.confirmation.actor == SpecificationActor("Amina Owner")
    assert document.manifest.backlog_approval is None
    assert document.functional_requirements is None
    assert document.objectives.desired_outcome == "Customers order in one visit."
    assert document.objectives.origin is StatementOrigin.CAPTURED
    assert document.scope.channels == ("Web",)
    assert [item.id for item in document.business_requirements] == ["BR-01"]
    assert [(item.id, item.origin) for item in document.business_rules] == [
        ("RL-01", StatementOrigin.CAPTURED),
        ("RL-02", StatementOrigin.ANALYSIS),
    ]
    assert document.constraints[0].text == "Go live before the fiscal year ends."
    # The fake answers every uncertainty; each stays in the record as resolved.
    for rows in (
        document.assumptions,
        document.dependencies,
        document.open_questions,
        document.ambiguities,
    ):
        assert rows and all(item.resolution is not None for item in rows)
    assert document.assumptions[0].resolution is not None
    assert document.assumptions[0].resolution.answer == "Confirmed."
    assert document.summary.counts.resolved_clarifications == 4
    assert [item.role for item in document.stakeholders] == ["Requirement Owner"]
    assert document.glossary == GLOSSARY


def test_approved_revision_adds_the_backlog_as_functional_requirements(
    client: TestClient, container: Container
) -> None:
    requirement_id, _, revision_number = approve_fake_breakdown(client)
    revisions = container.breakdown_repository
    breakdown = revisions.get_breakdown_revision(
        RequirementId(requirement_id), RevisionNumber(revision_number)
    )
    assert breakdown is not None
    requirement_revision = select_requirement_revision(
        revisions.list_requirement_revisions(RequirementId(requirement_id)), breakdown
    )
    assert requirement_revision is not None

    document = build_requirement_specification(
        breakdown,
        requirement_revision,
        revisions.list_requirement_revisions(RequirementId(requirement_id)),
        revisions.list_breakdown_revisions(RequirementId(requirement_id)),
    )

    backlog = document.functional_requirements
    assert backlog is not None
    assert backlog.approval.rationale == "Approved for portable export."
    assert document.manifest.backlog_approval == backlog.approval
    assert [feature.id for feature in backlog.features] == ["FR-01", "FR-02"]
    assert [story.id for story in backlog.features[0].stories] == ["FR-01.1", "FR-01.2"]
    assert backlog.features[0].stories[0].criteria[0].given == (
        "the customer is eligible for the approved Feature"
    )
    summaries = [item.summary for item in document.versions]
    assert summaries[0] == "Need captured"
    for expected in (
        "Analysis generated, round 1",
        "Analysis re-run, round 2; Clarification answered",
        "Analysis confirmed",
        "Epic generated",
        "Features generated (2)",
        "Review generated",
        "Stories generated (2)",
        "Submitted for review",
        "Final approval recorded",
    ):
        assert any(expected in summary for summary in summaries), (expected, summaries)
    confirmed = next(i for i, s in enumerate(summaries) if s == "Analysis confirmed")
    assert confirmed < summaries.index("Epic generated")
    assert summaries[-1] == "Final approval recorded"
    assert [item.kind for item in document.versions][0] is SpecificationVersionKind.REQUIREMENT
    assert document.versions[-1].number == revision_number


def test_unconfirmed_revision_has_no_document(client: TestClient, container: Container) -> None:
    requirement_id = client.post(
        "/requirements", json={"title": "Raw", "description": "Not yet confirmed."}
    ).json()["id"]
    assert post_analysis(client, requirement_id).succeeded
    revisions = container.breakdown_repository
    breakdown = revisions.list_breakdown_revisions(RequirementId(requirement_id))[-1]
    requirement_revision = revisions.list_requirement_revisions(RequirementId(requirement_id))[-1]

    assert not is_specification_available(breakdown)
    with pytest.raises(RequirementSpecificationNotAvailableError):
        build_requirement_specification(
            breakdown,
            requirement_revision,
            revisions.list_requirement_revisions(RequirementId(requirement_id)),
            revisions.list_breakdown_revisions(RequirementId(requirement_id)),
        )


def test_document_is_built_from_the_requirement_version_the_analysis_saw(
    client: TestClient, container: Container
) -> None:
    requirement_id, revision_number = _confirm_only(client)
    current = client.get(f"/requirements/{requirement_id}").json()
    edited = client.put(
        f"/requirements/{requirement_id}",
        json={
            "title": "Renamed after confirmation",
            "description": current["description"],
            "expected_version": current["version"],
            "impact_acknowledged": True,
        },
    )
    assert edited.status_code == 200, edited.text

    response = client.get(
        f"/requirements/{requirement_id}/revisions/{revision_number}/specification?format=md"
    )

    assert response.status_code == 200
    assert response.text.startswith("# Confirmed need\n")
    assert "Renamed after confirmation" not in response.text


# --- renderers ------------------------------------------------------------------------------


def test_word_document_carries_every_section_in_order_and_is_byte_stable() -> None:
    exporter = DocxSpecificationExporter()
    document = _specification(with_backlog=True)

    first = exporter.render(document)
    second = exporter.render(document)

    assert first == second
    assert first.startswith(b"PK")
    headings = _headings(first)
    assert headings[0] == "Order in one visit"
    numbered = [text for text in headings if text[:1].isdigit()]
    assert numbered == [
        "1. Executive summary",
        "2. Business objectives",
        "3. Scope",
        "4. Stakeholders",
        "5. Actors and scenarios",
        "6. Business requirements",
        "7. Business use cases",
        "8. Business rules",
        "9. Constraints",
        "10. Non-functional requirements",
        "11. Assumptions",
        "12. Dependencies",
        "13. Open questions",
        "14. Ambiguities",
        "15. Product impact",
        "16. Functional requirements",
        "17. Risks",
        "18. Source documents",
        "19. Approvals",
        "20. Glossary",
    ]
    assert "Version history" in headings
    assert "BUC-01 Customer orders a plan" in headings
    assert "FR-01 Ordering" in headings
    assert any(text.startswith("FR-01.1 As a SMB customer") for text in headings)
    word = Document(BytesIO(first))
    control = word.tables[0]
    assert [cell.text for cell in control.rows[0].cells] == ["Item", "Value"]
    assert ("Backlog approval", "Amina Owner, 2026-09-04 08:30 UTC — Ready for delivery.") in {
        (row.cells[0].text, row.cells[1].text) for row in control.rows
    }
    versions = word.tables[1]
    assert [cell.text for cell in versions.rows[0].cells] == [
        "Kind",
        "Version",
        "Recorded",
        "Change",
    ]
    assert [cell.text for cell in versions.rows[1].cells] == [
        "Need",
        "1",
        "2026-09-04 08:30 UTC",
        "Need captured",
    ]
    assert [cell.text for cell in versions.rows[2].cells][3] == "Analysis confirmed"
    actors = next(table for table in word.tables if table.rows[0].cells[0].text == "Actor")
    assert [cell.text for cell in actors.rows[1].cells] == [
        "SMB customer",
        "FR-01.1",
        "Ordering",
        "place an order",
    ]
    assert word.core_properties.author == "SMB AI Requirement Breakdown Agent"
    assert word.core_properties.created == NOW
    assert word.core_properties.modified == NOW
    assert word.core_properties.title == "Business Requirements Document — Order in one visit"


def test_word_document_without_approved_backlog_says_so() -> None:
    content = DocxSpecificationExporter().render(_specification(with_backlog=False, resolved=False))

    word = Document(BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    assert "functional requirements are not included in this document" in text
    assert "FR-01" not in text
    assert "Actors and their scenarios come from the approved backlog" in text
    assert "Risks come from the approved backlog" in text
    assert ("Backlog approval", "Not yet approved") in {
        (row.cells[0].text, row.cells[1].text) for row in word.tables[0].rows
    }
    statuses = {
        row.cells[-1].text
        for table in word.tables
        for row in table.rows
        if row.cells[0].text == "AS-01"
    }
    assert statuses == {"Open"}


def test_word_document_refuses_control_characters() -> None:
    document = _specification(with_backlog=False)
    broken = RequirementSpecification(
        **{
            **document.__dict__,
            "summary": SpecificationSummary("Bad \x07 text", document.summary.counts),
        }
    )

    with pytest.raises(SpecificationExportFormatError):
        DocxSpecificationExporter().render(broken)


def test_markdown_renders_the_same_outline_with_escaped_tables() -> None:
    content = MarkdownSpecificationExporter().render(_specification(with_backlog=True))

    text = content.decode("utf-8")
    assert not text.startswith("﻿")
    assert "\r" not in text
    assert text.endswith("\n") and not text.endswith("\n\n")
    assert text.startswith(
        "# Order in one visit\n\nBusiness Requirements Document — Requirement req-1"
    )
    assert "## 1. Executive summary\n\nCustomers must | be able to order\nwithout a visit." in text
    assert "## 16. Functional requirements" in text
    assert "## 7. Business use cases\n\n### BUC-01 Customer orders a plan" in text
    assert "| NF-01 | Order confirmation within 5 seconds. | Performance | — |" in text
    assert "1. Choose a plan\n2. Place the order" in text
    assert "### FR-01 Ordering" in text
    assert "### Version history\n\n| Kind | Version | Recorded | Change |" in text
    assert "| Need | 1 | 2026-09-04 08:30 UTC | Need captured |" in text
    assert "## 5. Actors and scenarios\n\n| Actor | Scenario | Feature | Goal |" in text
    assert "| SMB customer | FR-01.1 | Ordering | place an order |" in text
    assert "## 17. Risks\n\n| Severity | Risk |" in text
    assert (
        "#### FR-01.1 As a SMB customer, I want place an order, so that I do not visit a shop"
        in text
    )
    assert "| a signed contract | I order | it ships |" in text
    assert (
        "| BR-01 | Orders are placed online. | Found by the analysis | brief.docx (Section 2) |"
        in text
    )
    assert (
        "| OQ-01 | Which channels? | Scope of the first drop. | — | "
        "Resolved: Yes, in every channel. (Amina Owner, 2026-09-04 08:30 UTC) |"
    ) in text
    assert "| warning |" not in text and "| Warning | CRM release slips. |" in text
    lines = [line for line in text.splitlines() if line.startswith("## ")]
    assert lines[0] == "## Document control"
    assert lines[1:4] == ["## 1. Executive summary", "## 2. Business objectives", "## 3. Scope"]
    assert lines[-1] == "## 20. Glossary"


def test_markdown_and_word_hold_no_email_addresses() -> None:
    document = _specification(with_backlog=True)
    markdown = MarkdownSpecificationExporter().render(document).decode("utf-8")
    word = Document(BytesIO(DocxSpecificationExporter().render(document)))
    text = "\n".join(paragraph.text for paragraph in word.paragraphs) + "\n".join(
        cell.text for table in word.tables for row in table.rows for cell in row.cells
    )

    assert "@" not in markdown
    assert "@" not in text


def test_exporters_declare_distinct_formats_and_media_types() -> None:
    docx, markdown = DocxSpecificationExporter(), MarkdownSpecificationExporter()

    assert (docx.format, docx.extension, docx.media_type) == (
        SpecificationFormat.DOCX,
        "docx",
        DOCX_MEDIA_TYPE,
    )
    assert (markdown.format, markdown.extension, markdown.media_type) == (
        SpecificationFormat.MARKDOWN,
        "md",
        "text/markdown; charset=utf-8",
    )
    with pytest.raises(ValueError):
        ExportRequirementSpecification(
            requirements=None,  # type: ignore[arg-type]
            authorization=None,  # type: ignore[arg-type]
            revisions=None,  # type: ignore[arg-type]
            exporters=(docx, DocxSpecificationExporter()),
        )


# --- transport ------------------------------------------------------------------------------


def test_document_downloads_in_both_formats_with_attachment_headers(client: TestClient) -> None:
    requirement_id, revision = _confirm_only(client)
    base = f"/requirements/{requirement_id}/revisions/{revision}/specification"

    word = client.get(f"{base}?format=docx")
    markdown = client.get(f"{base}?format=md")

    assert word.status_code == 200
    assert word.headers["content-type"] == DOCX_MEDIA_TYPE
    assert word.headers["content-disposition"] == (
        'attachment; filename="'
        f'requirement-{requirement_id}-business-requirements-v{revision}.docx"'
    )
    assert word.headers["cache-control"] == "private, no-store"
    assert word.content.startswith(b"PK")
    assert word.content == client.get(f"{base}?format=docx").content
    assert markdown.status_code == 200
    assert markdown.headers["content-type"] == "text/markdown; charset=utf-8"
    assert markdown.headers["content-disposition"].endswith(
        f'business-requirements-v{revision}.md"'
    )
    assert markdown.text.startswith("# Confirmed need\n")
    assert "functional requirements are not included in this document" in markdown.text


def test_approved_revision_document_includes_functional_requirements(client: TestClient) -> None:
    requirement_id, _, revision = approve_fake_breakdown(client)

    response = client.get(
        f"/requirements/{requirement_id}/revisions/{revision}/specification?format=md"
    )

    assert response.status_code == 200
    assert "### FR-01 " in response.text
    assert "Approved for portable export." in response.text
    assert "functional requirements are not included" not in response.text


def test_document_is_reserved_for_the_requirement_team(client: TestClient) -> None:
    requirement_id, revision = _confirm_only(client)
    access = client.get(f"/requirements/{requirement_id}/assignments").json()
    assert (
        client.put(
            f"/requirements/{requirement_id}/reviewers/fake-reviewer",
            json={"expected_version": access["version"]},
        ).status_code
        == 200
    )
    path = f"/requirements/{requirement_id}/revisions/{revision}/specification?format=md"

    assert client.get(path, headers={"X-Fake-Actor-Id": "fake-reviewer"}).status_code == 200
    assert client.get(path, headers={"X-Fake-Actor-Id": "fake-observer"}).status_code == 403


def test_document_refuses_unconfirmed_missing_and_unknown_formats(client: TestClient) -> None:
    requirement_id = client.post(
        "/requirements", json={"title": "Raw", "description": "Not yet confirmed."}
    ).json()["id"]
    assert post_analysis(client, requirement_id).succeeded
    history = client.get(f"/requirements/{requirement_id}/revisions").json()
    revision = history["breakdown_revisions"][-1]
    assert revision["analysis_human_confirmed"] is False
    assert revision["specification_available"] is False
    assert revision["analysis_confirmed_by"] is None
    base = f"/requirements/{requirement_id}/revisions"

    unconfirmed = client.get(f"{base}/{revision['number']}/specification?format=md")
    assert unconfirmed.status_code == 409
    assert unconfirmed.json()["code"] == "requirement_specification_not_available"
    assert client.get(f"{base}/999/specification?format=md").status_code == 404
    assert client.get(f"{base}/{revision['number']}/specification?format=pdf").status_code == 422
    assert (
        client.get("/requirements/missing/revisions/1/specification?format=md").status_code == 404
    )


def test_revision_history_names_the_confirmation(client: TestClient) -> None:
    requirement_id, revision = _confirm_only(client)

    history = client.get(f"/requirements/{requirement_id}/revisions").json()

    item = next(row for row in history["breakdown_revisions"] if row["number"] == revision)
    assert item["specification_available"] is True
    assert item["exportable"] is False
    assert item["analysis_confirmed_by"]["display_name"] == "Amina Owner"
    assert item["analysis_confirmed_at"] is not None
    first = history["breakdown_revisions"][0]
    assert first["specification_available"] is False
    assert first["analysis_confirmed_at"] is None
