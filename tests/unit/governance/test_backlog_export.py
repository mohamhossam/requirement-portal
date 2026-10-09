"""Neutral JSON/XLSX export behavior and transport contracts."""

from __future__ import annotations

import json
from dataclasses import replace
from datetime import UTC, datetime
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from smb_requirement_agent.governance.application.errors import BacklogExportFormatError
from smb_requirement_agent.governance.application.exports import (
    ExportAcceptanceCriterion,
    ExportActor,
    ExportApproval,
    ExportArchitecture,
    ExportArchitectureDependency,
    ExportCapability,
    ExportCounts,
    ExportEpic,
    ExportFeature,
    ExportManifest,
    ExportOrganisationReference,
    ExportProvenance,
    ExportStory,
    ExportSystem,
    NeutralBacklogExport,
)
from smb_requirement_agent.governance.application.use_cases.export_breakdown import (
    ExportBreakdown,
    formal_final_approval,
    is_exportable_revision,
)
from smb_requirement_agent.governance.domain.review.entities import BreakdownStatus
from smb_requirement_agent.governance.domain.revision.entities import RevisionNumber
from smb_requirement_agent.governance.infrastructure.exports.json_exporter import (
    JsonBacklogExporter,
)
from smb_requirement_agent.governance.infrastructure.exports.xlsx_exporter import (
    XlsxBacklogExporter,
)
from smb_requirement_agent.identity.domain.entities import (
    DraftOwnership,
    RequirementAccess,
)
from smb_requirement_agent.identity.domain.errors import AuthorizationDeniedError
from smb_requirement_agent.interfaces.api.container import Container
from smb_requirement_agent.shared_kernel.actors import (
    ActorId,
    ActorProfile,
)
from smb_requirement_agent.shared_kernel.identifiers import RequirementId
from tests.unit.access_service import access_service_for
from tests.unit.workflow_helpers import approve_fake_breakdown, post_analysis

NOW = datetime(2026, 9, 4, 8, 30, tzinfo=UTC)


class UnownedAccessRepository:
    """Represent a legacy Requirement without an ownership record."""

    def get_requirement(self, requirement_id: RequirementId) -> RequirementAccess | None:
        return None

    def save_requirement(self, access: RequirementAccess) -> None:
        raise AssertionError("Export must not write access state.")

    def get_draft_ownership(self, draft_id: RequirementId) -> DraftOwnership | None:
        return None

    def save_draft_ownership(self, ownership: DraftOwnership) -> None:
        raise AssertionError("Export must not write draft state.")

    def delete_draft_ownership(self, draft_id: RequirementId) -> None:
        raise AssertionError("Export must not write draft state.")


def test_formally_approved_revision_exports_stable_nested_json(client: TestClient) -> None:
    requirement_id, _, revision = approve_fake_breakdown(client)
    path = f"/requirements/{requirement_id}/revisions/{revision}/export?format=json"

    first = client.get(path)
    second = client.get(path)

    assert first.status_code == 200
    assert first.content == second.content
    assert first.content.endswith(b"\n")
    assert first.headers["content-type"].startswith("application/json")
    assert first.headers["content-disposition"] == (
        f'attachment; filename="requirement-{requirement_id}-breakdown-v{revision}.json"'
    )
    assert first.headers["cache-control"] == "private, no-store"
    payload = json.loads(first.content)
    assert list(payload) == ["schema_version", "manifest", "epic"]
    assert list(payload["manifest"]) == [
        "requirement_id",
        "breakdown_revision",
        "revision_created_at",
        "final_approval",
        "counts",
    ]
    assert payload["schema_version"] == "1.5"
    assert payload["manifest"]["breakdown_revision"] == revision
    assert payload["manifest"]["final_approval"]["recorded_by"]["id"] == "fake-owner"
    assert payload["manifest"]["final_approval"]["rationale"] == ("Approved for portable export.")
    assert payload["epic"]["features"][0]["stories"][0]["acceptance_criteria"][0] == {
        "sequence": 1,
        "given": "the customer is eligible for the approved Feature",
        "when": payload["epic"]["features"][0]["stories"][0]["action"],
        "then": "the requested capability is completed successfully",
    }
    assert "analysis" not in payload
    assert "comments" not in payload
    assert "approvals" not in payload["epic"]


def test_xlsx_download_preserves_hierarchy_and_given_when_then(client: TestClient) -> None:
    requirement_id, _, revision = approve_fake_breakdown(client)

    response = client.get(f"/requirements/{requirement_id}/revisions/{revision}/export?format=xlsx")

    assert response.status_code == 200
    assert response.content.startswith(b"PK")
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == [
        "Manifest",
        "Epic",
        "Features",
        "Stories",
        "Acceptance Criteria",
        "Systems",
        "Dependencies",
        "Connected Systems",
        "Product Context",
        "Journey Steps",
    ]
    assert [cell.value for cell in workbook["Features"][1]] == [
        "sequence",
        "id",
        "epic_id",
        "name",
        "outcome",
        "delivery_drop",
        "splitting_pattern",
        "splitting_rationale",
        "status",
        "generated_at",
        "model",
        "prompt_version",
    ]
    for sheet in workbook.worksheets:
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == sheet.dimensions
    feature_ids = {row[1] for row in workbook["Features"].iter_rows(min_row=2, values_only=True)}
    story_rows = list(workbook["Stories"].iter_rows(min_row=2, values_only=True))
    assert story_rows
    assert {row[3] for row in story_rows} <= feature_ids
    criterion_rows = list(workbook["Acceptance Criteria"].iter_rows(min_row=2, values_only=True))
    assert {row[1] for row in criterion_rows} <= {row[2] for row in story_rows}
    assert criterion_rows[0][3:] == (
        "the customer is eligible for the approved Feature",
        story_rows[0][5],
        "the requested capability is completed successfully",
    )


def test_export_uses_current_team_access_and_keeps_historical_approval(
    client: TestClient,
) -> None:
    requirement_id, stories, revision = approve_fake_breakdown(client)
    access = client.get(f"/requirements/{requirement_id}/assignments").json()
    assert (
        client.put(
            f"/requirements/{requirement_id}/reviewers/fake-reviewer",
            json={"expected_version": access["version"]},
        ).status_code
        == 200
    )
    reviewer_headers = {"X-Fake-Actor-Id": "fake-reviewer"}
    observer_headers = {"X-Fake-Actor-Id": "fake-observer"}

    assert (
        client.get(
            f"/requirements/{requirement_id}/revisions/{revision}/export?format=json",
            headers=reviewer_headers,
        ).status_code
        == 200
    )
    assert (
        client.get(
            f"/requirements/{requirement_id}/revisions/{revision}/export?format=json",
            headers=observer_headers,
        ).status_code
        == 403
    )

    story = stories[0]
    edited = client.put(
        f"/requirements/{requirement_id}/features/{story['feature_id']}/stories/{story['id']}",
        json={
            "role": story["role"],
            "action": f"{story['action']} after final approval",
            "value": story["value"],
            "acceptance_criteria": story["acceptance_criteria"],
            "expected_version": int(str(story["version"])) + 1,
        },
    )
    assert edited.status_code == 200
    historical = client.get(
        f"/requirements/{requirement_id}/revisions/{revision}/export?format=json"
    )
    assert historical.status_code == 200
    assert (
        json.loads(historical.content)["epic"]["features"][0]["stories"][0]["action"]
        == story["action"]
    )


def test_unowned_legacy_requirement_must_be_claimed_before_export(
    client: TestClient, container: Container
) -> None:
    requirement_id, _, revision = approve_fake_breakdown(client)
    use_case = ExportBreakdown(
        container.requirement_repository,
        access_service_for(container.requirement_repository, UnownedAccessRepository()),
        container.breakdown_repository,
        container.backlog_exporters,
    )

    with pytest.raises(AuthorizationDeniedError):
        use_case.execute(
            RequirementId(requirement_id),
            RevisionNumber(revision),
            export_format=container.backlog_exporters[0].format,
            actor=ActorProfile(ActorId("legacy-user"), "Legacy User"),
        )


def test_ineligible_missing_and_invalid_format_errors_are_explicit(client: TestClient) -> None:
    requirement_id = client.post(
        "/requirements", json={"title": "Draft", "description": "Not approved."}
    ).json()["id"]
    assert post_analysis(client, requirement_id).status_code == 200
    history = client.get(f"/requirements/{requirement_id}/revisions").json()
    revision = history["breakdown_revisions"][-1]["number"]

    assert (
        client.get(
            f"/requirements/{requirement_id}/revisions/{revision}/export?format=json"
        ).status_code
        == 409
    )
    assert (
        client.get(f"/requirements/{requirement_id}/revisions/999/export?format=json").status_code
        == 404
    )
    assert (
        client.get(
            f"/requirements/{requirement_id}/revisions/{revision}/export?format=csv"
        ).status_code
        == 422
    )


def test_revision_summary_exposes_formal_approval_and_access_capability(
    client: TestClient,
) -> None:
    requirement_id, _, revision = approve_fake_breakdown(client)

    item = next(
        value
        for value in client.get(f"/requirements/{requirement_id}/revisions").json()[
            "breakdown_revisions"
        ]
        if value["number"] == revision
    )
    access = client.get(f"/requirements/{requirement_id}/assignments").json()

    assert item["exportable"] is True
    assert item["final_approved_by"]["id"] == "fake-owner"
    assert item["final_approved_at"]
    assert access["can_export_approved_revisions"] is True


def test_only_formal_final_approval_is_exportable(client: TestClient, container: Container) -> None:
    requirement_id, _, revision_number = approve_fake_breakdown(client)
    revision = container.breakdown_repository.get_breakdown_revision(
        RequirementId(requirement_id), RevisionNumber(revision_number)
    )
    assert revision is not None and revision.review is not None
    assert formal_final_approval(revision) is not None
    assert is_exportable_revision(revision)

    generated = replace(
        revision,
        review=replace(
            revision.review,
            status=BreakdownStatus.GENERATED,
            submitted_fingerprint=None,
            approvals=(),
        ),
    )
    under_review = replace(
        revision, review=replace(revision.review, status=BreakdownStatus.UNDER_REVIEW)
    )
    needs_revision = replace(
        revision, review=replace(revision.review, status=BreakdownStatus.NEEDS_REVISION)
    )
    legacy_label = replace(revision, review=replace(revision.review, approvals=()))
    mismatched_attestation = replace(
        revision,
        review=replace(revision.review, submitted_fingerprint="different-fingerprint"),
    )

    assert not is_exportable_revision(generated)
    assert not is_exportable_revision(under_review)
    assert not is_exportable_revision(needs_revision)
    assert not is_exportable_revision(legacy_label)
    assert not is_exportable_revision(mismatched_attestation)


def test_xlsx_treats_formula_like_text_as_text_and_rejects_lossy_cells() -> None:
    exporter = XlsxBacklogExporter()
    document = _document("=CMD()")

    workbook = load_workbook(BytesIO(exporter.render(document)), data_only=False)
    role_cell = workbook["Stories"]["E2"]
    assert role_cell.value == "=CMD()"
    assert role_cell.data_type == "s"

    oversized = replace(
        document,
        epic=replace(
            document.epic,
            features=(
                replace(
                    document.epic.features[0],
                    stories=(replace(document.epic.features[0].stories[0], role="x" * 32_768),),
                ),
            ),
        ),
    )
    with pytest.raises(BacklogExportFormatError, match="losslessly"):
        exporter.render(oversized)

    unsupported_control = replace(
        document,
        epic=replace(document.epic, name="Cannot encode\x01in Excel"),
    )
    with pytest.raises(BacklogExportFormatError, match="JSON"):
        exporter.render(unsupported_control)


def test_json_preserves_unicode_without_ascii_escaping() -> None:
    content = JsonBacklogExporter().render(_document("مراجع الأعمال"))

    assert "مراجع الأعمال".encode() in content
    assert b"\\u0645" not in content


def test_json_and_xlsx_preserve_architecture_references() -> None:
    document = _document("owner")

    payload = json.loads(JsonBacklogExporter().render(document))
    architecture = payload["epic"]["features"][0]["architecture"]
    assert architecture["knowledge_version"] == "catalog-v1"
    assert architecture["systems"][0]["squads"][0]["name"] == "Digital Squad"
    assert architecture["systems"][0]["value_streams"][0]["name"] == "Retail"
    assert architecture["dependencies"][0]["target_system_id"] == "billing"

    workbook = load_workbook(BytesIO(XlsxBacklogExporter().render(document)))
    system_rows = list(workbook["Systems"].iter_rows(min_row=2, values_only=True))
    dependency_rows = list(workbook["Dependencies"].iter_rows(min_row=2, values_only=True))
    assert system_rows[0][0:6] == (
        "feature",
        "feature-1",
        "catalog-v1",
        "2026-09-04T08:30:00Z",
        "portal",
        "SMB Portal",
    )
    assert dependency_rows[0][2:5] == ("portal", "billing", "Submit the order")
    assert len(system_rows) == 4  # mapped systems only, for the Feature and its Story
    connected_rows = list(workbook["Connected Systems"].iter_rows(min_row=2, values_only=True))
    assert connected_rows[0] == (
        "feature",
        "feature-1",
        "crm",
        "CRM",
        "Care Squad",
        "portal",
        "depends_on_mapped",
        "Customer data for ordering",
        "unspecified",
    )
    assert architecture["adjacent_systems"][0]["id"] == "crm"


def _document(role: str) -> NeutralBacklogExport:
    provenance = ExportProvenance(NOW, "fake-model", "prompt-v1")
    architecture = ExportArchitecture(
        knowledge_version="catalog-v1",
        mapped_at=NOW,
        systems=(
            ExportSystem(
                "portal",
                "SMB Portal",
                True,
                (ExportCapability("ordering", "Ordering"),),
                (ExportOrganisationReference("digital", "Digital Squad"),),
                (ExportOrganisationReference("retail", "Retail"),),
                (ExportOrganisationReference("ordering-app", "Ordering app"),),
            ),
            ExportSystem("billing", "Billing", True, (), (), (), ()),
        ),
        dependencies=(ExportArchitectureDependency("portal", "billing", "Submit the order"),),
        adjacent_systems=(
            ExportSystem(
                "crm",
                "CRM",
                True,
                (),
                (ExportOrganisationReference("care", "Care Squad"),),
                (),
                (),
            ),
        ),
        adjacent_dependencies=(
            ExportArchitectureDependency("crm", "portal", "Customer data for ordering"),
        ),
    )
    story = ExportStory(
        sequence=1,
        id="story-1",
        feature_id="feature-1",
        role=role,
        action="download content",
        value="the backlog is portable",
        voice=f"As a {role}, I want download content, so that the backlog is portable.",
        status="approved",
        provenance=provenance,
        acceptance_criteria=(ExportAcceptanceCriterion(1, "an approval", "I export", "it works"),),
        architecture=architecture,
    )
    feature = ExportFeature(
        sequence=1,
        id="feature-1",
        epic_id="epic-1",
        name="Portable export",
        outcome="A usable file",
        delivery_drop="mvp",
        splitting_pattern="component_system",
        splitting_rationale="Separate integration boundary",
        status="approved",
        provenance=provenance,
        architecture=architecture,
        stories=(story,),
    )
    return NeutralBacklogExport(
        schema_version="1.0",
        manifest=ExportManifest(
            requirement_id="requirement-1",
            breakdown_revision=4,
            revision_created_at=NOW,
            final_approval=ExportApproval(
                "approval-1",
                "sha256",
                ExportActor("owner-1", "Owner", None),
                NOW,
                None,
            ),
            counts=ExportCounts(1, 1, 1, 1),
        ),
        epic=ExportEpic(
            "epic-1",
            "requirement-1",
            "Export",
            "Portable backlog",
            "Enable integrations",
            "approved",
            provenance,
            (feature,),
        ),
    )
