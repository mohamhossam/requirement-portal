"""System components on a mapped impact, as requirement work receives them (ADR-0092).

Components in the catalogue, their files, suggestions and reading are the
knowledge service's job and are covered there. These tests cover what
requirement work does with the component it is told delivers a capability:
record, store and export it, and keep it out of approval fingerprints and
generation prompts.
"""

from __future__ import annotations

import io
import json
from dataclasses import replace
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from smb_requirement_agent.breakdown.application.ports.generation_guidance import GenerationGuidance
from smb_requirement_agent.breakdown.application.use_cases.architecture_mapping import (
    MapFeatureArchitecture,
)
from smb_requirement_agent.breakdown.domain.architecture.entities import ArchitectureImpact
from smb_requirement_agent.breakdown.infrastructure.backlog_codecs import (
    architecture_from_payload,
    architecture_to_payload,
)
from smb_requirement_agent.breakdown.infrastructure.llm.prompts.generation_guidance import (
    render_guidance,
)
from smb_requirement_agent.governance.domain.review.fingerprints import artifact_fingerprint
from smb_requirement_agent.governance.infrastructure.exports.json_exporter import (
    JsonBacklogExporter,
)
from smb_requirement_agent.governance.infrastructure.exports.xlsx_exporter import (
    XlsxBacklogExporter,
)
from smb_requirement_agent.references.application.ports.architecture_knowledge import (
    ArchitectureKnowledgeMatch,
    ArchitectureQuery,
)
from smb_requirement_agent.references.domain.architecture.catalogue import (
    InvalidArchitectureContentError,
    SystemCapability,
    SystemReference,
)
from smb_requirement_agent.requirements.domain.requirement.entities import Requirement
from smb_requirement_agent.requirements.domain.requirement.value_objects import (
    RequirementDescription,
    RequirementStatus,
    RequirementTitle,
)
from smb_requirement_agent.shared_kernel.identifiers import RequirementId
from tests.unit.breakdown.test_feature_domain import make_feature
from tests.unit.governance.test_backlog_export import _document

NOW = datetime(2026, 1, 1, 12, 0, tzinfo=UTC)


def _impact(capability: SystemCapability) -> ArchitectureImpact:
    return ArchitectureImpact("v1", NOW, (SystemReference("crm", "CRM", True, (capability,)),))


class _FixedKnowledge:
    """The knowledge service, answering every query with one fixed match."""

    def __init__(self, match: ArchitectureKnowledgeMatch) -> None:
        self.match_result = match

    def match(self, query: ArchitectureQuery) -> ArchitectureKnowledgeMatch:
        return self.match_result


def test_mapping_records_the_component_that_delivers_each_capability_as_given() -> None:
    placed = _impact(
        SystemCapability(
            "quote", "Quote", component_id="quote-engine", component_name="Quote engine"
        )
    )
    requirement = Requirement(
        id=RequirementId("req-1"),
        title=RequirementTitle("Quotes"),
        description=RequirementDescription("The assisted sales journey"),
        status=RequirementStatus.DRAFT,
    )

    mapped = MapFeatureArchitecture(
        _FixedKnowledge(ArchitectureKnowledgeMatch("v1", placed.systems, ()))
    ).execute(requirement, make_feature(), NOW)

    capability = mapped.impact.systems[0].capabilities[0]
    assert (capability.component_id, capability.component_name) == ("quote-engine", "Quote engine")
    assert mapped.feature.architecture == mapped.impact


def test_components_on_impacts_persist_and_stay_out_of_fingerprints_and_prompts() -> None:
    placed = _impact(SystemCapability("quote", "Quote", component_id="qe", component_name="QE"))
    bare = _impact(SystemCapability("quote", "Quote"))

    assert architecture_from_payload(architecture_to_payload(placed)) == placed
    legacy = architecture_to_payload(bare)
    assert isinstance(legacy, dict)
    assert legacy["systems"][0]["capabilities"] == [{"id": "quote", "name": "Quote"}]
    feature = make_feature()
    assert artifact_fingerprint(feature.with_architecture(placed)) == artifact_fingerprint(
        feature.with_architecture(bare)
    )
    assert render_guidance(
        GenerationGuidance(architecture=ArchitectureKnowledgeMatch("v1", placed.systems, ()))
    ) == render_guidance(
        GenerationGuidance(architecture=ArchitectureKnowledgeMatch("v1", bare.systems, ()))
    )
    with pytest.raises(InvalidArchitectureContentError, match="needs its name"):
        SystemCapability("quote", "Quote", component_id="qe")


def test_exports_carry_the_capability_component() -> None:
    document = _document("owner")
    feature = document.epic.features[0]
    assert feature.architecture is not None
    system = feature.architecture.systems[0]
    with_component = replace(
        system,
        capabilities=(replace(system.capabilities[0], component="Quote engine"),),
    )
    document = replace(
        document,
        epic=replace(
            document.epic,
            features=(
                replace(
                    feature,
                    architecture=replace(
                        feature.architecture,
                        systems=(with_component, *feature.architecture.systems[1:]),
                    ),
                ),
            ),
        ),
    )

    workbook = load_workbook(io.BytesIO(XlsxBacklogExporter().render(document)))
    header = [cell.value for cell in workbook["Systems"][1]]
    row = next(workbook["Systems"].iter_rows(min_row=2, values_only=True))
    capability_name = system.capabilities[0].name
    assert row[header.index("capability_components")] == f"{capability_name}: Quote engine"
    payload = json.loads(JsonBacklogExporter().render(document))
    capability = payload["epic"]["features"][0]["architecture"]["systems"][0]["capabilities"][0]
    assert capability["component"] == "Quote engine"
