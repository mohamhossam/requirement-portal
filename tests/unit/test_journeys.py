"""Journeys: the ordered activities that fulfil an order, and the flow derived from them (ADR-0096).

The reported gap: a New Activation journey's 18 activities, 7 flow rules and 18
activity integrations had nowhere to go, and its flow lived only in a diagram.
"""

from __future__ import annotations

import io
from dataclasses import replace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from smb_requirement_agent.application.ports.architecture_rag import EvidenceChunk
from smb_requirement_agent.application.ports.catalogue_file import CatalogueFileFormat
from smb_requirement_agent.application.ports.identity import Actor
from smb_requirement_agent.application.use_cases.architecture_index import BuildArchitectureIndex
from smb_requirement_agent.application.use_cases.architecture_knowledge import (
    ManageArchitectureKnowledge,
)
from smb_requirement_agent.domain.architecture.diff import ChangedItem, ChangeKind, diff_releases
from smb_requirement_agent.domain.architecture.journeys import (
    Activity,
    ActivityIntegration,
    FlowRule,
    FlowRuleKind,
    Journey,
    journey_edges,
)
from smb_requirement_agent.domain.architecture.knowledge import (
    ArchitectureKnowledge,
    InvalidKnowledgeError,
)
from smb_requirement_agent.domain.architecture.products import (
    OfferingComponent,
    OrderType,
    ProductOffering,
    SourceConfidence,
)
from smb_requirement_agent.infrastructure.architecture.catalogue_files import CatalogueFileAdapter
from smb_requirement_agent.infrastructure.architecture.embeddings import FakeEmbeddings
from smb_requirement_agent.infrastructure.architecture.evidence_index import InMemoryEvidenceIndex
from smb_requirement_agent.infrastructure.architecture.knowledge_yaml import seed_knowledge
from smb_requirement_agent.infrastructure.architecture.located_extractor import (
    LocatedDocumentExtractor,
)
from smb_requirement_agent.infrastructure.architecture.tokenizer import ApproximateTokenizer
from smb_requirement_agent.infrastructure.documents.text_extractor import SafeDocumentTextExtractor
from smb_requirement_agent.infrastructure.persistence.in_memory_architecture_knowledge import (
    InMemoryArchitectureKnowledgeRepository,
)
from smb_requirement_agent.infrastructure.persistence.in_memory_document_repository import (
    InMemoryDocumentStorage,
)

OWNER = {"X-Fake-Actor-Id": "fake-owner"}
MAINTAINER = Actor("amina", frozenset({"knowledge_maintainer"}))
ADAPTER = CatalogueFileAdapter()

OFFICE = ProductOffering(
    "office-connect",
    "Office Connect",
    order_types=(OrderType("NEW_ACTIVATION", "New Activation"),),
    components=(OfferingComponent("fibre", "Fibre Access"),),
)

# The reported journey's shape: a main track, a correction side track, four parallel tracks.
TRACKS = {
    "75": "CORRECTION",
    "120": "SERVICE",
    "130": "FIELD",
    "140": "BACKUP_5G",
    "150": "COMMERCIAL",
}
NUMBERS = ("10", "20", "30", "40", "50", "60", "70", "75", "80", "90", "100", "120", "130",
           "140", "150", "170", "180", "190")  # fmt: skip
RULES = (
    FlowRule(FlowRuleKind.DECISION, "70", "80", "PASS", "PASS"),
    FlowRule(FlowRuleKind.DECISION, "70", "75", "FAIL", "FAIL"),
    FlowRule(FlowRuleKind.LOOP, "75", "10", "Correct & Resubmit", "CORRECTION"),
    *(
        FlowRule(FlowRuleKind.PARALLEL, "100", number, None, TRACKS[number], "PG1", "170")
        for number in ("120", "130", "140", "150")
    ),
)
ACTIVATION = Journey(
    "new-activation",
    "New Activation",
    product_id="office-connect",
    order_type_code="NEW_ACTIVATION",
    # Listed out of order on purpose: the journey orders activities by number.
    activities=tuple(
        Activity(
            number,
            f"Activity {number}",
            track=TRACKS.get(number),
            performing_system_id="b2b-web" if number in {"10", "75"} else "bcrm",
            supporting_system_ids=("cim",) if number == "10" else (),
            system_function="Offer UI / basket" if number == "10" else None,
            component_ids=("fibre",) if number == "10" else (),
            confidence=SourceConfidence.INFERRED if number == "150" else None,
        )
        for number in reversed(NUMBERS)
    ),
    flow_rules=RULES,
    integrations=(
        ActivityIntegration(
            "10", "20", "INTERNAL_APP", "Digital Catalog", "Selected offer", "SYNC"
        ),
    ),
)


def _release(*journeys: Journey) -> ArchitectureKnowledge:
    return replace(seed_knowledge(), products=(OFFICE,), journeys=journeys)


# The flow ---------------------------------------------------------------------------------


def test_the_flow_is_derived_as_the_documents_diagram_draws_it() -> None:
    edges = {
        (edge.from_activity, edge.to_activity, edge.label) for edge in journey_edges(ACTIVATION)
    }

    assert edges == {
        ("10", "20", None),
        ("20", "30", None),
        ("30", "40", None),
        ("40", "50", None),
        ("50", "60", None),
        ("60", "70", None),
        ("80", "90", None),
        ("90", "100", None),
        ("170", "180", None),
        ("180", "190", None),
        ("70", "80", "PASS"),
        ("70", "75", "FAIL"),
        ("75", "10", "Correct & Resubmit"),
        ("100", "120", "SERVICE"),
        ("100", "130", "FIELD"),
        ("100", "140", "BACKUP_5G"),
        ("100", "150", "COMMERCIAL"),
        ("120", "170", "rejoin"),
        ("130", "170", "rejoin"),
        ("140", "170", "rejoin"),
        ("150", "170", "rejoin"),
    }
    assert [item.number for item in ACTIVATION.ordered][:9] == [
        "10", "20", "30", "40", "50", "60", "70", "75", "80",
    ]  # fmt: skip


@pytest.mark.parametrize(
    ("build", "message"),
    [
        (lambda: replace(ACTIVATION, activities=(*ACTIVATION.activities, Activity("10", "Again"))),
         "activity numbers must be unique"),
        (lambda: replace(ACTIVATION, flow_rules=(FlowRule("loop", "75", "5"),)),  # type: ignore[arg-type]
         "a loop rule names activity '5'"),
        (lambda: replace(ACTIVATION, integrations=(ActivityIntegration("10", "999"),)),
         "an integration names activity '999'"),
        (lambda: replace(ACTIVATION, product_id=None), "an order type needs its product offering"),
        (lambda: FlowRule("jump", "1", "2"), "decision, a loop or a parallel track"),  # type: ignore[arg-type]
    ],
)  # fmt: skip
def test_a_journey_is_refused_when_it_contradicts_itself(build: object, message: str) -> None:
    with pytest.raises(InvalidKnowledgeError, match=message):
        build()  # type: ignore[operator]


@pytest.mark.parametrize(
    ("change", "message"),
    [
        (lambda release: replace(release, systems=tuple(
            item for item in release.systems if item.id != "cim")),
         "Activity 10 names system 'cim'"),
        (lambda release: replace(release, products=()),
         "is for product offering 'office-connect', which is not"),
        (lambda release: replace(
            release, journeys=(replace(ACTIVATION, order_type_code="UPGRADE"),)),
         "order type 'UPGRADE' is not one of Office Connect's"),
        (lambda release: replace(release, products=(replace(OFFICE, components=()),)),
         "names component 'fibre', which is not part of its product offering"),
    ],
)  # fmt: skip
def test_what_a_journey_names_cannot_be_removed_from_under_it(change: object, message: str) -> None:
    with pytest.raises(InvalidKnowledgeError, match=message):
        change(_release(ACTIVATION))  # type: ignore[operator]


# Files ------------------------------------------------------------------------------------


@pytest.mark.parametrize("file_format", list(CatalogueFileFormat))
def test_every_format_keeps_journeys_whole(file_format: CatalogueFileFormat) -> None:
    content = ADAPTER.read(file_format, ADAPTER.write(file_format, _release(ACTIVATION)))

    (journey,) = content.journeys
    assert journey == ACTIVATION
    assert journey_edges(journey) == journey_edges(ACTIVATION)


def test_the_workbook_says_which_journey_row_is_wrong() -> None:
    workbook = load_workbook(
        io.BytesIO(ADAPTER.write(CatalogueFileFormat.XLSX, _release(ACTIVATION)))
    )
    rules = workbook["FlowRules"]
    headers = [cell.value for cell in rules[1]]
    rules.cell(2, headers.index("kind") + 1, "jump")
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(InvalidKnowledgeError, match="FlowRules row 2: A flow rule is a decision"):
        ADAPTER.read(CatalogueFileFormat.XLSX, buffer.getvalue())


# Diff and evidence ------------------------------------------------------------------------


def test_the_changes_view_names_what_changed_in_a_journey() -> None:
    base = _release(ACTIVATION)
    shorter = replace(ACTIVATION, flow_rules=ACTIVATION.flow_rules[:3], description="Fewer.")

    changed = diff_releases(base, replace(base, journeys=(shorter,))).changes

    assert [(item.item, item.change, item.fields) for item in changed] == [
        (ChangedItem.JOURNEY, ChangeKind.CHANGED, ("description", "flow_rules"))
    ]


class _Recording(InMemoryEvidenceIndex):
    def __init__(self) -> None:
        super().__init__(FakeEmbeddings(), ApproximateTokenizer())
        self.chunks: tuple[EvidenceChunk, ...] = ()

    def store(self, release_id: str, index_id: str, chunks: tuple[EvidenceChunk, ...]) -> None:
        self.chunks = chunks
        super().store(release_id, index_id, chunks)


def test_evidence_walks_the_journey_with_the_system_at_each_step() -> None:
    repository = InMemoryArchitectureKnowledgeRepository(seed_knowledge())
    index = _Recording()
    manage = ManageArchitectureKnowledge(repository, CatalogueFileAdapter(), index, 10_000_000)
    draft = manage.create_draft(MAINTAINER, "Next version")
    draft = manage.update(
        draft.id, draft.revision, MAINTAINER, products=(OFFICE,), journeys=(ACTIVATION,)
    )
    BuildArchitectureIndex(
        manage,
        index,
        InMemoryDocumentStorage(),
        LocatedDocumentExtractor(SafeDocumentTextExtractor()),
        ApproximateTokenizer(),
    ).execute(draft.id, draft.revision, "amina", fence=lambda: None)

    text = "\n".join(item.text for item in index.chunks if item.location.startswith("journey "))
    assert text.startswith("Journey: New Activation (Office Connect › New Activation)")
    assert "10. Activity 10 — B2B Web (with CIM): Offer UI / basket" in text
    assert "10 → 20: INTERNAL_APP, Digital Catalog, Selected offer" in text


# API --------------------------------------------------------------------------------------


def test_a_draft_saves_journeys_and_sends_their_flow(client: TestClient) -> None:
    draft = client.post(
        "/architecture-knowledge/releases", json={"name": "Next version"}, headers=OWNER
    ).json()
    url = f"/architecture-knowledge/releases/{draft['id']}"
    offering = {
        "id": "office-connect",
        "name": "Office Connect",
        "order_types": [{"code": "NEW_ACTIVATION", "name": "New Activation"}],
    }
    journey = {
        "id": "new-activation",
        "name": "New Activation",
        "product_id": "office-connect",
        "order_type_code": "NEW_ACTIVATION",
        "activities": [
            {"number": "10", "name": "Select offer", "performing_system_id": "b2b-web"},
            {"number": "20", "name": "Validate", "performing_system_id": "bcrm"},
            {"number": "25", "name": "Fix", "track": "CORRECTION", "performing_system_id": "cim"},
        ],
        "flow_rules": [
            {"kind": "decision", "from_activity": "20", "to_activity": "25", "condition": "FAIL"}
        ],
        # Sent edges are ignored: the flow is always derived.
        "edges": [{"from_activity": "10", "to_activity": "25", "kind": "sequence"}],
    }
    body = {"systems": draft["systems"], "relationships": draft["relationships"]}

    saved = client.put(
        url,
        json={
            **body,
            "expected_revision": draft["revision"],
            "products": [offering],
            "journeys": [journey],
        },
        headers=OWNER,
    )
    assert saved.status_code == 200, saved.text
    edges = saved.json()["journeys"][0]["edges"]
    assert [(item["from_activity"], item["to_activity"], item["label"]) for item in edges] == [
        ("10", "20", None),
        ("20", "25", "FAIL"),
    ]

    refused = client.put(
        url,
        json={**body, "expected_revision": saved.json()["revision"], "products": []},
        headers=OWNER,
    )
    assert refused.status_code == 422
    assert "New Activation is for product offering 'office-connect'" in refused.text
