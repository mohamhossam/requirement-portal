"""Product offerings and journeys around an impact, as requirement work receives them (ADR-0097).

Finding the offerings an item names and the journey steps of its mapped systems
is the knowledge service's job and is covered there. These tests cover what
requirement work does with the advice it is given: record, store and export it,
and keep it out of approval fingerprints and generation prompts.
"""

from __future__ import annotations

import io
import json
from dataclasses import replace
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from smb_requirement_agent.application.ports.architecture_knowledge import (
    ArchitectureKnowledgeMatch,
    ArchitectureQuery,
)
from smb_requirement_agent.application.ports.generation_guidance import GenerationGuidance
from smb_requirement_agent.application.use_cases.approval_policy import artifact_fingerprint
from smb_requirement_agent.application.use_cases.architecture_mapping import (
    MapFeatureArchitecture,
)
from smb_requirement_agent.application.use_cases.export_breakdown import (
    _journey_step,
    _product_context,
)
from smb_requirement_agent.domain.architecture.entities import (
    ArchitectureImpact,
    JourneyNeighbour,
    JourneyStep,
    OfferingDuty,
    ProductContext,
    SystemReference,
)
from smb_requirement_agent.domain.architecture.errors import InvalidArchitectureContentError
from smb_requirement_agent.domain.requirement.entities import Requirement
from smb_requirement_agent.domain.requirement.value_objects import (
    RequirementDescription,
    RequirementStatus,
    RequirementTitle,
)
from smb_requirement_agent.domain.shared.identifiers import RequirementId
from smb_requirement_agent.infrastructure.exports.json_exporter import JsonBacklogExporter
from smb_requirement_agent.infrastructure.exports.xlsx_exporter import XlsxBacklogExporter
from smb_requirement_agent.infrastructure.llm.prompts.generation_guidance import render_guidance
from smb_requirement_agent.infrastructure.persistence.shared_payloads import (
    architecture_from_payload,
    architecture_to_payload,
)
from tests.unit.test_backlog_export import _document
from tests.unit.test_feature_domain import make_feature

NOW = datetime(2026, 10, 1, tzinfo=UTC)

CONTEXT = ProductContext(
    "business-pro-plus",
    "Business Pro Plus",
    ("Business Pro Plus",),
    "New Activation",
    (OfferingDuty("static-ip", "Static IP", "rtf", "RTF", "VALIDATION", "Checks the IP."),),
)
STEP = JourneyStep(
    "crm",
    "bpp-new-activation",
    "New Activation",
    "10",
    "Select offer",
    True,
    ("Business Pro Plus", "New Activation"),
    (),
    (JourneyNeighbour("20", "Create order", "rtf", "RTF"), JourneyNeighbour("30", "Wait")),
)


def _impact() -> ArchitectureImpact:
    return ArchitectureImpact("v1", NOW, (SystemReference("crm", "CRM", True),))


class _FixedKnowledge:
    """The knowledge service, answering every query with one fixed match."""

    def __init__(self, match: ArchitectureKnowledgeMatch) -> None:
        self.match_result = match

    def match(self, query: ArchitectureQuery) -> ArchitectureKnowledgeMatch:
        return self.match_result


def test_mapping_records_the_product_context_and_journey_steps_as_given() -> None:
    knowledge = _FixedKnowledge(
        ArchitectureKnowledgeMatch(
            "v1",
            _impact().systems,
            (),
            product_contexts=(CONTEXT,),
            journey_steps=(STEP,),
        )
    )
    requirement = Requirement(
        id=RequirementId("req-1"),
        title=RequirementTitle("Business Pro Plus"),
        description=RequirementDescription("Business Pro Plus: the assisted sales journey"),
        status=RequirementStatus.DRAFT,
    )

    mapped = MapFeatureArchitecture(knowledge).execute(requirement, make_feature(), NOW)

    assert mapped.feature.architecture == mapped.impact
    (context,) = mapped.impact.product_contexts
    assert [(item.system_name, item.role) for item in context.responsibilities] == [
        ("RTF", "VALIDATION")
    ]
    (step,) = mapped.impact.journey_steps
    assert (step.system_id, step.number, [item.system_name for item in step.after]) == (
        "crm",
        "10",
        ["RTF", None],
    )


def test_advice_persists_and_stays_out_of_fingerprints_and_prompts() -> None:
    bare = _impact()
    advised = replace(bare, product_contexts=(CONTEXT,), journey_steps=(STEP,))

    assert architecture_from_payload(architecture_to_payload(advised)) == advised
    legacy = architecture_to_payload(bare)
    assert isinstance(legacy, dict)
    assert "product_contexts" not in legacy and "journey_steps" not in legacy
    feature = make_feature()
    assert artifact_fingerprint(feature.with_architecture(advised)) == artifact_fingerprint(
        feature.with_architecture(bare)
    )
    match = ArchitectureKnowledgeMatch(
        "v1", bare.systems, (), product_contexts=(CONTEXT,), journey_steps=(STEP,)
    )
    plain = ArchitectureKnowledgeMatch("v1", bare.systems, ())
    assert render_guidance(GenerationGuidance(architecture=match)) == render_guidance(
        GenerationGuidance(architecture=plain)
    )
    with pytest.raises(InvalidArchitectureContentError, match="mapped systems"):
        replace(bare, journey_steps=(replace(STEP, system_id="other"),))
    with pytest.raises(InvalidArchitectureContentError, match="words that matched"):
        replace(CONTEXT, matched_terms=())


def test_exports_carry_product_context_and_journey_steps() -> None:
    document = _document("owner")
    feature = document.epic.features[0]
    assert feature.architecture is not None
    mapped = feature.architecture.systems[0].id
    architecture = replace(
        feature.architecture,
        product_contexts=(_product_context(CONTEXT),),
        journey_steps=(_journey_step(replace(STEP, system_id=mapped)),),
    )
    document = replace(
        document,
        epic=replace(document.epic, features=(replace(feature, architecture=architecture),)),
    )

    workbook = load_workbook(io.BytesIO(XlsxBacklogExporter().render(document)))
    header = [cell.value for cell in workbook["Product Context"][1]]
    (row,) = workbook["Product Context"].iter_rows(min_row=2, values_only=True)
    assert row[header.index("product_name")] == "Business Pro Plus"
    assert row[header.index("system_name")] == "RTF"
    assert row[header.index("order_type")] == "New Activation"
    header = [cell.value for cell in workbook["Journey Steps"][1]]
    (row,) = workbook["Journey Steps"].iter_rows(min_row=2, values_only=True)
    assert row[header.index("fulfils")] == "Business Pro Plus > New Activation"
    assert row[header.index("part")] == "performs"
    assert row[header.index("after")] == "20. Create order (RTF); 30. Wait"
    payload = json.loads(JsonBacklogExporter().render(document))
    exported = payload["epic"]["features"][0]["architecture"]
    assert exported["product_contexts"][0]["responsibilities"][0]["role"] == "VALIDATION"
    assert exported["journey_steps"][0]["after"][1] == {
        "number": "30",
        "name": "Wait",
        "system_id": None,
        "system_name": None,
    }
