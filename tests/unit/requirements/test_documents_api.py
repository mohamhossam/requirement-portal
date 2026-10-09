"""Slice 5D source-document application and HTTP behavior."""

from __future__ import annotations

import io
import zipfile
from dataclasses import replace
from typing import Any, cast

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image

from smb_requirement_agent.application.errors import (
    UnsupportedDocumentError,
)
from smb_requirement_agent.identity.infrastructure.fake_identity import FAKE_ACTORS
from smb_requirement_agent.infrastructure.config.settings import Settings
from smb_requirement_agent.interfaces.api.container import Container, build_container
from smb_requirement_agent.requirements.application.use_cases.create_requirement import (
    CreateRequirementInput,
)
from smb_requirement_agent.requirements.application.use_cases.documents import UploadDocumentInput
from smb_requirement_agent.requirements.domain.document.value_objects import DocumentId
from smb_requirement_agent.shared_kernel.identifiers import RequirementId
from tests.conftest import FAKE_PROVIDER_SETTINGS
from tests.unit.workflow_helpers import post_analysis
from tests.word_table_fixtures import DOCX_MIME, reviewed_word_prose_document


def _create_requirement(client: TestClient) -> str:
    response = client.post(
        "/requirements",
        json={"title": "Bundle ordering", "description": "Order a bundle"},
    )
    assert response.status_code == 201
    return str(response.json()["id"])


def _upload_txt(
    client: TestClient, requirement_id: str, text: bytes = b"Policy A"
) -> dict[str, Any]:
    response = client.post(
        f"/requirements/{requirement_id}/attachments",
        files={"file": ("policy.txt", text, "text/plain")},
    )
    assert response.status_code == 201
    return cast(dict[str, Any], response.json())


def _docx_with_raster_image() -> bytes:
    image_stream = io.BytesIO()
    Image.new("RGB", (240, 120), color=(20, 80, 140)).save(image_stream, format="PNG")
    document_xml = (
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        "<w:body><w:p><w:r><w:t>Workflow screenshot</w:t><w:drawing>"
        '<a:blip r:embed="rId1"/></w:drawing></w:r></w:p></w:body></w:document>'
    )
    relationships = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships/image" Target="media/image1.png"/>'
        "</Relationships>"
    )
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/_rels/document.xml.rels", relationships)
        archive.writestr("word/media/image1.png", image_stream.getvalue())
    return stream.getvalue()


def test_document_journey_records_exact_analysis_version_and_checksum(
    client: TestClient, container: Container
) -> None:
    requirement_id = _create_requirement(client)
    uploaded = _upload_txt(client, requirement_id)
    document_id = uploaded["id"]
    version = uploaded["current_version"]

    included = client.put(
        f"/requirements/{requirement_id}/attachments/{document_id}/analysis-inclusion",
        json={"included": True, "expected_version": uploaded["version"]},
    )
    assert included.status_code == 200
    analysis = post_analysis(client, requirement_id)
    assert analysis.status_code == 200
    assert analysis.json()["document_references"] == [
        {
            "document_id": document_id,
            "version_id": version["id"],
            "filename": "policy.txt",
            "checksum_sha256": version["checksum_sha256"],
        }
    ]
    content = client.get(f"/documents/{document_id}/content")
    assert content.status_code == 200
    assert content.json()["extracted_text"] == "Policy A"
    assert any(item["id"] == document_id for item in client.get("/documents").json())


def test_xlsx_upload_exposes_structured_evidence_and_hidden_sheet_selection(
    client: TestClient,
) -> None:
    workbook = Workbook()
    visible = workbook.active
    assert visible is not None
    visible.title = "Workflow"
    visible["A1"] = "Step"
    visible["B2"] = "Validate customer"
    visible.merge_cells("A1:A2")
    hidden = workbook.create_sheet("Internal")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Unconfirmed rule"
    stream = io.BytesIO()
    workbook.save(stream)
    requirement_id = _create_requirement(client)

    response = client.post(
        f"/requirements/{requirement_id}/attachments",
        files={
            "file": (
                "workflow.xlsx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 201
    document = response.json()
    version = document["current_version"]
    assert version["analysis_readiness"] == "ready_with_warnings"
    assert version["evidence_summary"]["worksheet_count"] == 2
    assert version["evidence_summary"]["block_count"] >= 4
    assert document["hidden_worksheets"] == ["Worksheet 2"]
    extracted = document["versions"][0]
    assert extracted["extraction_version"] == "structured-xlsx-sections-v3"
    row = next(b for b in extracted["evidence_blocks"] if b["label"] == "Worksheet 1!2:2")
    assert row["text"] == "A2=[merged A1:A2 continuation; anchor A1] | B2=Validate customer"
    assert "Step" not in row["text"]
    selected = client.put(
        f"/requirements/{requirement_id}/attachments/{document['id']}/hidden-worksheets",
        json={"worksheet_names": ["Worksheet 2"], "expected_version": document["version"]},
    )
    assert selected.status_code == 200
    assert selected.json()["included_hidden_worksheets"] == ["Worksheet 2"]


def test_document_asset_endpoint_returns_only_sanitized_raster_content(
    client: TestClient,
) -> None:
    requirement_id = _create_requirement(client)
    response = client.post(
        f"/requirements/{requirement_id}/attachments",
        files={
            "file": (
                "workflow.docx",
                _docx_with_raster_image(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201
    document = response.json()
    version = document["versions"][0]
    image_block = next(item for item in version["evidence_blocks"] if item["kind"] == "image")

    asset = client.get(
        f"/documents/{document['id']}/versions/{version['id']}/assets/{image_block['asset_id']}"
    )

    assert asset.status_code == 200
    assert asset.headers["content-type"] == "image/jpeg"
    assert asset.headers["x-content-type-options"] == "nosniff"
    with Image.open(io.BytesIO(asset.content)) as image:
        assert image.size == (240, 120)


def test_new_version_detaches_inclusion_and_invalidates_current_analysis(
    client: TestClient, container: Container
) -> None:
    requirement_id = _create_requirement(client)
    uploaded = _upload_txt(client, requirement_id)
    document_id = uploaded["id"]
    client.put(
        f"/requirements/{requirement_id}/attachments/{document_id}/analysis-inclusion",
        json={"included": True, "expected_version": uploaded["version"]},
    )
    assert post_analysis(client, requirement_id).status_code == 200

    current = client.get(f"/documents/{document_id}").json()

    changed = client.post(
        f"/requirements/{requirement_id}/attachments/{document_id}/versions",
        data={"expected_version": current["version"]},
        files={"file": ("policy.txt", b"Policy B", "text/plain")},
    )

    assert changed.status_code == 201
    assert changed.json()["version_count"] == 2
    assert changed.json()["included_in_analysis"] is False
    assert client.get(f"/requirements/{requirement_id}/analysis").status_code == 404
    history = container.breakdown_repository.list_breakdown_revisions(RequirementId(requirement_id))
    assert any(
        revision.analysis is not None
        and revision.analysis.document_references[0].version_id == uploaded["current_version"]["id"]
        for revision in history
        if revision.analysis is not None and revision.analysis.document_references
    )


def test_remove_detaches_document_and_preserves_metadata_as_hidden_history(
    client: TestClient, container: Container
) -> None:
    requirement_id = _create_requirement(client)
    uploaded = _upload_txt(client, requirement_id)
    document_id = uploaded["id"]
    client.put(
        f"/requirements/{requirement_id}/attachments/{document_id}/analysis-inclusion",
        json={"included": True, "expected_version": uploaded["version"]},
    )

    current = client.get(f"/documents/{document_id}").json()
    removed = client.delete(
        f"/requirements/{requirement_id}/attachments/{document_id}",
        params={"expected_version": current["version"]},
    )

    assert removed.status_code == 204
    assert client.get(f"/documents/{document_id}").status_code == 404
    persisted = container.document_repository.get(DocumentId(document_id))
    assert persisted is not None
    assert persisted.removed is True
    assert persisted.versions[0].checksum_sha256 == uploaded["current_version"]["checksum_sha256"]


@pytest.mark.parametrize(
    ("filename", "content", "mime_type"),
    [
        ("../policy.txt", b"text", "text/plain"),
        ("policy.pdf", b"text", "text/plain"),
        ("policy.pdf", b"not-pdf", "application/pdf"),
        ("policy.exe", b"text", "application/octet-stream"),
    ],
)
def test_unsafe_or_unsupported_uploads_are_rejected(
    container: Container, filename: str, content: bytes, mime_type: str
) -> None:
    requirement = container.create_requirement.execute(
        CreateRequirementInput("Title", "Description"), FAKE_ACTORS[0]
    )

    with pytest.raises(UnsupportedDocumentError):
        container.upload_document.for_requirement(
            requirement.id, UploadDocumentInput(filename, mime_type, content), FAKE_ACTORS[0]
        )


def test_extraction_failure_is_retained_for_review_but_cannot_be_included(
    client: TestClient,
) -> None:
    requirement_id = _create_requirement(client)
    response = client.post(
        f"/requirements/{requirement_id}/attachments",
        files={"file": ("broken.pdf", b"%PDF-corrupt", "application/pdf")},
    )

    assert response.status_code == 201
    data = response.json()
    assert data["current_version"]["extraction_status"] == "failed"
    assert data["current_version"]["extraction_error"]
    include = client.put(
        f"/requirements/{requirement_id}/attachments/{data['id']}/analysis-inclusion",
        json={"included": True, "expected_version": data["version"]},
    )
    assert include.status_code == 409


def test_upload_size_limit_and_structured_context_is_packetized() -> None:
    settings: Settings = replace(
        FAKE_PROVIDER_SETTINGS,
        document_max_file_bytes=5,
        document_context_max_characters=4,
    )
    container = build_container(settings)
    requirement = container.create_requirement.execute(
        CreateRequirementInput("Title", "Description"), FAKE_ACTORS[0]
    )
    with pytest.raises(UnsupportedDocumentError, match="upload limit"):
        container.upload_document.for_requirement(
            requirement.id,
            UploadDocumentInput("large.txt", "text/plain", b"123456"),
            FAKE_ACTORS[0],
        )
    document = container.upload_document.for_requirement(
        requirement.id,
        UploadDocumentInput("limit.txt", "text/plain", b"12345"),
        FAKE_ACTORS[0],
    )
    container.set_document_inclusion.execute(
        FAKE_ACTORS[0], requirement.id, document.id, True, document.version_number
    )

    assert (
        container.analyze_requirement.execute(FAKE_ACTORS[0], requirement.id).requirement_id
        == requirement.id
    )


def test_draft_attachment_moves_to_promoted_requirement(client: TestClient) -> None:
    draft = client.post(
        "/requirements/drafts",
        json={
            "title": "Draft title",
            "description": "Draft description",
            "desired_outcome": "Draft outcome",
        },
    ).json()
    uploaded = client.post(
        f"/requirement-drafts/{draft['id']}/attachments",
        files={"file": ("source.txt", b"Draft source", "text/plain")},
    )
    assert uploaded.status_code == 201

    promoted = client.post(
        f"/requirements/drafts/{draft['id']}/promote",
        json={"expected_version": draft["version"]},
    )

    assert promoted.status_code == 201
    requirement_id = promoted.json()["id"]
    attachments = client.get(f"/requirements/{requirement_id}/attachments").json()
    assert [item["id"] for item in attachments] == [uploaded.json()["id"]]
    assert attachments[0]["draft_id"] is None


def test_txt_attachment_keeps_heading_wording_out_of_descendant_metadata(
    client: TestClient,
) -> None:
    requirement_id = _create_requirement(client)
    uploaded = _upload_txt(client, requirement_id, b"# Private heading\nXGPON coverage")
    version = uploaded["versions"][0]
    assert version["extraction_version"] == "structured-text-sections-v1"
    assert version["evidence_blocks"][1]["section_path"] == ["Heading at line 1"]
    assert version["evidence_blocks"][1]["text"] == "XGPON coverage"
    assert "Private" not in str(version["evidence_blocks"][1])


def test_word_attachment_prose_has_neutral_review_locations(client: TestClient) -> None:
    requirement_id = _create_requirement(client)
    response = client.post(
        f"/requirements/{requirement_id}/attachments",
        files={"file": ("policy.docx", reviewed_word_prose_document(), DOCX_MIME)},
    )
    assert response.status_code == 201
    version = response.json()["versions"][0]
    assert version["extraction_version"] == "structured-docx-sections-v3"
    assert version["evidence_blocks"][1]["label"] == "Paragraph 2"
    assert version["evidence_blocks"][1]["section_path"] == ["Heading at paragraph 1"]
    assert version["evidence_blocks"][1]["text"] == "Private paragraph XGPON التغطية مطلوبة."
