"""Excel, YAML and JSON catalogue files, and the changes view reviewed before publishing."""

from __future__ import annotations

import io
from dataclasses import replace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from smb_requirement_agent.application.ports.catalogue_file import CatalogueFileFormat
from smb_requirement_agent.domain.architecture.diff import ChangedItem, ChangeKind, diff_releases
from smb_requirement_agent.domain.architecture.knowledge import (
    InvalidKnowledgeError,
    KnowledgeCapability,
    SystemDefinition,
    SystemRelationship,
)
from smb_requirement_agent.infrastructure.architecture.catalogue_files import CatalogueFileAdapter
from smb_requirement_agent.infrastructure.architecture.knowledge_yaml import seed_knowledge

ADAPTER = CatalogueFileAdapter()
OWNER = {"X-Fake-Actor-Id": "fake-owner"}


@pytest.mark.parametrize("file_format", list(CatalogueFileFormat))
def test_every_format_round_trips_the_catalogue(file_format: CatalogueFileFormat) -> None:
    seed = replace(
        seed_knowledge(),
        systems=(
            replace(seed_knowledge().systems[0], constraints=("Read-only on weekends",)),
            *seed_knowledge().systems[1:],
        ),
    )

    content = ADAPTER.read(file_format, ADAPTER.write(file_format, seed))

    assert content.systems == seed.systems
    assert content.relationships == seed.relationships


def _workbook(**sheets: list[tuple[object, ...]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_excel_reads_separated_lists_and_ignores_blank_rows() -> None:
    content = ADAPTER.read(
        CatalogueFileFormat.XLSX,
        _workbook(
            Systems=[
                ("system_id", "name", "name_ar", "aliases"),
                ("crm", "CRM", None, "customer records; accounts"),
                (None, None, None, None),
                ("billing", "Billing", "الفوترة", None),
            ],
            Capabilities=[
                ("system_id", "capability_id", "name", "triggers"),
                ("crm", "lookup", "Customer lookup", "find customer;search account"),
            ],
            Relationships=[
                ("source_system_id", "target_system_id", "description"),
                ("crm", "billing", "Sends invoices"),
            ],
        ),
    )

    assert content.systems == (
        SystemDefinition(
            "crm",
            "CRM",
            ("customer records", "accounts"),
            (
                KnowledgeCapability(
                    "lookup", "Customer lookup", ("find customer", "search account")
                ),
            ),
        ),
        SystemDefinition("billing", "Billing", name_ar="الفوترة"),
    )
    assert content.relationships == (SystemRelationship("crm", "billing", "Sends invoices"),)


@pytest.mark.parametrize(
    ("sheets", "message"),
    [
        ({"Other": [("x",)]}, "needs a 'Systems' sheet"),
        ({"Systems": [("id", "title")]}, "Systems row 1 must have the headers"),
        (
            {"Systems": [("system_id", "name"), ("crm", None)]},
            "Systems row 2: name is required",
        ),
        (
            {
                "Systems": [("system_id", "name"), ("crm", "CRM")],
                "Capabilities": [
                    ("system_id", "capability_id", "name", "triggers"),
                    ("erp", "post", "Post", "post"),
                ],
            },
            "Capabilities row 2: system 'erp' is not listed",
        ),
        (
            {
                "Systems": [("system_id", "name"), ("crm", "CRM")],
                "Capabilities": [
                    ("system_id", "capability_id", "name", "triggers"),
                    ("crm", "lookup", "Lookup", None),
                ],
            },
            "Capabilities row 2: A capability needs non-empty matching phrases",
        ),
    ],
)
def test_excel_errors_name_the_sheet_and_row(
    sheets: dict[str, list[tuple[object, ...]]], message: str
) -> None:
    with pytest.raises(InvalidKnowledgeError, match=message):
        ADAPTER.read(CatalogueFileFormat.XLSX, _workbook(**sheets))


def test_text_formats_report_syntax_and_shape_errors() -> None:
    with pytest.raises(InvalidKnowledgeError, match="JSON file could not be parsed: line 1"):
        ADAPTER.read(CatalogueFileFormat.JSON, b"{nope")
    with pytest.raises(InvalidKnowledgeError, match="systems entry 1: name is required"):
        ADAPTER.read(CatalogueFileFormat.YAML, b"systems:\n  - id: crm\n")
    with pytest.raises(InvalidKnowledgeError, match="not an .xlsx"):
        ADAPTER.read(CatalogueFileFormat.XLSX, b"plain text")
    with pytest.raises(InvalidKnowledgeError, match=r"\.xlsx"):
        CatalogueFileFormat.from_filename("catalogue.csv")


def test_template_has_instructions_and_every_sheet() -> None:
    workbook = load_workbook(io.BytesIO(ADAPTER.template()))

    assert workbook.sheetnames == [
        "Instructions",
        "Systems",
        "Components",
        "Capabilities",
        "Constraints",
        "Relationships",
        "Domains",
        "LandscapeDomains",
        "Products",
        "OrderTypes",
        "OfferingComponents",
        "Responsibilities",
        "ProductPoints",
        "Journeys",
        "Activities",
        "FlowRules",
        "ActivityIntegrations",
    ]
    assert ADAPTER.read(CatalogueFileFormat.XLSX, ADAPTER.template()).systems == ()


def test_diff_lists_added_changed_and_removed_items() -> None:
    base = seed_knowledge()
    first, second, *rest = base.systems
    changed_first = replace(
        first,
        name="Renamed",
        capabilities=(replace(first.capabilities[0], triggers=("new phrase",)),),
    )
    draft = replace(
        base,
        id="draft",
        systems=(changed_first, *rest, SystemDefinition("new", "New system")),
        relationships=base.relationships[1:],
    )

    diff = diff_releases(base, draft)

    found = {(item.item, item.change, item.key): item for item in diff.changes}
    assert found[(ChangedItem.SYSTEM, ChangeKind.CHANGED, first.id)].fields == ("name",)
    assert (ChangedItem.SYSTEM, ChangeKind.ADDED, "new") in found
    assert (ChangedItem.SYSTEM, ChangeKind.REMOVED, second.id) in found
    capability_key = f"{first.id}/{first.capabilities[0].id}"
    assert found[(ChangedItem.CAPABILITY, ChangeKind.CHANGED, capability_key)].fields == (
        "triggers",
    )
    assert sum(item.item is ChangedItem.RELATIONSHIP for item in diff.changes) >= 1
    assert diff_releases(base, base).empty


def test_changes_and_template_endpoints_are_for_maintainers(client: TestClient) -> None:
    draft = client.post(
        "/architecture-knowledge/releases", json={"name": "Next version"}, headers=OWNER
    ).json()
    workbook = _workbook(
        Systems=[("system_id", "name"), ("crm", "CRM")],
    )

    preview = client.post(
        f"/architecture-knowledge/releases/{draft['id']}/catalogue-file/preview",
        files={"file": ("catalogue.xlsx", workbook, "application/octet-stream")},
        headers=OWNER,
    )
    applied = client.post(
        f"/architecture-knowledge/releases/{draft['id']}/catalogue-file",
        data={"expected_revision": draft["revision"]},
        files={"file": ("catalogue.xlsx", workbook, "application/octet-stream")},
        headers=OWNER,
    )
    changes = client.get(
        f"/architecture-knowledge/releases/{draft['id']}/changes", headers=OWNER
    ).json()

    assert preview.status_code == 200
    assert {"item": "system", "change": "added", "key": "crm", "label": "CRM", "fields": []} in (
        preview.json()["changes"]
    )
    assert applied.json()["systems"][0]["id"] == "crm"
    assert changes["changes"] == preview.json()["changes"]
    template = client.get("/architecture-knowledge/catalogue-template.xlsx", headers=OWNER)
    assert template.status_code == 200
    assert 'filename="catalogue-template.xlsx"' in template.headers["content-disposition"]
    reader = {"X-Fake-Actor-Id": "fake-reviewer"}
    assert (
        client.get("/architecture-knowledge/catalogue-template.xlsx", headers=reader).status_code
        == 403
    )
