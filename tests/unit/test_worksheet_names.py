"""Worksheet names occur in reviewable text, never copied location metadata."""

import hashlib
import io
from dataclasses import replace
from datetime import UTC, datetime
from typing import cast

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image
from smb_kernel.documents.text_extractor import SafeDocumentTextExtractor

from smb_requirement_agent.domain.document.value_objects import EvidenceBlockKind
from tests.spreadsheet_fixtures import XLSX_MIME, reviewed_worksheet_names_document


def test_sheet_names_are_independent_headings_and_positions_include_empty_hidden_tabs() -> None:
    extractor = SafeDocumentTextExtractor()
    content = reviewed_worksheet_names_document()
    result = extractor.extract_structured(XLSX_MIME, content)
    assert result == extractor.extract_structured(XLSX_MIME, content)
    assert result.extraction_version == "structured-xlsx-sections-v3"
    assert [b.label for b in result.evidence_blocks] == [
        "Worksheet 1",
        "Worksheet 1!1:1",
        "Worksheet 2",
        "Hidden worksheet: Worksheet 3",
        "Worksheet 3!1:1",
    ]
    assert [b.text for b in result.evidence_blocks if b.kind is EvidenceBlockKind.HEADING] == [
        "Private heading سري",
        "Private empty",
        "Private hidden",
    ]
    assert all("Private" not in b.label + "/".join(b.section_path) for b in result.evidence_blocks)
    assert all("Private" not in w.message for w in result.warnings)
    assert result.evidence_blocks[-1].section_path == ("Hidden worksheet: Worksheet 3",)
    assert any(w.code == "worksheet_name_structure" for w in result.warnings)
    assert any("veryHidden" in w.message for w in result.warnings)


def test_chart_and_image_paths_share_neutral_sheet_positions_and_source_text_is_retained() -> None:
    workbook = Workbook()
    assert workbook.active is not None
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "Private visible"
    sheet["A1"] = 1
    sheet["A2"] = "='Private visible'!A1"
    chart = BarChart()
    chart.title = "Private chart title"
    chart.add_data(Reference(sheet, min_col=1, min_row=1, max_row=2))
    sheet.add_chart(chart, "D2")
    hidden = workbook.create_sheet("Private hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Hidden evidence"
    for current in (sheet, hidden):
        stream = io.BytesIO()
        Image.new("RGB", (240, 120), (40, 90, 140)).save(stream, format="PNG")
        stream.seek(0)
        current.add_image(SpreadsheetImage(stream), "F2")
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    result = SafeDocumentTextExtractor().extract_structured(XLSX_MIME, output.getvalue())
    images = [b for b in result.evidence_blocks if b.kind is EvidenceBlockKind.IMAGE]
    assert {b.section_path for b in images} == {
        ("Worksheet 1",),
        ("Hidden worksheet: Worksheet 2",),
    }
    assert all("Private" not in b.label + "/".join(b.section_path) for b in result.evidence_blocks)
    chart_block = next(b for b in result.evidence_blocks if b.label == "Chart 1")
    assert chart_block.section_path == ("Worksheet 1",)
    assert "Private chart title" in (chart_block.text or "")
    assert "'Private visible'!" in (chart_block.text or "")
    assert "='Private visible'!A1 (cached value unavailable)" in result.text


def test_chart_tabs_do_not_shift_following_worksheet_media_context() -> None:
    workbook = Workbook()
    assert workbook.active is not None
    sheet = cast(Worksheet, workbook.active)
    sheet["A1"] = 1
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=1, min_row=1, max_row=1))
    chart_tab = workbook.create_chartsheet("Private chart")
    chart_tab.add_chart(chart)
    tail = workbook.create_sheet("Private tail")
    tail["A1"] = "Final row"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    result = SafeDocumentTextExtractor().extract_structured(XLSX_MIME, output.getvalue())
    final = next(b for b in result.evidence_blocks if b.text == "A1=Final row")
    assert final.label == "Worksheet 3!1:1"
    assert final.section_path == ("Worksheet 3",)
    assert next(b for b in result.evidence_blocks if b.label == "Chart 1").section_path == (
        "Worksheet 2",
    )


@pytest.mark.parametrize("legacy", [False, True])
def test_hidden_selection_keeps_new_positions_and_legacy_names_separate(legacy: bool) -> None:
    from smb_requirement_agent.application.use_cases.documents import AssembleAnalysisDocuments
    from smb_requirement_agent.domain.document.entities import SourceDocument, SourceDocumentVersion
    from smb_requirement_agent.domain.document.value_objects import (
        DocumentId,
        DocumentVersionId,
        ExtractionStatus,
    )
    from smb_requirement_agent.infrastructure.persistence.in_memory_document_repository import (
        InMemoryDocumentRepository,
        InMemoryDocumentStorage,
    )
    from smb_requirement_agent.shared_kernel.identifiers import RequirementId

    content = reviewed_worksheet_names_document()
    extractor = SafeDocumentTextExtractor()
    extracted = extractor.extract_structured(XLSX_MIME, content)
    blocks = extracted.evidence_blocks
    identifier = "Private hidden" if legacy else "Worksheet 3"
    if legacy:
        blocks = tuple(
            replace(b, section_path=("Hidden worksheet: Private hidden",))
            if b.section_path == ("Hidden worksheet: Worksheet 3",)
            else b
            for b in blocks
        )
    version = SourceDocumentVersion(
        id=DocumentVersionId("v1"),
        number=1,
        filename="policy.xlsx",
        mime_type=XLSX_MIME,
        size_bytes=len(content),
        checksum_sha256=hashlib.sha256(content).hexdigest(),
        extraction_status=ExtractionStatus.READY,
        created_at=datetime(2026, 9, 22, tzinfo=UTC),
        extracted_text=extracted.text,
        extraction_version="structured-xlsx-merges-v1" if legacy else extracted.extraction_version,
        evidence_blocks=blocks,
    )
    requirement_id = RequirementId("requirement")
    document = SourceDocument(
        id=DocumentId("document"),
        versions=(version,),
        requirement_id=requirement_id,
    ).set_included(True)
    repository = InMemoryDocumentRepository()
    repository.add(document)
    assembly = AssembleAnalysisDocuments(repository, InMemoryDocumentStorage(), extractor, 10000)
    assert "Never publish appendix" not in str(assembly.execute(requirement_id))
    selected = document.select_hidden_worksheets((identifier,))
    repository.save(selected)
    assert selected.included_hidden_worksheets == (identifier,)
    assert "Never publish appendix" in str(assembly.execute(requirement_id))
    repository.save(selected.select_hidden_worksheets(()))
    assert "Never publish appendix" not in str(assembly.execute(requirement_id))
    assert selected.versions == document.versions
