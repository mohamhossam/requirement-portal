"""Capability domains on a mapped impact, as requirement work receives them (ADR-0089).

The domain tree, placing capabilities in it and suggesting domains when nothing
maps are the knowledge service's job and are covered there. These tests cover
what requirement work does with the domains it is given: record, store and
export them, and keep them out of approval fingerprints and generation prompts.
"""

from __future__ import annotations

import io
import json
from dataclasses import replace
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from smb_requirement_agent.application.ports.architecture_knowledge import (
    ArchitectureKnowledgeMatch,
    ArchitectureQuery,
)
from smb_requirement_agent.application.ports.generation_guidance import GenerationGuidance
from smb_requirement_agent.application.use_cases.approval_policy import artifact_fingerprint
from smb_requirement_agent.application.use_cases.architecture_mapping import (
    MapFeatureArchitecture,
)
from smb_requirement_agent.domain.architecture.entities import (
    ArchitectureImpact,
    DomainSuggestion,
    SystemCapability,
    SystemReference,
)
from smb_requirement_agent.domain.architecture.errors import InvalidArchitectureContentError
from smb_requirement_agent.domain.requirement.entities import Requirement
from smb_requirement_agent.domain.requirement.value_objects import (
    RequirementDescription,
    RequirementStatus,
    RequirementTitle,
)
from smb_requirement_agent.domain.shared.identifiers import RequirementId
from smb_requirement_agent.infrastructure.exports.json_exporter import JsonBacklogExporter
from smb_requirement_agent.infrastructure.exports.xlsx_exporter import XlsxBacklogExporter
from smb_requirement_agent.infrastructure.llm.prompts.generation_guidance import render_guidance
from smb_requirement_agent.infrastructure.persistence.shared_payloads import (
    architecture_from_payload,
    architecture_to_payload,
)
from tests.unit.test_backlog_export import _document
from tests.unit.test_feature_domain import make_feature

NOW = datetime(2026, 1, 1, 12, 0, tzinfo=UTC)


class _FixedKnowledge:
    """The knowledge service, answering every query with one fixed match."""

    def __init__(self, match: ArchitectureKnowledgeMatch) -> None:
        self.match_result = match

    def match(self, query: ArchitectureQuery) -> ArchitectureKnowledgeMatch:
        return self.match_result


def _requirement() -> Requirement:
    return Requirement(
        id=RequirementId("req-1"),
        title=RequirementTitle("Billing"),
        description=RequirementDescription("Speed up billing for small business orders"),
        status=RequirementStatus.DRAFT,
    )


def test_mapping_records_capability_domains_and_suggested_domains_as_given() -> None:
    placed = SystemReference(
        "bcrm", "BCRM", True, (SystemCapability("sales", "Sales", "orders", ("Order capture",)),)
    )
    suggestion = DomainSuggestion("billing", ("Billing",), ("bscs",), ("billing",), ("BSCS",))

    mapped = MapFeatureArchitecture(
        _FixedKnowledge(ArchitectureKnowledgeMatch("v1", (placed,), ()))
    ).execute(_requirement(), make_feature(), NOW)
    unmapped = MapFeatureArchitecture(
        _FixedKnowledge(ArchitectureKnowledgeMatch("v1", (), (), suggested_domains=(suggestion,)))
    ).execute(_requirement(), make_feature(), NOW)

    assert mapped.impact.systems[0].capabilities[0].domain_path == ("Order capture",)
    assert mapped.impact.suggested_domains == ()
    assert unmapped.impact.systems == ()
    assert [item.path for item in unmapped.impact.suggested_domains] == [("Billing",)]
    assert unmapped.impact.suggested_domains[0].system_ids == ("bscs",)
    assert unmapped.impact.suggested_domains[0].system_names == ("BSCS",)
    assert unmapped.feature.architecture == unmapped.impact


def _impact(**changes: object) -> ArchitectureImpact:
    base = ArchitectureImpact(
        "v1",
        NOW,
        (
            SystemReference(
                "crm",
                "CRM",
                True,
                (SystemCapability("quote", "Quote", "quotes", ("Order capture", "Quoting")),),
            ),
        ),
    )
    return replace(base, **changes)  # type: ignore[arg-type]


def test_domains_on_impacts_persist_and_stay_out_of_fingerprints_and_prompts() -> None:
    impact = _impact()
    bare = replace(
        impact,
        systems=(replace(impact.systems[0], capabilities=(SystemCapability("quote", "Quote"),)),),
    )
    suggested = ArchitectureImpact(
        "v1",
        NOW,
        (),
        suggested_domains=(DomainSuggestion("billing", ("Billing",), ("bscs",), ("billing",)),),
    )

    assert architecture_from_payload(architecture_to_payload(impact)) == impact
    assert architecture_from_payload(architecture_to_payload(suggested)) == suggested
    legacy = architecture_to_payload(bare)
    assert isinstance(legacy, dict)
    assert legacy["systems"][0]["capabilities"] == [{"id": "quote", "name": "Quote"}]
    assert "suggested_domains" not in legacy

    feature = make_feature()
    assert artifact_fingerprint(feature.with_architecture(impact)) == artifact_fingerprint(
        feature.with_architecture(bare)
    )
    match = ArchitectureKnowledgeMatch(
        "v1", impact.systems, (), suggested_domains=suggested.suggested_domains
    )
    plain = ArchitectureKnowledgeMatch("v1", bare.systems, ())
    assert render_guidance(GenerationGuidance(architecture=match)) == render_guidance(
        GenerationGuidance(architecture=plain)
    )
    with pytest.raises(InvalidArchitectureContentError, match="no catalogued system"):
        replace(impact, suggested_domains=suggested.suggested_domains)
    with pytest.raises(InvalidArchitectureContentError, match="needs its path"):
        SystemCapability("quote", "Quote", "quotes")


def test_exports_carry_the_capability_domain() -> None:
    document = _document("owner")
    feature = document.epic.features[0]
    assert feature.architecture is not None
    system = feature.architecture.systems[0]
    with_domain = replace(
        system,
        capabilities=(replace(system.capabilities[0], domain_path=("Order capture", "Quoting")),),
    )
    document = replace(
        document,
        epic=replace(
            document.epic,
            features=(
                replace(
                    feature,
                    architecture=replace(
                        feature.architecture,
                        systems=(with_domain, *feature.architecture.systems[1:]),
                    ),
                ),
            ),
        ),
    )

    workbook = load_workbook(io.BytesIO(XlsxBacklogExporter().render(document)))
    header = [cell.value for cell in workbook["Systems"][1]]
    row = next(workbook["Systems"].iter_rows(min_row=2, values_only=True))
    assert row[header.index("capability_domains")] == "Order capture > Quoting"
    payload = json.loads(JsonBacklogExporter().render(document))
    capability = payload["epic"]["features"][0]["architecture"]["systems"][0]["capabilities"][0]
    assert capability["domain_path"] == ["Order capture", "Quoting"]
