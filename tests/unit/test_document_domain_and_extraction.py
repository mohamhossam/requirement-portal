"""Source-document invariants, safe extraction, storage, and snapshot tests."""

from __future__ import annotations

import hashlib
import io
import multiprocessing
import threading
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from queue import Empty
from typing import cast

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image
from pypdf import PdfWriter
from smb_kernel.documents.bounded_extractor import (
    BoundedSubprocessDocumentExtractor,
    ExtractionLimits,
)
from smb_kernel.documents.ports import ExtractedDocument
from smb_kernel.documents.text_extractor import (
    SafeDocumentTextExtractor,
)

from smb_requirement_agent.application.errors import (
    DocumentExtractionBusyError,
    DocumentExtractionError,
    DocumentExtractionTimeoutError,
    DocumentStorageError,
    UnsupportedDocumentError,
)
from smb_requirement_agent.domain.document.entities import (
    SourceDocument,
    SourceDocumentVersion,
)
from smb_requirement_agent.domain.document.errors import DocumentInclusionError
from smb_requirement_agent.domain.document.value_objects import (
    DocumentId,
    DocumentVersionId,
    ExtractionStatus,
)
from smb_requirement_agent.infrastructure.persistence.document_payloads import (
    document_from_payload,
    document_to_payload,
)
from smb_requirement_agent.infrastructure.persistence.in_memory_document_repository import (
    InMemoryDocumentStorage,
)
from smb_requirement_agent.shared_kernel.identifiers import RequirementId

NOW = datetime(2026, 9, 3, 12, 0, tzinfo=UTC)


def _version(
    identifier: str = "version-1",
    *,
    number: int = 1,
    text: str = "Source text",
) -> SourceDocumentVersion:
    content = text.encode()
    return SourceDocumentVersion(
        id=DocumentVersionId(identifier),
        number=number,
        filename="source.txt",
        mime_type="text/plain",
        size_bytes=len(content),
        checksum_sha256=hashlib.sha256(content).hexdigest(),
        extraction_status=ExtractionStatus.READY,
        created_at=NOW,
        extracted_text=text,
    )


def _docx(document_xml: bytes, *extra_names: str) -> bytes:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("word/document.xml", document_xml)
        for name in extra_names:
            archive.writestr(name, b"unsafe")
    return stream.getvalue()


def _docx_with_numbering(document_xml: bytes, numbering_xml: bytes) -> bytes:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("word/document.xml", document_xml)
        archive.writestr("word/numbering.xml", numbering_xml)
    return stream.getvalue()


def test_document_versions_are_immutable_and_new_version_requires_reselection() -> None:
    first = _version()
    document = SourceDocument(
        id=DocumentId("document-1"),
        requirement_id=RequirementId("requirement-1"),
        versions=(first,),
    ).set_included(True)
    second = _version("version-2", number=2, text="Changed source")

    updated = document.add_version(second)

    assert updated.versions == (first, second)
    assert document.versions == (first,)
    assert updated.included_version_id is None


def test_failed_extraction_version_cannot_be_included() -> None:
    content = b"bad"
    failed = SourceDocumentVersion(
        id=DocumentVersionId("failed-version"),
        number=1,
        filename="bad.pdf",
        mime_type="application/pdf",
        size_bytes=len(content),
        checksum_sha256=hashlib.sha256(content).hexdigest(),
        extraction_status=ExtractionStatus.FAILED,
        created_at=NOW,
        extraction_error="PDF text extraction failed.",
    )
    document = SourceDocument(
        id=DocumentId("document-1"),
        requirement_id=RequirementId("requirement-1"),
        versions=(failed,),
    )

    with pytest.raises(DocumentInclusionError):
        document.set_included(True)


def test_document_snapshot_round_trip_preserves_immutable_metadata() -> None:
    document = SourceDocument(
        id=DocumentId("document-1"),
        requirement_id=RequirementId("requirement-1"),
        versions=(_version(),),
    ).set_included(True)

    assert document_from_payload(document_to_payload(document)) == document


def test_txt_extraction_accepts_utf8_and_rejects_blank_or_binary() -> None:
    extractor = SafeDocumentTextExtractor()

    assert extractor.extract("text/plain", b"  Customer rule  ") == "Customer rule"
    with pytest.raises(DocumentExtractionError, match="no usable text"):
        extractor.extract("text/plain", b"  \n")
    with pytest.raises(UnsupportedDocumentError, match="NUL"):
        extractor.extract("text/plain", b"text\x00data")


def test_extraction_rejects_character_pdf_page_and_xml_node_limits() -> None:
    with pytest.raises(UnsupportedDocumentError, match="characters"):
        SafeDocumentTextExtractor(max_extracted_characters=4).extract("text/plain", b"12345")

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_blank_page(width=100, height=100)
    pdf = io.BytesIO()
    writer.write(pdf)
    with pytest.raises(UnsupportedDocumentError, match="more than 1 pages"):
        SafeDocumentTextExtractor(max_pdf_pages=1).extract_structured(
            "application/pdf", pdf.getvalue()
        )

    xml = (
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/'
        b'wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Text</w:t>'
        b"</w:r></w:p></w:body></w:document>"
    )
    with pytest.raises(UnsupportedDocumentError, match="XML nodes"):
        SafeDocumentTextExtractor(max_xml_nodes=2).extract_structured(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            _docx(xml),
        )


def test_pdf_and_docx_signatures_must_match_declared_type() -> None:
    extractor = SafeDocumentTextExtractor()

    with pytest.raises(UnsupportedDocumentError, match="declared PDF"):
        extractor.extract("application/pdf", b"not-a-pdf")
    with pytest.raises(UnsupportedDocumentError, match="declared DOCX"):
        extractor.extract(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            b"not-a-docx",
        )


def test_corrupt_pdf_and_docx_are_explicit_extraction_failures() -> None:
    extractor = SafeDocumentTextExtractor()

    with pytest.raises(DocumentExtractionError, match="PDF text extraction failed"):
        extractor.extract("application/pdf", b"%PDF-corrupt")
    with pytest.raises(DocumentExtractionError, match="DOCX package is corrupt"):
        extractor.extract(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            b"PK-corrupt",
        )


def test_docx_extracts_plain_text_and_rejects_macros_and_traversal() -> None:
    extractor = SafeDocumentTextExtractor()
    xml = (
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b"<w:body><w:p><w:r><w:t>&lt;customer&gt;</w:t></w:r></w:p></w:body>"
        b"</w:document>"
    )
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    assert extractor.extract(mime, _docx(xml)) == "<customer>"
    with pytest.raises(UnsupportedDocumentError, match="Macro-enabled"):
        extractor.extract(mime, _docx(xml, "word/vbaProject.bin"))
    with pytest.raises(UnsupportedDocumentError, match="unsafe archive path"):
        extractor.extract(mime, _docx(xml, "../payload.bin"))


def test_docx_preserves_numbering_and_marks_merges_without_copying_anchor_text() -> None:
    document_xml = (
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/'
        b'wordprocessingml/2006/main"><w:body><w:p><w:pPr><w:numPr>'
        b'<w:ilvl w:val="0"/><w:numId w:val="7"/></w:numPr></w:pPr>'
        b"<w:r><w:t>First step</w:t></w:r></w:p><w:tbl><w:tr>"
        b"<w:tc><w:p><w:r><w:t>Field</w:t></w:r></w:p></w:tc>"
        b"<w:tc><w:p><w:r><w:t>Value</w:t></w:r></w:p></w:tc></w:tr><w:tr>"
        b'<w:tc><w:tcPr><w:vMerge w:val="restart"/></w:tcPr>'
        b"<w:p><w:r><w:t>Scope</w:t></w:r></w:p></w:tc>"
        b"<w:tc><w:p><w:r><w:t>A</w:t></w:r></w:p></w:tc></w:tr><w:tr>"
        b"<w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc>"
        b"<w:tc><w:p><w:r><w:t>B</w:t></w:r></w:p></w:tc></w:tr><w:tr>"
        b'<w:tc><w:tcPr><w:gridSpan w:val="2"/></w:tcPr>'
        b"<w:p><w:r><w:t>Shared note</w:t></w:r></w:p></w:tc></w:tr>"
        b"</w:tbl></w:body></w:document>"
    )
    numbering_xml = (
        b'<w:numbering xmlns:w="http://schemas.openxmlformats.org/'
        b'wordprocessingml/2006/main"><w:abstractNum w:abstractNumId="1">'
        b'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/>'
        b'<w:lvlText w:val="%1."/></w:lvl></w:abstractNum><w:num w:numId="7">'
        b'<w:abstractNumId w:val="1"/></w:num></w:numbering>'
    )

    result = SafeDocumentTextExtractor().extract_structured(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        _docx_with_numbering(document_xml, numbering_xml),
    )

    assert any(block.text == "1. First step" for block in result.evidence_blocks)
    assert any(
        block.text is not None
        and "[vertical merge continuation]" in block.text
        and "Scope" not in block.text
        for block in result.evidence_blocks
    )
    assert any(
        block.text is not None and "Shared note [spans 2 columns]" in block.text
        for block in result.evidence_blocks
    )


def test_xlsx_preserves_coordinates_formulas_images_charts_and_hidden_sheet_warnings(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    assert workbook.active is not None
    visible = cast(Worksheet, workbook.active)
    visible.title = "Pricing"
    visible["A1"] = "Product"
    visible["B1"] = 1
    visible["B2"] = "=1+1"
    chart = BarChart()
    chart.title = "Product volume"
    chart.add_data(Reference(visible, min_col=2, min_row=1, max_row=2))
    visible.add_chart(chart, "D2")
    image_path = tmp_path / "workflow.png"
    Image.new("RGB", (240, 120), color=(40, 90, 140)).save(image_path)
    visible.add_image(SpreadsheetImage(str(image_path)), "F2")
    hidden = workbook.create_sheet("Internal")
    hidden.sheet_state = "hidden"
    hidden["C3"] = "Sensitive rule"
    hidden.add_image(SpreadsheetImage(str(image_path)), "D4")
    stream = io.BytesIO()
    workbook.save(stream)

    result = SafeDocumentTextExtractor().extract_structured(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        stream.getvalue(),
    )

    assert any(
        block.text is not None and block.text.startswith("B2==1+1")
        for block in result.evidence_blocks
    )
    assert any(
        block.section_path == ("Hidden worksheet: Worksheet 2",)
        and block.text == "C3=Sensitive rule"
        for block in result.evidence_blocks
    )
    assert any(warning.code == "hidden_worksheet" for warning in result.warnings)
    assert any(warning.code == "formula_cache_unavailable" for warning in result.warnings)
    assert any(
        block.text is not None and "Product volume" in block.text
        for block in result.evidence_blocks
    )
    assert any(
        block.kind.value == "image" and block.section_path == ("Worksheet 1",)
        for block in result.evidence_blocks
    )
    assert any(
        block.kind.value == "image" and block.section_path == ("Hidden worksheet: Worksheet 2",)
        for block in result.evidence_blocks
    )
    assert any(
        block.text is not None
        and "Product volume" in block.text
        and block.section_path == ("Worksheet 1",)
        for block in result.evidence_blocks
    )


def test_xlsx_rejects_macros_and_archive_traversal() -> None:
    extractor = SafeDocumentTextExtractor()
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook_xml = b"<workbook />"

    with pytest.raises(UnsupportedDocumentError, match="Macro-enabled"):
        extractor.extract(
            mime,
            _office_package("xl/workbook.xml", workbook_xml, "xl/vbaProject.bin"),
        )
    with pytest.raises(UnsupportedDocumentError, match="unsafe archive path"):
        extractor.extract(mime, _office_package("xl/workbook.xml", workbook_xml, "../x"))


def test_xlsx_counts_blank_cells_and_rejects_decoded_images_over_limits(tmp_path: Path) -> None:
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    sparse = Workbook()
    assert sparse.active is not None
    sparse_sheet = cast(Worksheet, sparse.active)
    sparse_sheet["Z100"] = "far cell"
    sparse_bytes = io.BytesIO()
    sparse.save(sparse_bytes)
    with pytest.raises(UnsupportedDocumentError, match="visited cells"):
        SafeDocumentTextExtractor(max_visited_spreadsheet_cells=100).extract_structured(
            mime, sparse_bytes.getvalue()
        )

    with_image = Workbook()
    assert with_image.active is not None
    image_sheet = cast(Worksheet, with_image.active)
    image_sheet["A1"] = "evidence"
    image_path = tmp_path / "over-limit.png"
    Image.new("RGB", (200, 200), color=(10, 20, 30)).save(image_path)
    image_sheet.add_image(SpreadsheetImage(str(image_path)), "B2")
    image_bytes = io.BytesIO()
    with_image.save(image_bytes)
    with pytest.raises(UnsupportedDocumentError, match="pixel limit"):
        SafeDocumentTextExtractor(max_decoded_image_pixels=1_000).extract_structured(
            mime, image_bytes.getvalue()
        )


def _office_package(main_name: str, main_content: bytes, extra_name: str) -> bytes:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr(main_name, main_content)
        archive.writestr(extra_name, b"unsafe")
    return stream.getvalue()


def test_in_memory_document_bytes_are_immutable_and_idempotent() -> None:
    storage = InMemoryDocumentStorage()
    version_id = DocumentVersionId("immutable-version")

    storage.put(version_id, b"original")
    storage.put(version_id, b"original")

    assert storage.get(version_id) == b"original"
    with pytest.raises(DocumentStorageError, match="immutable"):
        storage.put(version_id, b"replacement")


class FakeExtractionQueue:
    def __init__(self, payload: object | None, *, block: bool = False) -> None:
        self.payload = payload
        self.block = block
        self.entered = threading.Event()
        self.release_result = threading.Event()

    def get(self, timeout: float) -> object:
        del timeout
        self.entered.set()
        if self.block:
            self.release_result.wait(timeout=2)
        if self.payload is None:
            raise Empty
        return self.payload

    def close(self) -> None:
        return None

    def join_thread(self) -> None:
        return None


class FakeExtractionEvent:
    def set(self) -> None:
        return None


class FakeExtractionProcess:
    pid = 41

    def __init__(self) -> None:
        self.alive = False
        self.terminated = False

    def start(self) -> None:
        self.alive = True

    def is_alive(self) -> bool:
        return self.alive

    def terminate(self) -> None:
        self.terminated = True
        self.alive = False

    def join(self, timeout: float | None = None) -> None:
        del timeout
        self.alive = False


class FakeExtractionContext:
    def __init__(self, queue: FakeExtractionQueue) -> None:
        self.queue = queue
        self.process = FakeExtractionProcess()

    def Queue(self, maxsize: int) -> FakeExtractionQueue:
        assert maxsize == 1
        return self.queue

    def Event(self) -> FakeExtractionEvent:
        return FakeExtractionEvent()

    def Process(self, **kwargs: object) -> FakeExtractionProcess:
        assert kwargs["daemon"] is True
        return self.process


class RecordingResourceLimiter:
    def __init__(self, *, fail: bool = False) -> None:
        self.fail = fail
        self.applied: list[tuple[int, int]] = []
        self.released: list[object | None] = []

    def apply(self, process_id: int, memory_bytes: int) -> object:
        self.applied.append((process_id, memory_bytes))
        if self.fail:
            raise OSError("platform detail")
        return "limit-handle"

    def release(self, handle: object | None) -> None:
        self.released.append(handle)


def _bounded_extractor(
    context: FakeExtractionContext,
    limiter: RecordingResourceLimiter,
    *,
    concurrency: int = 1,
    waiting_requests: int = 0,
) -> BoundedSubprocessDocumentExtractor:
    return BoundedSubprocessDocumentExtractor(
        concurrency=concurrency,
        waiting_requests=waiting_requests,
        deadline_seconds=0.01,
        memory_bytes=512 * 1024 * 1024,
        limits=ExtractionLimits(200, 1_000_000, 250_000, 20_000_000, 1_000_000),
        resource_limiter=limiter,
        process_context=cast(multiprocessing.context.BaseContext, context),
    )


def test_bounded_extractor_rejects_saturation_and_releases_capacity() -> None:
    extracted = ExtractedDocument("text", "v1", (), (), ())
    queue = FakeExtractionQueue(("ok", extracted), block=True)
    context = FakeExtractionContext(queue)
    extractor = _bounded_extractor(context, RecordingResourceLimiter())
    completed: list[ExtractedDocument] = []
    worker = threading.Thread(
        target=lambda: completed.append(extractor.extract_structured("text/plain", b"one"))
    )
    worker.start()
    assert queue.entered.wait(timeout=1)

    with pytest.raises(DocumentExtractionBusyError, match="busy"):
        extractor.extract_structured("text/plain", b"two")

    queue.release_result.set()
    worker.join(timeout=2)
    assert completed == [extracted]


def test_bounded_extractor_terminates_timed_out_child_and_releases_limit() -> None:
    context = FakeExtractionContext(FakeExtractionQueue(None))
    limiter = RecordingResourceLimiter()
    extractor = _bounded_extractor(context, limiter)

    with pytest.raises(DocumentExtractionTimeoutError, match="deadline"):
        extractor.extract_structured("text/plain", b"content")

    assert context.process.terminated is True
    assert limiter.applied == [(41, 512 * 1024 * 1024)]
    assert limiter.released == ["limit-handle"]


def test_bounded_extractor_hides_resource_limiter_diagnostics_and_stops_child() -> None:
    context = FakeExtractionContext(FakeExtractionQueue(("ok", b"unused")))
    limiter = RecordingResourceLimiter(fail=True)
    extractor = _bounded_extractor(context, limiter)

    with pytest.raises(DocumentExtractionError, match="resource constrained") as captured:
        extractor.extract_structured("text/plain", b"content")

    assert "platform detail" not in str(captured.value)
    assert context.process.terminated is True
    assert limiter.released == [None]
