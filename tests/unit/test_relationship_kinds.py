"""How one system depends on another: kinds on catalogue relationships (ADR-0088)."""

from __future__ import annotations

import io
from dataclasses import replace
from datetime import UTC, datetime

import pytest
from openpyxl import Workbook
from pydantic import TypeAdapter

from smb_requirement_agent.application.ports.architecture_knowledge import ArchitectureQuery
from smb_requirement_agent.application.ports.catalogue_extractor import (
    ExtractionRequest,
    ExtractionSegment,
)
from smb_requirement_agent.application.ports.catalogue_file import CatalogueFileFormat
from smb_requirement_agent.application.use_cases.approval_policy import artifact_fingerprint
from smb_requirement_agent.application.use_cases.catalogue_candidates import _settled
from smb_requirement_agent.domain.architecture.candidates import (
    CandidateContent,
    CandidateKind,
    CandidateMatch,
    apply_candidate,
    classify,
)
from smb_requirement_agent.domain.architecture.diff import ChangedItem, ChangeKind, diff_releases
from smb_requirement_agent.domain.architecture.entities import (
    ArchitectureDependency,
    ArchitectureImpact,
    SystemReference,
)
from smb_requirement_agent.domain.architecture.errors import InvalidArchitectureContentError
from smb_requirement_agent.domain.architecture.knowledge import (
    ArchitectureKnowledge,
    InvalidKnowledgeError,
    KnowledgeReleaseStatus,
    RelationshipKind,
    SystemRelationship,
)
from smb_requirement_agent.domain.architecture.neighbours import adjacent
from smb_requirement_agent.infrastructure.architecture.catalogue_files import CatalogueFileAdapter
from smb_requirement_agent.infrastructure.architecture.knowledge_yaml import seed_knowledge
from smb_requirement_agent.infrastructure.architecture.yaml_knowledge import (
    YamlArchitectureKnowledge,
    default_knowledge_path,
)
from smb_requirement_agent.infrastructure.llm.catalogue_extraction import FakeCatalogueExtractor
from smb_requirement_agent.infrastructure.persistence.shared_payloads import (
    architecture_from_payload,
    architecture_to_payload,
)
from tests.unit.test_feature_domain import make_feature

NOW = datetime(2026, 1, 1, 12, 0, tzinfo=UTC)
ADAPTER = CatalogueFileAdapter()


def _draft() -> ArchitectureKnowledge:
    return replace(
        seed_knowledge(),
        id="draft",
        status=KnowledgeReleaseStatus.DRAFT,
        published_at=None,
        published_by=None,
    )


def test_a_relationship_is_unspecified_until_a_source_says_how() -> None:
    assert SystemRelationship("a", "b", "uses").kind is RelationshipKind.UNSPECIFIED
    stored = SystemRelationship("a", "b", "uses", "calls_api")  # type: ignore[arg-type]
    assert stored.kind is RelationshipKind.CALLS_API
    with pytest.raises(InvalidKnowledgeError, match="Unknown relationship kind"):
        SystemRelationship("a", "b", "uses", "telepathy")  # type: ignore[arg-type]
    with pytest.raises(InvalidArchitectureContentError, match="Unknown relationship kind"):
        ArchitectureDependency("a", "b", "uses", "telepathy")  # type: ignore[arg-type]


def test_releases_stored_before_kinds_existed_still_load() -> None:
    stored = TypeAdapter(ArchitectureKnowledge).dump_python(seed_knowledge(), mode="json")
    for item in stored["relationships"]:
        item.pop("kind")

    loaded = TypeAdapter(ArchitectureKnowledge).validate_python(stored)

    assert {item.kind for item in loaded.relationships} == {RelationshipKind.UNSPECIFIED}


def test_the_seed_names_kinds_only_where_its_descriptions_show_them() -> None:
    kinds = {
        (item.source_system_id, item.target_system_id): item.kind
        for item in seed_knowledge().relationships
    }

    assert kinds[("b2b-web", "b2b-bff")] is RelationshipKind.CALLS_API
    assert kinds[("rtf", "cwom")] is RelationshipKind.ORCHESTRATES
    assert kinds[("dcrm", "cbcm-crmgw")] is RelationshipKind.UNSPECIFIED


def test_changing_a_kind_is_an_edit_to_the_same_relationship() -> None:
    base = seed_knowledge()
    first, *rest = base.relationships
    draft = replace(
        base, id="draft", relationships=(replace(first, kind=RelationshipKind.ORCHESTRATES), *rest)
    )

    changes = diff_releases(base, draft).changes

    assert [(item.item, item.change, item.fields) for item in changes] == [
        (ChangedItem.RELATIONSHIP, ChangeKind.CHANGED, ("kind",))
    ]
    assert changes[0].key == (
        f"{first.source_system_id}->{first.target_system_id}:{first.description.casefold()}"
    )


def test_a_suggestion_sets_a_stated_kind_but_never_clears_one() -> None:
    draft = _draft()
    link = next(item for item in draft.relationships if item.source_system_id == "dcrm")

    def suggestion(kind: RelationshipKind | None) -> CandidateContent:
        return CandidateContent(
            CandidateKind.RELATIONSHIP,
            "dcrm",
            target_system_id="cbcm-crmgw",
            text=link.description,
            relationship_kind=kind,
        )

    assert classify(suggestion(None), draft) is CandidateMatch.ALREADY_PRESENT
    assert classify(suggestion(RelationshipKind.UNSPECIFIED), draft) is (
        CandidateMatch.ALREADY_PRESENT
    )
    stated = suggestion(RelationshipKind.CALLS_API)
    assert classify(stated, draft) is CandidateMatch.UPDATES_EXISTING

    updated = apply_candidate(stated, draft)
    changed = next(item for item in updated.relationships if item.source_system_id == "dcrm")
    assert changed.kind is RelationshipKind.CALLS_API
    assert len(updated.relationships) == len(draft.relationships)
    kept = apply_candidate(suggestion(RelationshipKind.UNSPECIFIED), updated)
    assert next(item for item in kept.relationships if item.source_system_id == "dcrm").kind is (
        RelationshipKind.CALLS_API
    )
    with pytest.raises(InvalidKnowledgeError, match="Only a dependency"):
        CandidateContent(
            CandidateKind.CONSTRAINT, "dcrm", text="x", relationship_kind=RelationshipKind.CALLS_API
        )


def test_merged_proposals_keep_one_kind_and_never_pick_between_two() -> None:
    content = CandidateContent(CandidateKind.RELATIONSHIP, "a", target_system_id="b", text="orders")
    calls, events = RelationshipKind.CALLS_API, RelationshipKind.PUBLISHES_EVENTS_TO

    assert _settled(content, [(True, None), (False, calls)]).relationship_kind is calls
    assert _settled(content, [(True, calls), (False, events)]).relationship_kind is calls
    assert _settled(content, [(False, calls), (False, events)]).relationship_kind is (
        RelationshipKind.UNSPECIFIED
    )


@pytest.mark.parametrize("file_format", list(CatalogueFileFormat))
def test_every_file_format_keeps_the_kind(file_format: CatalogueFileFormat) -> None:
    seed = seed_knowledge()

    content = ADAPTER.read(file_format, ADAPTER.write(file_format, seed))

    assert content.relationships == seed.relationships


def _workbook(relationships: list[tuple[object, ...]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    systems = workbook.create_sheet("Systems")
    for row in (("system_id", "name"), ("crm", "CRM"), ("billing", "Billing")):
        systems.append(list(row))
    sheet = workbook.create_sheet("Relationships")
    for relationship in relationships:
        sheet.append(list(relationship))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_workbooks_without_the_kind_column_still_import_and_labels_are_accepted() -> None:
    old = ADAPTER.read(
        CatalogueFileFormat.XLSX,
        _workbook(
            [("source_system_id", "target_system_id", "description"), ("crm", "billing", "x")]
        ),
    )
    labelled = ADAPTER.read(
        CatalogueFileFormat.XLSX,
        _workbook(
            [
                ("source_system_id", "target_system_id", "description", "kind"),
                ("crm", "billing", "Invoices", "Calls API"),
                ("billing", "crm", "Receipts", None),
            ]
        ),
    )

    assert old.relationships == (SystemRelationship("crm", "billing", "x"),)
    assert [item.kind for item in labelled.relationships] == [
        RelationshipKind.CALLS_API,
        RelationshipKind.UNSPECIFIED,
    ]
    with pytest.raises(InvalidKnowledgeError, match="Relationships row 2: kind must be one of"):
        ADAPTER.read(
            CatalogueFileFormat.XLSX,
            _workbook(
                [
                    ("source_system_id", "target_system_id", "description", "kind"),
                    ("crm", "billing", "x", "gossip"),
                ]
            ),
        )


def test_mapping_and_connected_systems_carry_the_kind() -> None:
    seed = YamlArchitectureKnowledge(default_knowledge_path())

    result = seed.match(ArchitectureQuery(text=(), declared_systems=("B2B Web", "B2B BFF")))

    assert [item.kind for item in result.dependencies] == [RelationshipKind.CALLS_API]
    nearby = adjacent({"rtf"}, seed_knowledge().relationships)
    assert {item.kind for item in nearby.dependencies} == {
        RelationshipKind.TRANSFERS_DATA_TO,
        RelationshipKind.ORCHESTRATES,
    }


def _impact(kind: RelationshipKind) -> ArchitectureImpact:
    return ArchitectureImpact(
        "v1",
        NOW,
        (SystemReference("a", "A", True), SystemReference("b", "B", True)),
        (ArchitectureDependency("a", "b", "uses", kind),),
    )


def test_an_unspecified_kind_leaves_payloads_and_fingerprints_as_they_were() -> None:
    unspecified = _impact(RelationshipKind.UNSPECIFIED)
    stated = _impact(RelationshipKind.CALLS_API)

    payload = architecture_to_payload(unspecified)
    assert isinstance(payload, dict)
    assert payload["dependencies"] == [
        {"source_system_id": "a", "target_system_id": "b", "description": "uses"}
    ]
    assert architecture_from_payload(payload) == unspecified
    assert architecture_from_payload(architecture_to_payload(stated)) == stated

    feature = make_feature()
    before_kinds = replace(
        unspecified,
        dependencies=(ArchitectureDependency("a", "b", "uses"),),
    )
    assert artifact_fingerprint(feature.with_architecture(unspecified)) == artifact_fingerprint(
        feature.with_architecture(before_kinds)
    )
    assert artifact_fingerprint(feature.with_architecture(stated)) != artifact_fingerprint(
        feature.with_architecture(unspecified)
    )


def test_the_offline_extractor_reads_a_call_as_an_api_dependency() -> None:
    proposal = FakeCatalogueExtractor().propose(
        ExtractionRequest(
            "Integration notes",
            (
                ExtractionSegment(
                    1,
                    "line 1",
                    "Order Hub calls BCRM for quotes\nOrder Hub sends invoices to Billing\n"
                    "Order Hub depends on CRM for customers",
                ),
            ),
            (),
        )
    )

    assert [item.content.relationship_kind for item in proposal.changes] == [
        RelationshipKind.CALLS_API,
        RelationshipKind.TRANSFERS_DATA_TO,
        RelationshipKind.UNSPECIFIED,
    ]
