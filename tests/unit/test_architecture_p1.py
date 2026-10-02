"""Architecture mapping: named systems are always found, Word sections are chunked
with their headings, and models come from the application's configuration."""

from __future__ import annotations

import io
import json
import zipfile
from dataclasses import replace
from datetime import UTC, datetime
from types import SimpleNamespace
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from smb_requirement_agent.application.errors import (
    DocumentExtractionError,
    ModelTransportError,
    UnsupportedDocumentError,
)
from smb_requirement_agent.application.ports.architecture_knowledge import ArchitectureQuery
from smb_requirement_agent.application.ports.architecture_rag import (
    ArchitectureEvidenceError,
    EvidenceChunk,
    EvidenceSelection,
)
from smb_requirement_agent.application.ports.identity import Actor
from smb_requirement_agent.application.use_cases.architecture_evidence import gather_evidence
from smb_requirement_agent.application.use_cases.architecture_index import BuildArchitectureIndex
from smb_requirement_agent.application.use_cases.architecture_knowledge import (
    ManageArchitectureKnowledge,
)
from smb_requirement_agent.application.use_cases.resolve_architecture_knowledge import (
    ResolveArchitectureKnowledge,
)
from smb_requirement_agent.domain.architecture.knowledge import (
    ArchitectureKnowledge,
    KnowledgeDocumentVersion,
    SystemDefinition,
)
from smb_requirement_agent.domain.document.value_objects import DocumentVersionId
from smb_requirement_agent.infrastructure.architecture.catalogue_files import CatalogueFileAdapter
from smb_requirement_agent.infrastructure.architecture.embeddings import (
    DIMENSIONS,
    ArchitectureEmbeddings,
    FakeEmbeddings,
)
from smb_requirement_agent.infrastructure.architecture.evidence_index import InMemoryEvidenceIndex
from smb_requirement_agent.infrastructure.architecture.knowledge_yaml import seed_knowledge
from smb_requirement_agent.infrastructure.architecture.located_extractor import (
    LocatedDocumentExtractor,
)
from smb_requirement_agent.infrastructure.architecture.reasoning import (
    FakeArchitectureReasoner,
    StructuredArchitectureReasoner,
)
from smb_requirement_agent.infrastructure.architecture.tokenizer import (
    ApproximateTokenizer,
    FakeWordTokenizer,
)
from smb_requirement_agent.infrastructure.config.options import LLMProvider
from smb_requirement_agent.infrastructure.config.settings import Settings
from smb_requirement_agent.infrastructure.documents.text_extractor import SafeDocumentTextExtractor
from smb_requirement_agent.infrastructure.observability.metrics import Metrics
from smb_requirement_agent.infrastructure.persistence.in_memory_architecture_knowledge import (
    InMemoryArchitectureKnowledgeRepository,
)
from smb_requirement_agent.infrastructure.persistence.in_memory_document_repository import (
    InMemoryDocumentStorage,
)
from smb_requirement_agent.infrastructure.persistence.in_memory_organisation import (
    InMemoryOrganisationRepository,
)
from smb_requirement_agent.infrastructure.time.fixed_clock import FixedClock
from smb_requirement_agent.interfaces.api.composition.llm import build_llm_adapters
from tests.spreadsheet_fixtures import XLSX_MIME

NOW = datetime(2026, 9, 30, 9, 0, tzinfo=UTC)
MAINTAINER = Actor("amina", frozenset({"knowledge_maintainer"}))
OWNER = {"X-Fake-Actor-Id": "fake-owner"}
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _published(index: InMemoryEvidenceIndex) -> ArchitectureKnowledge:
    seed = seed_knowledge()
    return replace(
        seed,
        id="rel-2",
        built_revision=seed.revision,
        index_profile=index.profile,
        index_id="idx",
    )


def _crowded_index() -> tuple[InMemoryEvidenceIndex, ArchitectureKnowledge]:
    """BCC's own record, outranked by twelve passages about the same flow."""
    index = InMemoryEvidenceIndex(FakeEmbeddings(), FakeWordTokenizer())
    release = _published(index)
    passage = "the catalogue upload flow publishes offers nightly to the portal " * 3
    index.store(
        release.id,
        "idx",
        (
            EvidenceChunk("own", "Oracle ATG BCC", "system oracle-atg-bcc", "Oracle ATG BCC\nBCC"),
            *(
                EvidenceChunk(f"d{n}", "Reference", f"paragraph {n}", passage, "doc")
                for n in range(12)
            ),
        ),
    )
    return index, release


QUERY = "The catalogue upload flow publishes offers nightly to the portal via BCC"


def test_a_named_system_is_found_even_when_search_ranks_its_record_out() -> None:
    index, release = _crowded_index()

    assert all(chunk.id != "own" for chunk in index.retrieve("idx", QUERY, 8))
    evidence = gather_evidence(index, release, "idx", QUERY)

    assert evidence[0].id == "own"
    assert len(evidence) == 8


def test_mapping_adds_the_named_system_when_the_model_selects_nothing() -> None:
    index, release = _crowded_index()

    class Silent:
        model = "silent"

        def select(self, *_: object) -> EvidenceSelection:
            return EvidenceSelection(
                (), (), "Insufficient evidence to identify a catalogue system."
            )

    resolver = ResolveArchitectureKnowledge(
        InMemoryArchitectureKnowledgeRepository(release),
        index,
        Silent(),
        FakeArchitectureReasoner(),  # type: ignore[arg-type]
        InMemoryOrganisationRepository(FixedClock(NOW)),
    )

    match = resolver.match(ArchitectureQuery((QUERY,), release_id=release.id))

    assert [system.id for system in match.systems] == ["oracle-atg-bcc"]
    assert match.citation_ids == ("own",)
    assert match.uncertainty is None


def _paragraph(text: str, style: str | None = None) -> str:
    properties = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    return f"<w:p>{properties}<w:r><w:t>{text}</w:t></w:r></w:p>"


def _docx() -> bytes:
    body = "".join(
        (
            _paragraph("Ordering", "Heading1"),
            # A localised style id; styles.xml says it is a second-level heading.
            _paragraph("Catalogue", "Ueberschrift2"),
            _paragraph("Offers are authored in BCC."),
            _paragraph("It is validated nightly."),
            _paragraph("Billing", "Heading1"),
            _paragraph("Invoices close the order."),
            "<w:tbl><w:tr>"
            f"<w:tc>{_paragraph('BCRM')}</w:tc><w:tc>{_paragraph('Sales CRM')}</w:tc>"
            "</w:tr></w:tbl>",
        )
    )
    styles = (
        f'<w:styles xmlns:w="{W}"><w:style w:type="paragraph" w:styleId="Ueberschrift2">'
        '<w:name w:val="heading 2"/></w:style></w:styles>'
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr(
            "word/document.xml", f'<w:document xmlns:w="{W}"><w:body>{body}</w:body></w:document>'
        )
        archive.writestr("word/styles.xml", styles)
    return buffer.getvalue()


class _RecordingIndex(InMemoryEvidenceIndex):
    def __init__(self) -> None:
        super().__init__(FakeEmbeddings(), ApproximateTokenizer())
        self.chunks: tuple[EvidenceChunk, ...] = ()

    def store(self, release_id: str, index_id: str, chunks: tuple[EvidenceChunk, ...]) -> None:
        self.chunks = chunks
        super().store(release_id, index_id, chunks)


def _build(content: bytes, mime_type: str = DOCX) -> tuple[EvidenceChunk, ...]:
    repository = InMemoryArchitectureKnowledgeRepository(seed_knowledge())
    index = _RecordingIndex()
    manage = ManageArchitectureKnowledge(repository, CatalogueFileAdapter(), index, 10_000_000)
    storage = InMemoryDocumentStorage()
    draft = manage.create_draft(MAINTAINER, "Next version")
    storage.put(DocumentVersionId("doc"), content)
    version = KnowledgeDocumentVersion(
        "doc", "Reference", "reference", mime_type, "en", "sum", "doc", "amina", NOW
    )
    repository.save(draft.updated(documents=(version,)), draft.revision, "amina", "upload")
    current = manage.get(draft.id)
    BuildArchitectureIndex(
        manage,
        index,
        storage,
        LocatedDocumentExtractor(SafeDocumentTextExtractor()),
        ApproximateTokenizer(),
    ).execute(current.id, current.revision, "amina", fence=lambda: None)
    return index.chunks


def test_word_passages_are_packed_under_their_headings() -> None:
    documents = [chunk for chunk in _build(_docx()) if chunk.document_version_id == "doc"]

    assert [(chunk.location, chunk.text) for chunk in documents] == [
        (
            "paragraphs 3–4",
            "Reference / Ordering / Catalogue\n"
            "Offers are authored in BCC.\nIt is validated nightly.",
        ),
        ("paragraph 6", "Reference / Billing\nInvoices close the order."),
        ("table 1, row 1", "Reference / Billing\nBCRM | Sales CRM"),
    ]


def test_reading_for_suggestions_keeps_every_passage_and_its_location() -> None:
    located = LocatedDocumentExtractor(SafeDocumentTextExtractor()).extract(DOCX, _docx())

    assert [item.location for item in located] == [
        "paragraph 1",
        "paragraph 2",
        "paragraph 3",
        "paragraph 4",
        "paragraph 5",
        "paragraph 6",
        "table 1, row 1, cell 1, paragraph 7",
        "table 1, row 1, cell 2, paragraph 8",
    ]
    assert [item.text for item in located if item.heading] == ["Ordering", "Catalogue", "Billing"]
    assert located[2].heading_path == ("Ordering", "Catalogue")


def _inventory() -> bytes:
    workbook = Workbook()
    systems = workbook.active
    assert systems is not None
    systems.title = "Systems"
    systems.append(["System", "Owner"])
    systems.append(["Order Hub", "Sales"])
    systems.append(["BCRM", "Care"])
    links = workbook.create_sheet("Links")
    links.append(["From", "To"])
    links.append(["Order Hub", "BCRM"])
    hidden = workbook.create_sheet("Scratch")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Draft note nobody reviewed"
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_a_workbook_is_read_row_by_row_under_its_sheet_names() -> None:
    located = LocatedDocumentExtractor(SafeDocumentTextExtractor()).extract(XLSX_MIME, _inventory())

    assert [(item.location, item.text, item.heading_path, item.heading) for item in located] == [
        ("sheet 1", "Systems", (), True),
        ("sheet 1, row 1", "A1=System | B1=Owner", ("Systems",), False),
        ("sheet 1, row 2", "A2=Order Hub | B2=Sales", ("Systems",), False),
        ("sheet 1, row 3", "A3=BCRM | B3=Care", ("Systems",), False),
        ("sheet 2", "Links", (), True),
        ("sheet 2, row 1", "A1=From | B1=To", ("Links",), False),
        ("sheet 2, row 2", "A2=Order Hub | B2=BCRM", ("Links",), False),
    ]


@pytest.mark.parametrize(
    ("mime_type", "content"),
    [
        ("text/csv", b"System,Owner\nOrder Hub,Sales\n"),
        ("text/tab-separated-values", b"System\tOwner\nOrder Hub\tSales\n"),
    ],
)
def test_a_delimited_file_is_read_row_by_row(mime_type: str, content: bytes) -> None:
    located = LocatedDocumentExtractor(SafeDocumentTextExtractor()).extract(mime_type, content)

    assert [(item.location, item.text) for item in located] == [
        ("row 1", "R1C1: System | R1C2: Owner"),
        ("row 2", "R2C1: Order Hub | R2C2: Sales"),
    ]


def test_markdown_is_read_by_its_headings_while_plain_text_keeps_line_windows() -> None:
    content = b"# Systems\n\n- Order Hub\n" + b"line\n" * 40
    extractor = LocatedDocumentExtractor(SafeDocumentTextExtractor())

    markdown = extractor.extract("text/markdown", content)
    plain = extractor.extract("text/plain", content)

    assert [(item.location, item.heading) for item in markdown] == [
        ("line 1", True),
        ("lines 3-42", False),
        ("line 43", False),
    ]
    assert markdown[1].heading_path == ("Systems",)
    assert markdown[1].text.startswith("- Order Hub\nline")
    assert [item.location for item in plain] == ["lines 1-40", "lines 41-43"]


@pytest.mark.parametrize(
    ("content", "error"),
    [(b"# Systems\x00", UnsupportedDocumentError), (b"# Syst\xffems", DocumentExtractionError)],
)
def test_markdown_that_is_not_utf8_text_is_refused(content: bytes, error: type[Exception]) -> None:
    with pytest.raises(error):
        LocatedDocumentExtractor(SafeDocumentTextExtractor()).extract("text/markdown", content)


def test_markdown_rows_are_packed_under_their_heading_in_the_index() -> None:
    content = (
        b"# Systems\n\nCore systems.\n\n| System | Owner |\n|---|---|\n"
        b"| Order Hub | Sales |\n| BCRM | Care |\n\n## Links\n\nOrder Hub calls BCRM.\n"
    )

    documents = [
        (chunk.location, chunk.text)
        for chunk in _build(content, "text/markdown")
        if chunk.document_version_id == "doc"
    ]

    assert documents == [
        (
            "lines 3-8",
            "Reference / Systems\nCore systems.\nSystem: Order Hub | Owner: Sales\n"
            "System: BCRM | Owner: Care",
        ),
        ("line 12", "Reference / Systems / Links\nOrder Hub calls BCRM."),
    ]


def test_workbook_rows_are_packed_per_sheet_in_the_index() -> None:
    documents = [
        chunk for chunk in _build(_inventory(), XLSX_MIME) if chunk.document_version_id == "doc"
    ]

    assert [(chunk.location, chunk.text) for chunk in documents] == [
        (
            "sheet 1, rows 1–3",
            "Reference / Systems\nA1=System | B1=Owner\nA2=Order Hub | B2=Sales\nA3=BCRM | B3=Care",
        ),
        ("sheet 2, rows 1–2", "Reference / Links\nA1=From | B1=To\nA2=Order Hub | B2=BCRM"),
    ]


def test_system_records_carry_their_dependencies_both_ways() -> None:
    own = {chunk.location: chunk.text for chunk in _build(_docx())}

    assert (
        "Depends on: Oracle ATG BCC (API call) — Digital channels read the shared product "
        "catalogue through the BFF." in own["system b2b-bff"]
    )
    assert "Used by: B2B BFF (API call) — " in own["system oracle-atg-bcc"]
    # An unspecified kind adds nothing, so older chunk text is unchanged.
    assert "Depends on: CBCM / CRMGW — DCRM captures" in own["system dcrm"]
    # A placed capability names its domain, so retrieval can match domain words.
    assert "Assisted sales, quoting, and order capture (Order capture):" in own["system bcrm"]
    assert "\n\n" not in own["system oracle-atg-bcc"]


def test_approximate_tokens_cut_long_words_every_four_characters() -> None:
    assert ApproximateTokenizer().spans("ab abcdefghij") == ((0, 2), (3, 7), (7, 11), (11, 13))


class _Recording:
    model = "configured-embedding"

    def __init__(self, width: int = DIMENSIONS, fail: bool = False) -> None:
        self.batches: list[int] = []
        self.width = width
        self.fail = fail

    def embed(self, texts: tuple[str, ...]) -> tuple[tuple[float, ...], ...]:
        if self.fail:
            raise ModelTransportError("down")
        self.batches.append(len(texts))
        return tuple((1.0,) * self.width for _ in texts)


def test_architecture_embeddings_use_the_configured_model_in_bounded_batches() -> None:
    recording = _Recording()

    vectors = ArchitectureEmbeddings(recording).embed(tuple(f"text {n}" for n in range(70)))

    assert recording.batches == [32, 32, 6]
    assert len(vectors) == 70
    assert ArchitectureEmbeddings(recording).model == "configured-embedding"
    with pytest.raises(ArchitectureEvidenceError):
        ArchitectureEmbeddings(_Recording(fail=True)).embed(("x",))
    with pytest.raises(ArchitectureEvidenceError, match="768"):
        ArchitectureEmbeddings(_Recording(width=1024)).embed(("x",))


def test_the_reasoner_offers_named_systems_first_within_a_small_budget() -> None:
    prompts: list[str] = []

    class Client:
        model = "small"

        def parse(self, *, user_prompt: str, **_: Any) -> SimpleNamespace:
            prompts.append(user_prompt)
            return SimpleNamespace(impacts=[], uncertainty=None, conflicts=[])

    release = replace(
        seed_knowledge(),
        systems=(
            *(SystemDefinition(f"system-{n}", f"System number {n}") for n in range(300)),
            *seed_knowledge().systems,
        ),
    )

    reasoner = StructuredArchitectureReasoner(Client(), max_input_tokens=2000)  # type: ignore[arg-type]
    selected = reasoner.select(
        ArchitectureQuery(("Change the BCRM quote screen",)),
        release,
        (EvidenceChunk("e1", "BCRM", "system bcrm", "BCRM"),),
    )

    offered = json.loads(prompts[0])["catalogue_systems"]
    assert offered[0]["id"] == "bcrm"
    assert len(offered) < len(release.systems)
    assert selected.uncertainty is not None
    assert "catalogue systems did not fit the model's input" in selected.uncertainty


def test_architecture_models_follow_the_application_provider() -> None:
    fake = build_llm_adapters(Settings(llm_provider=LLMProvider.FAKE), Metrics())
    local = build_llm_adapters(
        Settings(llm_provider=LLMProvider.LOCAL, local_llm_model="m", local_embedding_model="e"),
        Metrics(),
    )
    try:
        assert isinstance(fake.architecture_reasoner, FakeArchitectureReasoner)
        assert isinstance(local.architecture_reasoner, StructuredArchitectureReasoner)
        assert local.architecture_reasoner.model == "m"
    finally:
        fake.close()
        local.close()


def test_the_latest_build_is_read_back_for_maintainers_only(client: TestClient) -> None:
    draft = client.post(
        "/architecture-knowledge/releases", json={"name": "Next version"}, headers=OWNER
    ).json()
    status = f"/architecture-knowledge/releases/{draft['id']}/build"
    assert client.get(status, headers=OWNER).json() is None

    started = client.post(
        status, json={"expected_revision": draft["revision"]}, headers=OWNER
    ).json()

    assert client.get(status, headers=OWNER).json() == started
    assert started["status"] == "succeeded"
    reader = {"X-Fake-Actor-Id": "fake-reviewer"}
    assert client.get(status, headers=reader).status_code == 403
    assert (
        client.get("/architecture-knowledge/releases/missing/build", headers=OWNER).status_code
        == 404
    )
