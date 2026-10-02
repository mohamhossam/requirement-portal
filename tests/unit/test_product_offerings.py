"""Product offerings: what a commercial product is made of and which systems deliver it (ADR-0095).

The reported gap: a landscape document's product, its order types, components
and each component's responsibilities per system had nowhere to go in the
catalogue.
"""

from __future__ import annotations

import io
from dataclasses import replace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from smb_kernel.documents.text_extractor import SafeDocumentTextExtractor

from smb_requirement_agent.application.ports.architecture_rag import EvidenceChunk
from smb_requirement_agent.application.ports.catalogue_file import CatalogueFileFormat
from smb_requirement_agent.application.ports.identity import Actor
from smb_requirement_agent.application.use_cases.architecture_index import BuildArchitectureIndex
from smb_requirement_agent.application.use_cases.architecture_knowledge import (
    ManageArchitectureKnowledge,
)
from smb_requirement_agent.domain.architecture.diff import ChangedItem, ChangeKind, diff_releases
from smb_requirement_agent.domain.architecture.knowledge import (
    ArchitectureKnowledge,
    InvalidKnowledgeError,
)
from smb_requirement_agent.domain.architecture.products import (
    ComponentResponsibility,
    OfferingComponent,
    OfferingPoint,
    OrderType,
    ProductOffering,
    SourceConfidence,
)
from smb_requirement_agent.infrastructure.architecture.catalogue_files import CatalogueFileAdapter
from smb_requirement_agent.infrastructure.architecture.embeddings import FakeEmbeddings
from smb_requirement_agent.infrastructure.architecture.evidence_index import InMemoryEvidenceIndex
from smb_requirement_agent.infrastructure.architecture.knowledge_yaml import seed_knowledge
from smb_requirement_agent.infrastructure.architecture.located_extractor import (
    LocatedDocumentExtractor,
)
from smb_requirement_agent.infrastructure.architecture.tokenizer import ApproximateTokenizer
from smb_requirement_agent.infrastructure.persistence.in_memory_architecture_knowledge import (
    InMemoryArchitectureKnowledgeRepository,
)
from smb_requirement_agent.infrastructure.persistence.in_memory_document_repository import (
    InMemoryDocumentStorage,
)

OWNER = {"X-Fake-Actor-Id": "fake-owner"}
MAINTAINER = Actor("amina", frozenset({"knowledge_maintainer"}))
ADAPTER = CatalogueFileAdapter()

FIBRE = OfferingComponent(
    "fibre",
    "Fibre Access",
    code="PO_FIBRE",
    kind="SERVICE",
    mandatory=True,
    customer_visible=True,
    description="Primary fixed connection.",
    technical_spec="CFSS_FIBRE",
    responsibilities=(
        ComponentResponsibility(
            "b2b-web",
            "Primary channel",
            "Captures the order and its options.",
            ("NEW_ACTIVATION",),
            SourceConfidence.CONFIRMED,
            "Synthetic design §3",
        ),
        ComponentResponsibility("bcrm", "CUSTOMER_CONTEXT", "Supplies the account."),
    ),
    confidence=SourceConfidence.CONFIRMED,
)
OFFICE = ProductOffering(
    "office-connect",
    "Office Connect",
    code="OFFICE_CONNECT",
    family="Connect",
    version="1.0.0",
    lifecycle="REFERENCE",
    proposition="A fixed connection with a managed router.",
    rules=("The router is mandatory.",),
    order_types=(
        OrderType("NEW_ACTIVATION", "New Activation"),
        OrderType("UPGRADE", "Upgrade", enabled=False, confidence=SourceConfidence.GAP),
    ),
    components=(FIBRE,),
    values=(OfferingPoint("Secure connectivity", "Firewall included.", SourceConfidence.INFERRED),),
    audiences=(OfferingPoint("Branch offices"),),
    confidence=SourceConfidence.CONFIRMED,
    source="Synthetic design",
)


def _with_offering(release: ArchitectureKnowledge | None = None) -> ArchitectureKnowledge:
    return replace(release or seed_knowledge(), products=(OFFICE,))


# Domain -----------------------------------------------------------------------------------


def test_a_role_is_one_code_however_it_is_written() -> None:
    assert FIBRE.responsibilities[0].role == "PRIMARY_CHANNEL"
    assert ComponentResponsibility("bcrm", "primary-orchestrator", "x").role == (
        "PRIMARY_ORCHESTRATOR"
    )
    # A document's confidence arrives as text and is kept as one of three.
    assert OrderType("A", "A", confidence="Inferred").confidence is SourceConfidence.INFERRED  # type: ignore[arg-type]


@pytest.mark.parametrize(
    ("build", "message"),
    [
        (
            lambda: replace(OFFICE, order_types=(*OFFICE.order_types, OrderType("upgrade", "x"))),
            "order type codes must be unique",
        ),
        (lambda: replace(OFFICE, components=(FIBRE, FIBRE)), "component ids must be unique"),
        (
            lambda: replace(
                FIBRE,
                responsibilities=(*FIBRE.responsibilities, FIBRE.responsibilities[0]),
            ),
            "same role in it more than once",
        ),
        (
            lambda: replace(
                OFFICE,
                components=(
                    replace(
                        FIBRE,
                        responsibilities=(
                            ComponentResponsibility("bcrm", "ROLE", "x", ("MIGRATION",)),
                        ),
                    ),
                ),
            ),
            "order type 'MIGRATION' is not one of the offering's",
        ),
        (lambda: OrderType("A", "A", confidence="likely"), "Confidence must be confirmed"),  # type: ignore[arg-type]
    ],
)
def test_an_offering_is_refused_when_it_contradicts_itself(build: object, message: str) -> None:
    with pytest.raises(InvalidKnowledgeError, match=message):
        build()  # type: ignore[operator]


def test_a_system_an_offering_names_cannot_be_removed() -> None:
    release = _with_offering()

    assert OFFICE.systems == ("b2b-web", "bcrm")
    with pytest.raises(InvalidKnowledgeError, match="Office Connect › Fibre Access names system"):
        replace(release, systems=tuple(item for item in release.systems if item.id != "bcrm"))


# Files ------------------------------------------------------------------------------------


@pytest.mark.parametrize("file_format", list(CatalogueFileFormat))
def test_every_format_keeps_offerings_whole(file_format: CatalogueFileFormat) -> None:
    release = _with_offering()

    content = ADAPTER.read(file_format, ADAPTER.write(file_format, release))

    assert content.products == (OFFICE,)


def test_the_workbook_says_which_row_is_wrong() -> None:
    workbook = load_workbook(io.BytesIO(ADAPTER.write(CatalogueFileFormat.XLSX, _with_offering())))

    def broken(sheet: str, column: str, value: str) -> bytes:
        rows = workbook[sheet]
        headers = [cell.value for cell in rows[1]]
        rows.cell(2, headers.index(column) + 1, value)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    with pytest.raises(
        InvalidKnowledgeError, match="OrderTypes row 2: product 'nope' is not listed"
    ):
        ADAPTER.read(CatalogueFileFormat.XLSX, broken("OrderTypes", "product_id", "nope"))
    workbook = load_workbook(io.BytesIO(ADAPTER.write(CatalogueFileFormat.XLSX, _with_offering())))
    with pytest.raises(
        InvalidKnowledgeError, match="OfferingComponents row 2: mandatory must be yes"
    ):
        ADAPTER.read(CatalogueFileFormat.XLSX, broken("OfferingComponents", "mandatory", "maybe"))
    workbook = load_workbook(io.BytesIO(ADAPTER.write(CatalogueFileFormat.XLSX, _with_offering())))
    with pytest.raises(InvalidKnowledgeError, match="Responsibilities row 2: component 'ghost'"):
        ADAPTER.read(CatalogueFileFormat.XLSX, broken("Responsibilities", "component_id", "ghost"))


def test_a_file_from_before_offerings_reads_as_it_did() -> None:
    content = ADAPTER.read(CatalogueFileFormat.YAML, b"systems:\n  - id: crm\n    name: CRM\n")

    assert content.products == ()


# Diff and evidence ------------------------------------------------------------------------


def test_the_changes_view_tells_components_from_responsibilities() -> None:
    base = seed_knowledge()
    added = diff_releases(base, _with_offering())
    assert [(item.item, item.change, item.label) for item in added.changes] == [
        (ChangedItem.PRODUCT, ChangeKind.ADDED, "Office Connect")
    ]

    moved = replace(
        OFFICE,
        version="1.1.0",
        components=(replace(FIBRE, responsibilities=FIBRE.responsibilities[:1]),),
    )
    changed = diff_releases(_with_offering(), replace(base, products=(moved,))).changes
    assert [(item.change, item.fields) for item in changed] == [
        (ChangeKind.CHANGED, ("version", "responsibilities"))
    ]


class _Recording(InMemoryEvidenceIndex):
    def __init__(self) -> None:
        super().__init__(FakeEmbeddings(), ApproximateTokenizer())
        self.chunks: tuple[EvidenceChunk, ...] = ()

    def store(self, release_id: str, index_id: str, chunks: tuple[EvidenceChunk, ...]) -> None:
        self.chunks = chunks
        super().store(release_id, index_id, chunks)


def test_evidence_names_the_systems_behind_each_component() -> None:
    repository = InMemoryArchitectureKnowledgeRepository(seed_knowledge())
    index = _Recording()
    manage = ManageArchitectureKnowledge(repository, CatalogueFileAdapter(), index, 10_000_000)
    draft = manage.create_draft(MAINTAINER, "Next version")
    draft = manage.update(draft.id, draft.revision, MAINTAINER, products=(OFFICE,))
    BuildArchitectureIndex(
        manage,
        index,
        InMemoryDocumentStorage(),
        LocatedDocumentExtractor(SafeDocumentTextExtractor()),
        ApproximateTokenizer(),
    ).execute(draft.id, draft.revision, "amina", fence=lambda: None)

    (chunk,) = [item for item in index.chunks if item.location == "product office-connect"]
    assert chunk.source_label == "Office Connect" and chunk.document_version_id is None
    assert chunk.text.startswith(
        "Product offering: Office Connect (OFFICE_CONNECT), Connect family"
    )
    assert "Order types: New Activation; Upgrade (not offered)" in chunk.text
    assert "Component Fibre Access (service, mandatory): Primary fixed connection." in chunk.text
    assert "Fibre Access — B2B Web (primary channel): Captures the order" in chunk.text


# API --------------------------------------------------------------------------------------


def test_a_draft_saves_offerings_and_refuses_removing_a_system_they_name(
    client: TestClient,
) -> None:
    draft = client.post(
        "/architecture-knowledge/releases", json={"name": "Next version"}, headers=OWNER
    ).json()
    offering = {
        "id": "office-connect",
        "name": "Office Connect",
        "order_types": [{"code": "NEW_ACTIVATION", "name": "New Activation"}],
        "components": [
            {
                "id": "fibre",
                "name": "Fibre Access",
                "mandatory": True,
                "responsibilities": [
                    {
                        "system_id": "bcrm",
                        "role": "Customer context",
                        "description": "Supplies the account.",
                        "order_types": ["NEW_ACTIVATION"],
                        "confidence": "confirmed",
                    }
                ],
            }
        ],
    }
    url = f"/architecture-knowledge/releases/{draft['id']}"
    body = {"systems": draft["systems"], "relationships": draft["relationships"]}

    saved = client.put(
        url,
        json={**body, "expected_revision": draft["revision"], "products": [offering]},
        headers=OWNER,
    )
    assert saved.status_code == 200, saved.text
    duty = saved.json()["products"][0]["components"][0]["responsibilities"][0]
    assert (duty["role"], duty["confidence"]) == ("CUSTOMER_CONTEXT", "confirmed")

    kept = client.put(
        url, json={**body, "expected_revision": saved.json()["revision"]}, headers=OWNER
    )
    assert [item["id"] for item in kept.json()["products"]] == ["office-connect"]

    without_bcrm = [item for item in draft["systems"] if item["id"] != "bcrm"]
    refused = client.put(
        url,
        json={
            "expected_revision": kept.json()["revision"],
            "systems": without_bcrm,
            "relationships": [
                item
                for item in draft["relationships"]
                if "bcrm" not in {item["source_system_id"], item["target_system_id"]}
            ],
        },
        headers=OWNER,
    )
    assert refused.status_code == 422
    assert "Office Connect › Fibre Access names system 'bcrm'" in refused.text
