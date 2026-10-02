"""Small real XLSX packages plus deliberately malformed merge metadata."""

import io
import zipfile
from xml.etree import ElementTree

from openpyxl import Workbook

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def spreadsheet_document(cells: dict[str, str | int], merges: tuple[str, ...] = ()) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Policy"
    for coordinate, value in cells.items():
        sheet[coordinate] = value
    for reference in merges:
        sheet.merge_cells(reference)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def reviewed_spreadsheet_document() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Policy"
    sheet["A1"] = "Private header"
    sheet.merge_cells("A1:C1")
    sheet["A2"] = "Private anchor"
    sheet["B2"] = "Private condition"
    sheet.merge_cells("A2:A3")
    sheet["B3"] = "XGPON التغطية مطلوبة. " * 80
    sheet["C3"] = "Public note"
    hidden = workbook.create_sheet("Internal")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "Private hidden rule"
    hidden.merge_cells("A1:B1")
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def with_merge_ranges(content: bytes, references: tuple[str, ...], *, dimension: str = "") -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(output, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(data)
                previous = root.find(f"{{{SHEET_NS}}}mergeCells")
                if previous is not None:
                    root.remove(previous)
                merged = ElementTree.SubElement(root, f"{{{SHEET_NS}}}mergeCells")
                for reference in references:
                    ElementTree.SubElement(merged, f"{{{SHEET_NS}}}mergeCell", ref=reference)
                if dimension:
                    dim = root.find(f"{{{SHEET_NS}}}dimension")
                    assert dim is not None
                    dim.set("ref", dimension)
                data = ElementTree.tostring(root)
            target.writestr(entry, data)
    return output.getvalue()


def reviewed_worksheet_names_document() -> bytes:
    """Names must be reviewed independently of the rows they locate."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Private heading سري"
    sheet["A1"] = "XGPON التغطية مطلوبة."
    sheet["B1"] = "Public note"
    workbook.create_sheet("Private empty")
    hidden = workbook.create_sheet("Private hidden")
    hidden.sheet_state = "veryHidden"
    hidden["A1"] = "Never publish appendix"
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()
